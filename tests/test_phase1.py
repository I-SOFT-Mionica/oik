"""Tests for core.phase1_generate."""

from __future__ import annotations

from pathlib import Path

import openpyxl

from biracki_odbor.config.oik import OikConstants, Municipality
from biracki_odbor.core.models import PollingStation
from biracki_odbor.core.phase1_generate import generate_party_files
from tests.fixtures.stalni_sastav import build_template


def _make_oik(n_stations: int = 3) -> OikConstants:
    return OikConstants(
        municipality=Municipality(
            name="Уб", name_genitive="Уба", name_locative="Убу",
            full_committee_name="ОПШТИНСКА ИЗБОРНА КОМИСИЈА ОПШТИНЕ УБ",
        ),
        polling_stations=[
            PollingStation(
                bm_number=i + 1,
                name=f"БМ {i+1}",
                address=f"Адреса {i+1}",
                full_name=f"БМ {i+1} (пуни назив)",
            )
            for i in range(n_stations)
        ],
        lock_password="test1234",
    )


def test_generates_one_file_per_party(tmp_path: Path) -> None:
    template = build_template(tmp_path / "t.xlsx", n_polling_stations=3,
                              parties=("СНС", "СПС", "ДС"))
    out = tmp_path / "stranke"
    result = generate_party_files(template, out, _make_oik(3))
    assert result.report.ok, [i.message_sr for i in result.report.issues]
    assert len(result.files_written) == 3
    for p in result.files_written:
        assert p.exists() and p.suffix == ".xlsx"


def test_locked_cells_protected_by_password(tmp_path: Path) -> None:
    template = build_template(tmp_path / "t.xlsx", n_polling_stations=2,
                              parties=("СНС", "СПС"))
    out = tmp_path / "stranke"
    oik = _make_oik(2)
    oik.lock_password = "secret9"
    result = generate_party_files(template, out, oik)
    party_file = result.files_written[0]
    wb = openpyxl.load_workbook(party_file)
    ws = wb.active
    assert ws.protection.sheet is True
    # openpyxl hashes the password, so we just check sheet is protected.
    wb.close()


def test_missing_bm_in_oik_warns(tmp_path: Path) -> None:
    template = build_template(tmp_path / "t.xlsx", n_polling_stations=3,
                              parties=("СНС",))
    out = tmp_path / "stranke"
    oik = _make_oik(1)  # template has 3 BMs, OIK only knows 1
    result = generate_party_files(template, out, oik)
    from biracki_odbor.validation.reports import Codes
    codes = {i.code for i in result.report.issues}
    assert Codes.BM_UNKNOWN in codes
    # Files still written; this is a warning, not an error
    assert len(result.files_written) == 1


def test_invalid_template_returns_report(tmp_path: Path) -> None:
    out = tmp_path / "stranke"
    result = generate_party_files(tmp_path / "nope.xlsx", out, _make_oik())
    assert not result.report.ok
    assert result.files_written == []


def test_party_file_has_correct_structure(tmp_path: Path) -> None:
    template = build_template(tmp_path / "t.xlsx", n_polling_stations=2,
                              parties=("СНС",))
    result = generate_party_files(template, tmp_path / "stranke", _make_oik(2))
    party_file = result.files_written[0]
    wb = openpyxl.load_workbook(party_file)
    ws = wb.active
    # Header row at row 3
    assert ws.cell(row=3, column=1).value == "Р.бр."
    assert ws.cell(row=3, column=2).value == "БМ бр."
    assert ws.cell(row=3, column=4).value == "Функција"
    # Schema post-split: E=Име, F=Презиме, G=ЈМБГ
    assert ws.cell(row=3, column=5).value == "Име"
    assert ws.cell(row=3, column=6).value == "Презиме"
    assert ws.cell(row=3, column=7).value == "ЈМБГ"
    # Row 4 starts data — first slot is (BM 1, Председник) → r.br 1
    assert ws.cell(row=4, column=1).value == 1
    # Party in column 12 (was 11; shifted by 1 for the new Презиме column)
    assert ws.cell(row=4, column=12).value == "СНС"
    # Top-of-sheet summary block (B6): row 1 labels in A:C, row 2 values.
    assert ws.cell(row=1, column=1).value == "Укупно"
    assert ws.cell(row=1, column=2).value == "Попуњени"
    assert ws.cell(row=1, column=3).value == "Неисправни"
    # A2 = total slot count for the single party in this fixture (2 BMs × 6 funcs).
    assert ws.cell(row=2, column=1).value == 12
    # B2 / C2 are COUNTA / COUNTIF formulas pointing at the data range.
    assert "COUNTA" in str(ws.cell(row=2, column=2).value)
    assert "COUNTIF" in str(ws.cell(row=2, column=3).value)
    assert "✗" in str(ws.cell(row=2, column=3).value)
    # Paste-resistant validity column at N (col 14): header + locked formula
    assert ws.cell(row=3, column=14).value == "Провера"
    n4 = ws.cell(row=4, column=14)
    assert isinstance(n4.value, str)
    assert n4.value.startswith("=IF(G4=")
    assert "✓" in n4.value and "✗" in n4.value
    assert n4.protection.locked is True
    # `hidden=True` keeps the formula off the formula bar so party clerks
    # can't read the underlying JMBG-validation logic.
    assert n4.protection.hidden is True
    wb.close()


def test_combined_name_header_off_by_default_keeps_plain_labels(tmp_path: Path) -> None:
    from biracki_odbor.config.election import ElectionConfig

    template = build_template(tmp_path / "t.xlsx", n_polling_stations=1, parties=("СНС",))
    result = generate_party_files(
        template, tmp_path / "stranke", _make_oik(1),
        election=ElectionConfig(),  # combined_name_header defaults to False
    )
    ws = openpyxl.load_workbook(result.files_written[0]).active
    assert ws.cell(row=3, column=5).value == "Име"
    assert ws.cell(row=3, column=6).value == "Презиме"
    wb = ws.parent
    wb.close()


def test_combined_name_header_truly_merges_ime_prezime_cells(tmp_path: Path) -> None:
    """Official forms list "Име и презиме" as one field — explicitly
    requested, accepted tradeoff (3+-word names can't auto-split, see
    popunjeni_io tests) over the earlier cosmetic-only version that kept
    two real columns behind a shared-looking header. E:F are a genuine
    merged range on every data row; everything after (ЈМБГ at G, phone
    at J, the hidden marker) stays at its original column, unmoved."""
    from biracki_odbor.config.election import ElectionConfig

    template = build_template(tmp_path / "t.xlsx", n_polling_stations=1, parties=("СNS",))
    result = generate_party_files(
        template, tmp_path / "stranke", _make_oik(1),
        election=ElectionConfig(combined_name_header=True),
    )
    wb = openpyxl.load_workbook(result.files_written[0])
    ws = wb.active
    assert ws.cell(row=3, column=5).value == "Име и презиме"
    assert ws.cell(row=3, column=6).value is None
    assert "E3:F3" in {str(r) for r in ws.merged_cells.ranges}
    assert "E4:F4" in {str(r) for r in ws.merged_cells.ranges}
    # Nothing downstream moved.
    assert ws.cell(row=3, column=7).value == "ЈМБГ"
    wb.close()


def test_combined_name_header_off_keeps_the_old_two_column_layout(tmp_path: Path) -> None:
    """The default (off) path must stay completely unmerged — pins that
    the merge logic is gated correctly, not applied unconditionally."""
    from biracki_odbor.config.election import ElectionConfig

    template = build_template(tmp_path / "t.xlsx", n_polling_stations=1, parties=("СНС",))
    result = generate_party_files(
        template, tmp_path / "stranke", _make_oik(1),
        election=ElectionConfig(combined_name_header=False),
    )
    wb = openpyxl.load_workbook(result.files_written[0])
    ws = wb.active
    assert "E4:F4" not in {str(r) for r in ws.merged_cells.ranges}
    assert ws.cell(row=3, column=6).value == "Презиме"
    wb.close()


def test_combined_name_header_round_trips_through_election_config_json() -> None:
    from biracki_odbor.config.election import ElectionConfig

    saved = ElectionConfig(combined_name_header=True)
    loaded = ElectionConfig.from_dict(saved.to_dict())

    assert loaded.combined_name_header is True


def test_combined_name_header_defaults_false_for_pre_existing_configs() -> None:
    """A config saved before this toggle existed has no such key at
    all — must not crash, must default to the old (unmerged) layout."""
    from biracki_odbor.config.election import ElectionConfig

    loaded = ElectionConfig.from_dict({"election_uuid": "abc123"})

    assert loaded.combined_name_header is False


def test_sheet_protection_is_hardened(tmp_path: Path) -> None:
    """Format/insert/delete actions must be blocked (True in OOXML).

    Previously the booleans were inverted, so users could format columns
    and unhide the hidden marker column M, exposing the UID stamp.
    """
    template = build_template(tmp_path / "t.xlsx", n_polling_stations=1,
                              parties=("СНС",))
    result = generate_party_files(template, tmp_path / "stranke", _make_oik(1))
    wb = openpyxl.load_workbook(result.files_written[0])
    ws = wb.active
    prot = ws.protection
    assert prot.sheet is True
    # True = locked in OOXML; user can't perform the action.
    assert prot.formatCells is True
    assert prot.formatColumns is True   # blocks "Unhide column" → marker stays hidden
    assert prot.formatRows is True
    assert prot.insertColumns is True
    assert prot.insertRows is True
    assert prot.deleteColumns is True
    assert prot.deleteRows is True
    # Sort/filter stay allowed for the party's convenience.
    assert prot.sort is False
    assert prot.autoFilter is False
    wb.close()


def test_marker_cell_is_hidden_and_locked(tmp_path: Path) -> None:
    """M1 must be both locked and hidden so the UID can't be peeked at
    even if the user somehow selects the cell."""
    from biracki_odbor.config.election import (
        DocumentInfo, ElectionConfig, ElectionDates, ElectionInfo,
    )
    election = ElectionConfig(
        election=ElectionInfo(
            dates=ElectionDates(
                date_full="17. децембар 2026. године",
                session_date="6. децембра 2026. године",
                session_number="5",
            ),
            types=["parliamentary"],
            rik_decision="X",
        ),
        document=DocumentInfo(number_prefix="013-5-30", legal_basis="Y"),
        party_full_names={"СНС": "A"},
        election_uuid="hardening-uuid",
    )
    template = build_template(tmp_path / "t.xlsx", n_polling_stations=1,
                              parties=("СНС",))
    result = generate_party_files(template, tmp_path / "stranke", _make_oik(1),
                                  election=election)
    wb = openpyxl.load_workbook(result.files_written[0])
    ws = wb.active
    m1 = ws["M1"]
    assert m1.value and m1.value.startswith("BO/hardening-uuid/")
    assert m1.protection.locked is True
    assert m1.protection.hidden is True
    wb.close()


def test_phone_validation_is_custom_with_leading_zero_rule(tmp_path: Path) -> None:
    """Phone column J must enforce digits-only and starting with '0'."""
    template = build_template(tmp_path / "t.xlsx", n_polling_stations=1,
                              parties=("СНС",))
    result = generate_party_files(template, tmp_path / "stranke", _make_oik(1))
    wb = openpyxl.load_workbook(result.files_written[0])
    ws = wb.active
    # Find the phone DV by inspecting its formula.
    phone_dvs = [dv for dv in ws.data_validations.dataValidation
                 if dv.formula1 and 'LEFT(J' in dv.formula1]
    assert phone_dvs, "expected a phone DV with LEFT-prefix check"
    dv = phone_dvs[0]
    assert dv.type == "custom"
    assert 'ISNUMBER' in dv.formula1
    assert 'LEFT(J4,1)="0"' in dv.formula1
    # And the cell itself is text-formatted so leading zero sticks.
    assert ws["J4"].number_format == "@"
    wb.close()


def test_email_validation_uses_at_and_dot(tmp_path: Path) -> None:
    template = build_template(tmp_path / "t.xlsx", n_polling_stations=1,
                              parties=("СНС",))
    result = generate_party_files(template, tmp_path / "stranke", _make_oik(1))
    wb = openpyxl.load_workbook(result.files_written[0])
    ws = wb.active
    email_dvs = [dv for dv in ws.data_validations.dataValidation
                 if dv.formula1 and 'SEARCH("@"' in dv.formula1]
    assert email_dvs, "expected an email DV with @-search"
    dv = email_dvs[0]
    assert dv.type == "custom"
    assert 'SEARCH(".",K4' in dv.formula1
    wb.close()


def test_identification_columns_have_per_row_cf(tmp_path: Path) -> None:
    """CF on A:D must reference each row's own N cell, not N4 for everyone.

    Evolution of this test:
    - v0.3.1: rule used `=$N4="✓"` over the whole range. Field reports said
      every row tinted based on N4's state.
    - v0.4.0: tried `=INDIRECT("N"&ROW())="✓"`. Also failed in the field
      ("the indirect doesn't work as well", 2026-05-26).
    - v0.4.1: switched to per-row CF with fully-absolute references
      (`=$N$5="✓"` for row 5, `=$N$6="✓"` for row 6, etc.). One CF rule
      per (data row × {✓,✗}) — no relative adjustment for Excel to misread.
    - v0.5.1: still per-row, but expanded to 3 states (empty / ✓ / ✗) so
      the locked-context blue is also delivered via CF — see
      test_identification_columns_three_state_cf_with_paired_fonts.

    This test asserts the per-row shape: each data row has its own CF
    range AND its formula references that row's own N cell as an
    absolute reference.
    """
    template = build_template(tmp_path / "t.xlsx", n_polling_stations=1,
                              parties=("СНС",))
    result = generate_party_files(template, tmp_path / "stranke", _make_oik(1))
    wb = openpyxl.load_workbook(result.files_written[0])
    ws = wb.active
    # Single-station × 6 functions → data rows 4..9. Each row should have
    # its own CF range "A{r}:D{r}" with three rules (empty/✓/✗).
    # `ConditionalFormatting` keys repr as `<ConditionalFormatting A4:D4>`,
    # so we extract the range substring.
    found_rows: set[str] = set()
    for cf_range, rules in ws.conditional_formatting._cf_rules.items():
        key_repr = str(cf_range)
        for r in range(4, 10):
            target = f"A{r}:D{r}"
            if target in key_repr:
                found_rows.add(str(r))
                for rule in rules:
                    for formula in (rule.formula or []):
                        # Formula must reference this specific row's N cell.
                        assert f"$N${r}" in formula, (
                            f"row {r}'s CF should reference $N${r}; got {formula!r}"
                        )
                break
    assert found_rows == {"4", "5", "6", "7", "8", "9"}, (
        f"expected per-row CF for rows 4..9; got {sorted(found_rows)}"
    )
    wb.close()


def test_identification_columns_three_state_cf_with_paired_fonts(tmp_path: Path) -> None:
    """v0.5.1 fix for A:D rows that don't visibly tint.

    In v0.5.0 the formula `=$N$5="✓"` was correct (user verified: copying
    the same formula as a new CF rule on a regular cell tinted that cell
    fine), but A:D cells didn't visibly tint. Cause: A:D cells were
    written with an explicit `LOCKED_FILL` PatternFill, and openpyxl's
    CF DifferentialStyle doesn't fully repaint over a solid native fill.

    Fix: drop the static fill+font on cols 1-4 entirely, and drive ALL
    three visual states (empty, ✓, ✗) from CF. The empty-state rule
    paints the "locked-context blue" so users still see at a glance
    which columns aren't theirs to edit.

    Each fill is paired with a font color so text contrast stays
    legible across all three backgrounds (Excel's Good/Bad/Neutral
    palette: dark blue on #D9E2F3, dark green on #C6EFCE, dark red on
    #FFC7CE).
    """
    template = build_template(tmp_path / "t.xlsx", n_polling_stations=1,
                              parties=("СНС",))
    result = generate_party_files(template, tmp_path / "stranke", _make_oik(1))
    wb = openpyxl.load_workbook(result.files_written[0])
    ws = wb.active

    # Cols 1-4 on a data row must have NO solid static fill. If someone
    # re-adds `c.fill = LOCKED_FILL` to that loop, the CF tint will
    # silently stop displaying again.
    for col in range(1, 5):
        cell = ws.cell(row=4, column=col)
        fill = cell.fill
        # openpyxl reports an unset fill as patternType=None.
        assert fill.patternType in (None, "none"), (
            f"col {col} row 4 must have no static fill (got {fill.patternType!r})"
            " — CF DifferentialStyle won't fully repaint over a solid fill,"
            " which was the v0.5.0 A:D tinting bug."
        )

    # Each data row should carry 3 CF rules (empty / ✓ / ✗).
    for cf_range, rules in ws.conditional_formatting._cf_rules.items():
        key_repr = str(cf_range)
        if "A4:D4" not in key_repr:
            continue
        formulas = [f for rule in rules for f in (rule.formula or [])]
        assert len(rules) == 3, (
            f"row 4 should have 3 CF rules (empty/✓/✗); got {len(rules)}: {formulas}"
        )

        # Map state → font color (read from the rule's DifferentialStyle).
        # FormulaRule stores font on rule.dxf.font.
        state_colors: dict[str, str] = {}
        for rule in rules:
            formula = (rule.formula or [""])[0]
            if '=""' in formula:
                state = "empty"
            elif '="✓"' in formula:
                state = "valid"
            elif '="✗"' in formula:
                state = "invalid"
            else:
                continue
            assert rule.dxf is not None and rule.dxf.font is not None, (
                f"{state} rule must pair font with fill for legible contrast;"
                " dark text on green/red is jarring"
            )
            state_colors[state] = rule.dxf.font.color.value
        assert state_colors.get("empty") == "002F5496", (
            f"empty-state font must be dark blue (00 + 2F5496); got {state_colors.get('empty')!r}"
        )
        assert state_colors.get("valid") == "00006100", (
            f"valid-state font must be dark green; got {state_colors.get('valid')!r}"
        )
        assert state_colors.get("invalid") == "009C0006", (
            f"invalid-state font must be dark red; got {state_colors.get('invalid')!r}"
        )
        break
    else:
        raise AssertionError("no CF rules found on A4:D4")

    wb.close()


def test_name_and_lastname_have_sentence_case_warning(tmp_path: Path) -> None:
    template = build_template(tmp_path / "t.xlsx", n_polling_stations=1,
                              parties=("СНС",))
    result = generate_party_files(template, tmp_path / "stranke", _make_oik(1))
    wb = openpyxl.load_workbook(result.files_written[0])
    ws = wb.active
    name_dvs = [dv for dv in ws.data_validations.dataValidation
                if dv.formula1 and 'PROPER(E' in dv.formula1]
    last_dvs = [dv for dv in ws.data_validations.dataValidation
                if dv.formula1 and 'PROPER(F' in dv.formula1]
    assert name_dvs, "expected a name DV with PROPER check"
    assert last_dvs, "expected a last-name DV with PROPER check"
    # Warning style so paste-from-CRM with weird caps doesn't hard-block.
    assert name_dvs[0].errorStyle == "warning"
    assert last_dvs[0].errorStyle == "warning"
    wb.close()


def test_slot_rbr_is_global_uid() -> None:
    """r.br must encode (bm, function) deterministically so the same slot
    gets the same number in master AND in any party file containing it."""
    from biracki_odbor.core.models import slot_rbr

    # BM 1 occupies r.br 1..6
    assert slot_rbr(1, "Председник") == 1
    assert slot_rbr(1, "Заменик члана 2") == 6
    # BM 2 occupies r.br 7..12 — the user-stated invariant
    assert slot_rbr(2, "Председник") == 7
    assert slot_rbr(2, "Заменик члана 2") == 12
    # BM 3 occupies r.br 13..18
    assert slot_rbr(3, "Председник") == 13


def test_rbr_consistent_across_party_files(tmp_path: Path) -> None:
    """Same (BM, function) slot → same r.br in every party file that owns it."""
    template = build_template(tmp_path / "t.xlsx", n_polling_stations=3,
                              parties=("СНС", "СПС", "ДС"))
    result = generate_party_files(template, tmp_path / "stranke", _make_oik(3))

    # Build {(bm, function): rbr} from each file, then check they agree.
    seen: dict[tuple[int, str], int] = {}
    for fp in result.files_written:
        wb = openpyxl.load_workbook(fp)
        ws = wb.active
        for r in range(4, ws.max_row + 1):
            bm = ws.cell(row=r, column=2).value
            func = ws.cell(row=r, column=4).value
            rbr = ws.cell(row=r, column=1).value
            if bm is None or func is None:
                break
            key = (int(bm), str(func))
            if key in seen:
                assert seen[key] == rbr, (
                    f"slot {key} got r.br {rbr} in {fp.name}, "
                    f"expected {seen[key]} (already seen in another file)"
                )
            else:
                seen[key] = rbr
        wb.close()

    # And the formula holds at the boundary: BM 2 Председник = 7.
    assert seen[(2, "Председник")] == 7


def test_generate_party_files_uses_process_pool_above_six_parties(tmp_path: Path) -> None:
    """Field report: "Покрени" on Стални Phase 1 crashed with
    "too many values to unpack (expected 7)" as soon as a real
    municipality's RIK import produced 6+ parties. Root cause: the
    ProcessPoolExecutor branch (only taken when max_workers > 1 AND
    total >= 6 -- see the "Below ~6 parties the pool is a net loss"
    comment) unpacked each task tuple into exactly 7 names, but the
    tuple itself carries 11 elements (title_prefix/rbr_list/
    display_name/combined_name_header were added to _write_party_workbook
    later without updating this branch to match). Every test before this
    one used 3 or fewer parties, so the pool branch was never actually
    exercised. max_workers=2 here forces it deterministically regardless
    of this machine's CPU count."""
    template = build_template(tmp_path / "t.xlsx", n_polling_stations=2)  # 6 parties (fixture default)
    out = tmp_path / "stranke"
    result = generate_party_files(template, out, _make_oik(2), max_workers=2)
    assert result.report.ok, [i.message_sr for i in result.report.issues]
    assert len(result.files_written) == 6
    for p in result.files_written:
        assert p.exists() and p.suffix == ".xlsx"
