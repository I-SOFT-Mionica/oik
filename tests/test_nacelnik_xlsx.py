"""Tests for core.nacelnik_xlsx — the chief's takeover Excel template.

Covers the column-layout change: "Странка" (which party originally held
the slot) was dropped — the chief doesn't need or want to know that,
it's confusing/irrelevant to their job of just filling the slot — and
"Р.бр." (slot_rbr, the same canonical numbering used in master.xlsx and
every per-party workbook) was added. Slot matching itself is via
(bm_number, function), read per row — never row position or Р.бр. — so
sorting/reordering rows in the file can never break the match.
"""

from __future__ import annotations

from pathlib import Path

from biracki_odbor.core.models import Person, PollingStation, slot_rbr
from biracki_odbor.core.nacelnik_xlsx import (
    _ALL_COLS, export_nacelnik_template, read_nacelnik_xlsx,
)


def _stations() -> dict[int, PollingStation]:
    return {
        1: PollingStation(bm_number=1, name="Школа А"),
        2: PollingStation(bm_number=2, name="Школа Б"),
    }


def test_export_headers_have_rbr_and_no_party_column(tmp_path: Path) -> None:
    out = export_nacelnik_template(
        [Person(bm_number=1, function="Председник", party="СНС")],
        _stations(), tmp_path / "n.xlsx",
    )
    import openpyxl
    wb = openpyxl.load_workbook(out, read_only=True)
    header = [c.value for c in next(wb.active.iter_rows(min_row=1, max_row=1))]
    wb.close()
    assert header == _ALL_COLS
    assert header[0] == "Р.бр."
    assert "Странка" not in header


def test_export_rbr_matches_canonical_slot_rbr(tmp_path: Path) -> None:
    persons = [
        Person(bm_number=2, function="Члан 1", party="СПС"),
        Person(bm_number=1, function="Председник", party="СНС"),
    ]
    out = export_nacelnik_template(persons, _stations(), tmp_path / "n.xlsx")
    import openpyxl
    wb = openpyxl.load_workbook(out, read_only=True)
    rows = list(wb.active.iter_rows(min_row=3, values_only=True))
    wb.close()
    by_slot = {(r[1], r[3]): r[0] for r in rows if r[1] is not None}
    assert by_slot[(1, "Председник")] == slot_rbr(1, "Председник")
    assert by_slot[(2, "Члан 1")] == slot_rbr(2, "Члан 1")


def test_roundtrip_matches_by_bm_and_function_not_row_position(tmp_path: Path) -> None:
    """Rows reordered (e.g. by an Excel sort — explicitly allowed by the
    sheet protection settings) must still match by their own (bm,
    function) cell values, never by row position or Р.бр."""
    persons = [
        Person(bm_number=1, function="Председник", party="СНС"),
        Person(bm_number=2, function="Члан 1", party="СПС"),
    ]
    path = tmp_path / "n.xlsx"
    export_nacelnik_template(persons, _stations(), path)

    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws.protection.sheet = False  # test only needs to write, not respect the lock

    # Export sorts by (bm, function), so row 3 = БМ1/Председник, row 4 =
    # БМ2/Члан 1. Swap the locked context cells between the two rows —
    # simulating a descending sort — then fill in the (now-swapped) БМ1
    # row, which physically sits at row 4.
    for col in range(1, 5):
        r3, r4 = ws.cell(row=3, column=col), ws.cell(row=4, column=col)
        r3.value, r4.value = r4.value, r3.value
    ws.cell(row=4, column=5, value="Ана")   # Ime, on the row now holding БМ1
    ws.cell(row=4, column=6, value="Анић")  # Prezime
    wb.save(path)
    wb.close()

    rows, errors = read_nacelnik_xlsx(path)
    filled = [r for r in rows if r.is_filled]
    assert len(filled) == 1
    assert filled[0].bm_number == 1
    assert filled[0].function == "Председник"
    assert filled[0].first_name == "Ана"
    # The unfilled BM 2 slot (now physically at row 3) is still reported.
    assert any("2" in e and "Члан 1" in e for e in errors)


def test_nacelnik_row_has_no_party_field() -> None:
    from biracki_odbor.core.nacelnik_xlsx import NacelnikRow
    import dataclasses
    fields = {f.name for f in dataclasses.fields(NacelnikRow)}
    assert "party" not in fields


def test_jmbg_and_telefon_columns_are_text_formatted(tmp_path: Path) -> None:
    """Same fix as the regular party-submission templates
    (phase1_generate.py): without an explicit text format, Excel treats
    a leading-zero JMBG/phone as a number and silently drops the 0 the
    moment the chief types it in. Column 7 = JMBG, column 10 = Telefon
    (see _INPUT_COLS)."""
    out = export_nacelnik_template(
        [Person(bm_number=1, function="Председник", party="СНС")],
        _stations(), tmp_path / "n.xlsx",
    )
    import openpyxl
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    assert ws.cell(row=3, column=7).number_format == "@"
    assert ws.cell(row=3, column=10).number_format == "@"
    # Mesto (col 8) is ordinary text anyway -- must NOT be force-formatted,
    # confirms the fix is scoped to the two numeric-looking columns only.
    assert ws.cell(row=3, column=8).number_format != "@"
    wb.close()


def test_jmbg_and_telefon_have_inline_data_validation(tmp_path: Path) -> None:
    """Same in-Excel error popup as the regular templates -- catches a
    bad entry the moment the chief types it, not just on import."""
    out = export_nacelnik_template(
        [Person(bm_number=1, function="Председник", party="СНС")],
        _stations(), tmp_path / "n.xlsx",
    )
    import openpyxl
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    titles = {dv.errorTitle for dv in ws.data_validations.dataValidation}
    assert "Неисправан ЈМБГ" in titles
    assert "Број телефона" in titles
    wb.close()


def test_no_data_validation_crash_on_an_empty_slot_list(tmp_path: Path) -> None:
    """export_nacelnik_template([], ...) -- no empty slots to fill --
    must not crash computing a last_data_row from zero rows."""
    out = export_nacelnik_template([], _stations(), tmp_path / "n.xlsx")
    assert out.exists()


def _election(*, combined_name_header: bool):
    from biracki_odbor.config.election import ElectionConfig
    ec = ElectionConfig()
    ec.combined_name_header = combined_name_header
    return ec


def test_validity_column_formula_reads_this_rows_jmbg(tmp_path: Path) -> None:
    """The paste-resistant 'Провера' column is the real fix for '0 in
    JMBG silently stripped' -- DataValidation alone doesn't fire on a
    paste, this formula-driven column does."""
    out = export_nacelnik_template(
        [Person(bm_number=1, function="Председник", party="СНС")],
        _stations(), tmp_path / "n.xlsx",
    )
    import openpyxl
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    assert ws.cell(row=1, column=12).value == "Провера"
    formula = ws.cell(row=3, column=12).value
    assert formula.startswith("=IF(G3=")
    assert ws.cell(row=3, column=12).protection.hidden is True
    wb.close()


def test_locked_context_columns_have_no_static_fill(tmp_path: Path) -> None:
    """A static PatternFill on the locked columns blocks openpyxl's
    conditional-formatting DifferentialStyle from fully repainting them
    -- the exact bug phase1_generate.py's history documents. The
    empty/valid/invalid row coloring must come entirely from CF."""
    out = export_nacelnik_template(
        [Person(bm_number=1, function="Председник", party="СНС")],
        _stations(), tmp_path / "n.xlsx",
    )
    import openpyxl
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    for col in range(1, 5):
        assert ws.cell(row=3, column=col).fill.fgColor.rgb in (None, "00000000")
    wb.close()


def test_conditional_formatting_rules_cover_jmbg_and_validity_column(tmp_path: Path) -> None:
    out = export_nacelnik_template(
        [Person(bm_number=1, function="Председник", party="СНС")],
        _stations(), tmp_path / "n.xlsx",
    )
    import openpyxl
    wb = openpyxl.load_workbook(out)
    ws = wb.active
    ranges = {str(r) for r in ws.conditional_formatting._cf_rules.keys()}
    assert any("G3" in r for r in ranges)   # JMBG column
    assert any("L3" in r for r in ranges)   # Провера column
    assert any("A3:D3" in r for r in ranges)  # row-level context coloring
    wb.close()


def test_combined_name_header_merges_and_round_trips(tmp_path: Path) -> None:
    """Same election.combined_name_header toggle the party templates
    already use -- this is the same kind of file, just a different
    assigning mechanism."""
    persons = [Person(bm_number=1, function="Председник", party="СНС")]
    path = tmp_path / "n.xlsx"
    export_nacelnik_template(
        persons, _stations(), path, election=_election(combined_name_header=True),
    )

    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    assert ws.cell(row=1, column=5).value == "Име и презиме"
    assert ws.cell(row=1, column=6).value is None
    assert "E1:F1" in ws.merged_cells
    ws.protection.sheet = False
    ws.cell(row=3, column=5, value="Ана Анић")
    wb.save(path)
    wb.close()

    rows, errors = read_nacelnik_xlsx(path)
    filled = [r for r in rows if r.is_filled]
    assert len(filled) == 1
    assert filled[0].first_name == "Ана"
    assert filled[0].last_name == "Анић"
    assert not any("заглавља" in e for e in errors)


def test_combined_name_header_three_word_name_splits_first_word_as_given_name(tmp_path: Path) -> None:
    """3+ words: first word is the given name, the rest joins as the
    surname -- never left blank, same user-specified behavior as
    popunjeni_io.parse_party_file for the party files."""
    persons = [Person(bm_number=1, function="Председник", party="СНС")]
    path = tmp_path / "n.xlsx"
    export_nacelnik_template(
        persons, _stations(), path, election=_election(combined_name_header=True),
    )
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws.protection.sheet = False
    ws.cell(row=3, column=5, value="Ана Марија Анић")
    wb.save(path)
    wb.close()

    rows, _errors = read_nacelnik_xlsx(path)
    filled = [r for r in rows if r.is_filled]
    assert filled[0].first_name == "Ана"
    assert filled[0].last_name == "Марија Анић"


def test_legacy_11_column_file_without_provera_still_imports(tmp_path: Path) -> None:
    """Files generated before the Провера column / combined_name_header
    support was added must keep importing -- explicit user requirement,
    not just a nice-to-have. Hand-builds the OLD (pre-change) shape
    directly rather than via export_nacelnik_template, since that
    function only ever writes the current shape now."""
    import openpyxl
    from biracki_odbor.core.nacelnik_xlsx import _LEGACY_ALL_COLS

    path = tmp_path / "legacy.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for col_idx, label in enumerate(_LEGACY_ALL_COLS, start=1):
        ws.cell(row=1, column=col_idx, value=label)
    ws.cell(row=2, column=1, value="Упутство...")
    ws.cell(row=3, column=1, value=1)
    ws.cell(row=3, column=2, value=1)
    ws.cell(row=3, column=3, value="Школа А")
    ws.cell(row=3, column=4, value="Председник")
    ws.cell(row=3, column=5, value="Ана")
    ws.cell(row=3, column=6, value="Анић")
    ws.cell(row=3, column=7, value="0101990800019")
    wb.save(path)
    wb.close()

    rows, errors = read_nacelnik_xlsx(path)
    filled = [r for r in rows if r.is_filled]
    assert len(filled) == 1
    assert filled[0].first_name == "Ана"
    assert filled[0].last_name == "Анић"
    assert not any("заглавља" in e for e in errors)


def test_split_name_file_still_reads_correctly_when_combined_header_is_off(tmp_path: Path) -> None:
    """Exporting without `election` (or with the toggle off) keeps the
    original split Ime/Prezime columns -- default/legacy-caller shape
    must still round-trip."""
    persons = [Person(bm_number=1, function="Председник", party="СНС")]
    path = tmp_path / "n.xlsx"
    export_nacelnik_template(
        persons, _stations(), path, election=_election(combined_name_header=False),
    )
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws.protection.sheet = False
    ws.cell(row=3, column=5, value="Ана")
    ws.cell(row=3, column=6, value="Анић")
    wb.save(path)
    wb.close()

    rows, errors = read_nacelnik_xlsx(path)
    filled = [r for r in rows if r.is_filled]
    assert filled[0].first_name == "Ана"
    assert filled[0].last_name == "Анић"
    assert not errors
