"""Tests for core.phase3_docx — basic structural checks on the .docx output."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from biracki_odbor.config.election import ElectionConfig, ElectionInfo, ElectionDates, DocumentInfo
from biracki_odbor.core.models import Person
from biracki_odbor.core.phase3_docx import generate_documents
from biracki_odbor.core.workflow import PROSIRENI_SPEC
from biracki_odbor.validation.reports import Codes
from tests.test_jmbg import VALID_SAMPLES
from tests.test_phase1 import _make_oik


def _make_election() -> ElectionConfig:
    return ElectionConfig(
        election=ElectionInfo(
            dates=ElectionDates(
                date_full="17. децембар 2026. године",
                session_date="6. децембра 2026. године",
                session_number="5",
            ),
            types=["parliamentary"],
            rik_decision="Одлуком РИК-а 02 број 013-376/26",
        ),
        document=DocumentInfo(
            number_prefix="013-5-30",
            legal_basis="На основу члана 32. Закона ...",
        ),
        party_full_names={
            "СНС": "Посланичке групе Србија",
        },
    )


def _make_persons() -> list[Person]:
    """Two BMs, full slate of 6 members each."""
    funcs = [
        "Председник", "Заменик председника",
        "Члан 1", "Заменик члана 1",
        "Члан 2", "Заменик члана 2",
    ]
    persons = []
    for bm in (1, 2):
        for i, func in enumerate(funcs):
            persons.append(Person(
                bm_number=bm, function=func, party="СНС",
                first_name=f"Особа{bm}", last_name=f"Презименић{i + 1}",
                jmbg=VALID_SAMPLES[i % len(VALID_SAMPLES)],
                mesto="Уб", adresa="Адреса", telefon="0641234567",
            ))
    return persons


def test_normal_style_uses_single_line_spacing(tmp_path: Path) -> None:
    """python-docx's blank-document template defaults Normal to 1.15
    ("Вишеструки") line spacing -- never overridden, so every paragraph
    that didn't pin its own spacing silently inherited that instead of
    single spacing. Caught via live testing (Word's "Поставити
    подразумеване вредности" dialog showing Вишеструки/1.15)."""
    from docx.enum.text import WD_LINE_SPACING

    out = tmp_path / "resenja"
    result = generate_documents(_make_persons(), _make_oik(1), _make_election(), out)
    doc = Document(result.files_written[0])
    assert doc.styles["Normal"].paragraph_format.line_spacing_rule == WD_LINE_SPACING.SINGLE


def test_header_and_footer_distance_is_1cm(tmp_path: Path) -> None:
    out = tmp_path / "resenja"
    result = generate_documents(_make_persons(), _make_oik(1), _make_election(), out)
    doc = Document(result.files_written[0])
    section = doc.sections[0]
    assert section.header_distance.cm == pytest.approx(1.0, abs=0.01)
    assert section.footer_distance.cm == pytest.approx(1.0, abs=0.01)


def test_document_format_is_operator_configurable(tmp_path: Path) -> None:
    """Поставке ОИК → 3. Систем → Формат докумената — not just a fixed
    default: a non-default DocumentFormat on the OIK must actually change
    the generated document, not just the (already-covered) default case."""
    from docx.enum.text import WD_LINE_SPACING

    out = tmp_path / "resenja"
    oik = _make_oik(1)
    oik.document_format.line_spacing = "double"
    oik.document_format.header_distance_cm = 2.0
    oik.document_format.footer_distance_cm = 1.5
    result = generate_documents(_make_persons(), oik, _make_election(), out)
    doc = Document(result.files_written[0])
    assert doc.styles["Normal"].paragraph_format.line_spacing_rule == WD_LINE_SPACING.DOUBLE
    section = doc.sections[0]
    assert section.header_distance.cm == pytest.approx(2.0, abs=0.01)
    assert section.footer_distance.cm == pytest.approx(1.5, abs=0.01)


def test_ordinary_body_paragraphs_use_the_configured_spacing(tmp_path: Path) -> None:
    """тачка 1, РИК provision, тачка 2/3, образложење, and the legal
    remedy all used a mix of hardcoded 3pt/6pt space_after (no evident
    pattern to which got which — the inconsistency a customer flagged)
    or, for the other two generators, silently fell through to python-
    docx's blank-document template default (10pt). All of these must now
    follow ONE configured value instead."""
    out = tmp_path / "resenja"
    oik = _make_oik(1)
    oik.document_format.body_before_pt = 4.0
    oik.document_format.body_after_pt = 8.0
    result = generate_documents(_make_persons(), oik, _make_election(), out)
    doc = Document(result.files_written[0])

    # Point to a paragraph that's definitely ordinary running text, not
    # part of the title block/headings/signature block -- the legal
    # remedy notice, already used as a fixed marker elsewhere in this file.
    body_para = next(p for p in doc.paragraphs if "УПУТСТВО О ПРАВНОМ СРЕДСТВУ" in p.text)
    assert body_para.paragraph_format.space_after is None  # inherits from style
    normal = doc.styles["Normal"].paragraph_format
    assert normal.space_before.pt == 4.0
    assert normal.space_after.pt == 8.0


def test_title_block_follows_its_own_separate_spacing_setting(tmp_path: Path) -> None:
    """Title-block lines (subtitle, composition mark, election type,
    РАСПИСАНИХ ЗА, ЗА БИРАЧКО МЕСТО) respond to title_before_pt/
    title_after_pt -- a SEPARATE knob from body spacing, exposed next to
    it on the generation screen so the operator can shrink either one
    independently when a решење overflows onto an unwanted extra page."""
    out = tmp_path / "resenja"
    oik = _make_oik(1)
    oik.document_format.title_before_pt = 2.0
    oik.document_format.title_after_pt = 1.0
    # Body spacing must NOT leak into the title block.
    oik.document_format.body_before_pt = 20.0
    oik.document_format.body_after_pt = 20.0
    result = generate_documents(_make_persons(), oik, _make_election(), out)
    doc = Document(result.files_written[0])

    from biracki_odbor.core.workflow import STALNI_SPEC
    subtitle = next(
        p for p in doc.paragraphs if p.text == STALNI_SPEC.resenja_title_subtitle
    )
    assert subtitle.paragraph_format.space_before.pt == 2.0
    assert subtitle.paragraph_format.space_after.pt == 1.0


def test_main_heading_keeps_its_own_fixed_spacing_regardless_of_title_setting(
    tmp_path: Path,
) -> None:
    """The "РЕШЕЊЕ" heading itself is a fixed structural anchor, like the
    "Образложење" section heading and the signature block -- never
    collapsed into the title-block setting, even when that's reconfigured."""
    out = tmp_path / "resenja"
    oik = _make_oik(1)
    oik.document_format.title_before_pt = 20.0
    oik.document_format.title_after_pt = 20.0
    result = generate_documents(_make_persons(), oik, _make_election(), out)
    doc = Document(result.files_written[0])

    title_para = next(p for p in doc.paragraphs if p.text == "РЕШЕЊЕ")
    assert title_para.paragraph_format.space_before.pt == 12.0
    assert title_para.paragraph_format.space_after.pt == 12.0


def test_document_format_multiple_1_15_sets_both_rule_and_multiplier(tmp_path: Path) -> None:
    """The one choice that needs BOTH a rule (MULTIPLE) and a numeric
    multiplier (1.15), not just a fixed ratio -- the case the ordering
    bug (setting line_spacing=None after line_spacing_rule silently wiped
    the rule back out) would have broken first, since this is the only
    branch where line_spacing is actually set to a non-None value."""
    from docx.enum.text import WD_LINE_SPACING

    out = tmp_path / "resenja"
    oik = _make_oik(1)
    oik.document_format.line_spacing = "multiple_1_15"
    result = generate_documents(_make_persons(), oik, _make_election(), out)
    doc = Document(result.files_written[0])
    pf = doc.styles["Normal"].paragraph_format
    assert pf.line_spacing_rule == WD_LINE_SPACING.MULTIPLE
    assert pf.line_spacing == pytest.approx(1.15, abs=0.001)


def test_signature_block_uses_chairman_by_default(tmp_path: Path) -> None:
    out = tmp_path / "resenja"
    oik = _make_oik(1)
    oik.chairman.name = "Петар Петровић"
    oik.chairman.title = "ПРЕДСЕДНИК"
    result = generate_documents(_make_persons(), oik, _make_election(), out)
    text = "\n".join(p.text for p in Document(result.files_written[0]).paragraphs)
    assert "Петар Петровић" in text
    assert "ПРЕДСЕДНИК" in text


def test_signature_block_uses_delegated_signer_when_authorized(tmp_path: Path) -> None:
    """Delegation is ONE shared setting (config.election.DocumentInfo) —
    when a chairman has delegated signing authority for the period, the
    delegate signs the ORIGINAL appointment decision too, not just
    amendments — this used to be izmena-only."""
    out = tmp_path / "resenja"
    oik = _make_oik(1)
    oik.chairman.name = "Петар Петровић"
    election = _make_election()
    election.document.use_authorization = True
    election.document.authorized_name = "Ивана Тешић"
    election.document.authorized_role = "заменик председника"
    result = generate_documents(_make_persons(), oik, election, out)
    text = "\n".join(p.text for p in Document(result.files_written[0]).paragraphs)
    assert "Ивана Тешић" in text
    assert "заменик председника" in text
    assert "Петар Петровић" not in text


def test_signature_block_ignores_delegation_when_toggled_off(tmp_path: Path) -> None:
    out = tmp_path / "resenja"
    oik = _make_oik(1)
    oik.chairman.name = "Петар Петровић"
    election = _make_election()
    election.document.use_authorization = False
    election.document.authorized_name = "Ивана Тешић"
    result = generate_documents(_make_persons(), oik, election, out)
    text = "\n".join(p.text for p in Document(result.files_written[0]).paragraphs)
    assert "Петар Петровић" in text
    assert "Ивана Тешић" not in text


def test_generates_one_docx_per_bm(tmp_path: Path) -> None:
    out = tmp_path / "resenja"
    result = generate_documents(_make_persons(), _make_oik(2), _make_election(), out)
    assert len(result.files_written) == 2
    for p in result.files_written:
        assert p.exists() and p.suffix == ".docx"
        assert p.name.startswith("БМ_")
        assert "STALNI" in p.name


def test_docx_contains_expected_text(tmp_path: Path) -> None:
    out = tmp_path / "resenja"
    result = generate_documents(_make_persons(), _make_oik(2), _make_election(), out)
    doc = Document(result.files_written[0])
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "РЕШЕЊЕ" in all_text
    assert "Образложење" in all_text
    assert "БИРАЧКО МЕСТО БРОЈ 1" in all_text
    assert "Уб" in all_text


def test_member_names_present_in_table(tmp_path: Path) -> None:
    persons = _make_persons()
    out = tmp_path / "resenja"
    generate_documents(persons, _make_oik(2), _make_election(), out)
    doc = Document(out / "БМ_01_STALNI.docx")
    table_text = "\n".join(
        cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    )
    for i in range(6):
        # Members render as "{first_name} {last_name}" via Person.name property.
        assert f"Особа1 Презименић{i + 1}" in table_text


def test_unmatched_party_warning(tmp_path: Path) -> None:
    persons = _make_persons()
    # Introduce an unmatched party
    persons[0].party = "СПС"
    out = tmp_path / "resenja"
    result = generate_documents(persons, _make_oik(2), _make_election(), out)
    codes = {i.code for i in result.report.issues}
    assert Codes.PARTY_UNMATCHED in codes


def test_prosireni_uses_prosireni_obrazlozenje_not_stalni(tmp_path: Path, monkeypatch) -> None:
    """Field report: a prosireni решење's Образложење text was literally
    stalni's ("...сталном саставу...", "Чланом 35."), even though the
    decision's own title correctly said "У ПРОШИРЕНОМ САСТАВУ" —
    _build_document called obrazlozenje_mod.render() without a `path`,
    which always defaulted to stalni's obrazlozenje.json regardless of
    which spec was passed in."""
    monkeypatch.setenv("BIRACKI_ODBOR_HOME", str(tmp_path / "home"))
    out = tmp_path / "resenja"
    result = generate_documents(
        _make_persons(), _make_oik(2), _make_election(), out, spec=PROSIRENI_SPEC,
    )
    doc = Document(result.files_written[0])
    paras = [p.text for p in doc.paragraphs]
    obr_idx = next(i for i, t in enumerate(paras) if t.strip() == "Образложење")
    obrazlozenje_text = "\n".join(paras[obr_idx + 1:])
    assert "сталном саставу" not in obrazlozenje_text
    assert "Чланом 39. став 1." in obrazlozenje_text


def _make_prosireni_persons() -> list[Person]:
    """One BM, two participating entities (slots 1 and 2), each with a
    Члан + Заменик члана pair -- prosireni's actual function-string shape,
    unlike stalni's fixed Председник/Члан 1/Члан 2 schema."""
    from biracki_odbor.core.prosireni_allocation import (
        MEMBER_LABEL, SUBSTITUTE_LABEL, slot_function,
    )
    persons = []
    for slot_idx, party in ((1, "СНС"), (2, "СПС")):
        for j, (label, first) in enumerate((
            (MEMBER_LABEL, f"Члан{slot_idx}"), (SUBSTITUTE_LABEL, f"Заменик{slot_idx}"),
        )):
            persons.append(Person(
                bm_number=1, function=slot_function(label, slot_idx), party=party,
                first_name=first, last_name="Презименић",
                jmbg=VALID_SAMPLES[(slot_idx * 2 + j) % len(VALID_SAMPLES)],
                mesto="Уб", adresa="Адреса", telefon="0641234567",
            ))
    return persons


def test_prosireni_members_table_not_blank(tmp_path: Path, monkeypatch) -> None:
    """Field report: prosireni резенја generated with a completely empty
    members table. _build_document's table was hardcoded to stalni's 3-group
    (Председник/Члан 1/Члан 2) schema, so member_by_func.get(...) never
    matched prosireni's compound "Члан [N]"/"Заменик члана [N]" function
    strings and every row rendered blank. The table must now be derived
    from the actual slot indices present in the data."""
    monkeypatch.setenv("BIRACKI_ODBOR_HOME", str(tmp_path / "home"))
    out = tmp_path / "resenja"
    result = generate_documents(
        _make_prosireni_persons(), _make_oik(1), _make_election(), out,
        spec=PROSIRENI_SPEC,
    )
    doc = Document(result.files_written[0])
    table_text = "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    assert "Члан1 Презименић" in table_text
    assert "Заменик1 Презименић" in table_text
    assert "Члан2 Презименић" in table_text
    assert "Заменик2 Презименић" in table_text
    assert "члан" in table_text.lower()


def test_prosireni_numbering_is_slot_index_not_fill_order(tmp_path: Path, monkeypatch) -> None:
    """Domain rule, confirmed by the user: a list's член/заменик number IS
    its permanent ballot-order slot_index (see prosireni_allocation.
    canonical_slots). If list 2's and list 3's slots are never filled,
    "2." and "3." simply don't appear in the решење at all — list 6's
    member stays "6.", never renumbered down to fill the gap.

    An earlier version used enumerate()'s position over the sorted,
    already-filled-only slot indices as the printed number instead of
    the slot index itself — correct with no gaps (as in
    test_prosireni_members_table_not_blank's slots 1,2), but silently
    wrong the moment any slot in between was unfilled: with slots
    {1, 6} actually filled, it printed "1)" and "2)" — misattributing
    list 6's real seat to list 2's ballot number in the actual legal
    document."""
    from biracki_odbor.core.prosireni_allocation import (
        MEMBER_LABEL, SUBSTITUTE_LABEL, slot_function,
    )
    monkeypatch.setenv("BIRACKI_ODBOR_HOME", str(tmp_path / "home"))
    persons = []
    for slot_idx, party in ((1, "СНС"), (6, "СПС")):
        for label, first in (
            (MEMBER_LABEL, f"Члан{slot_idx}"), (SUBSTITUTE_LABEL, f"Заменик{slot_idx}"),
        ):
            persons.append(Person(
                bm_number=1, function=slot_function(label, slot_idx), party=party,
                first_name=first, last_name="Презименић",
                jmbg=VALID_SAMPLES[slot_idx % len(VALID_SAMPLES)],
                mesto="Уб", adresa="Адреса", telefon="0641234567",
            ))
    out = tmp_path / "resenja"
    result = generate_documents(
        persons, _make_oik(1), _make_election(), out, spec=PROSIRENI_SPEC,
    )
    doc = Document(result.files_written[0])
    table_text = [
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    ]
    assert "1)" in table_text
    assert "6)" in table_text
    assert "2)" not in table_text


def test_empty_bm_warning(tmp_path: Path) -> None:
    persons = _make_persons()
    # Strip all members of BM 2
    persons = [p for p in persons if p.bm_number == 1]
    # Add unfilled placeholder for BM 2
    persons.append(Person(bm_number=2, function="Председник", party="СНС"))
    out = tmp_path / "resenja"
    result = generate_documents(persons, _make_oik(2), _make_election(), out)
    codes = {i.code for i in result.report.issues}
    assert Codes.EMPTY_BM in codes


def test_docx_opens_cleanly(tmp_path: Path) -> None:
    """Sanity: a generated file should be readable by python-docx without error."""
    out = tmp_path / "resenja"
    result = generate_documents(_make_persons(), _make_oik(2), _make_election(), out)
    for p in result.files_written:
        d = Document(p)
        assert len(d.paragraphs) > 10
        # BM-name signature table + members table + stamp spacer.
        assert len(d.tables) == 3


def test_obrazlozenje_has_page_break_before(tmp_path: Path) -> None:
    """Образложење must start on page 2 — uses page_break_before, not 6 empty paragraphs."""
    out = tmp_path / "resenja"
    result = generate_documents(_make_persons(), _make_oik(2), _make_election(), out)
    doc = Document(result.files_written[0])
    found_pbb = False
    for p in doc.paragraphs:
        if "Образложење" in p.text or "О б р а з л о ж е њ е" in p.text:
            pbb = p.paragraph_format.page_break_before
            if pbb is True:
                found_pbb = True
            break
    assert found_pbb, "Образложење paragraph must have page_break_before=True"


def test_days_before_election_excludes_election_day(tmp_path: Path, monkeypatch) -> None:
    """Field report: a cutoff of Oct 9 for an Oct 25 election must render
    as "15 dana pre izbora", not 16 -- election day itself isn't counted
    as one of the "days before". Raw calendar-day subtraction
    (25 - 9 = 16) overcounts by exactly one."""
    import biracki_odbor.core.phase3_docx as phase3_docx_mod

    captured: dict = {}
    original_render = phase3_docx_mod.obrazlozenje_mod.render

    def spy_render(*args, **kwargs):
        captured["days_before_proposal"] = kwargs.get("days_before_proposal")
        captured["days_before_appointment"] = kwargs.get("days_before_appointment")
        return original_render(*args, **kwargs)

    monkeypatch.setattr(phase3_docx_mod.obrazlozenje_mod, "render", spy_render)

    election = ElectionConfig(
        election=ElectionInfo(
            dates=ElectionDates(
                date_full="25. октобра 2026. године",
                session_date="6. октобра 2026. године",
                session_number="5",
                cutoff_stalni="2026-10-09",
                cutoff_prosireni="2026-10-14",
            ),
            types=["parliamentary"],
            rik_decision="Одлуком РИК-а 02 број 013-376/26",
        ),
        document=DocumentInfo(
            number_prefix="013-5-30",
            legal_basis="На основу члана 32. Закона ...",
        ),
    )
    out = tmp_path / "resenja"
    generate_documents(_make_persons(), _make_oik(2), election, out)

    assert captured["days_before_proposal"] == 15
    assert captured["days_before_appointment"] == 10


def test_bm_table_only_middle_line_visible(tmp_path: Path) -> None:
    """BM table renders as a single underline; only the top cell's bottom border draws."""
    from docx.oxml.ns import qn
    out = tmp_path / "resenja"
    result = generate_documents(_make_persons(), _make_oik(2), _make_election(), out)
    doc = Document(result.files_written[0])
    bm_table = doc.tables[0]  # First table is the BM-name one.
    assert len(bm_table.rows) == 2

    def border_val(cell, side: str) -> str | None:
        tc_pr = cell._tc.find(qn("w:tcPr"))
        if tc_pr is None:
            return None
        borders = tc_pr.find(qn("w:tcBorders"))
        if borders is None:
            return None
        el = borders.find(qn(f"w:{side}"))
        return None if el is None else el.get(qn("w:val"))

    top_cell = bm_table.rows[0].cells[0]
    bot_cell = bm_table.rows[1].cells[0]
    # Only the top cell's bottom border is "single"; everything else is "nil".
    assert border_val(top_cell, "bottom") == "single"
    for side in ("top", "left", "right"):
        assert border_val(top_cell, side) == "nil"
    for side in ("top", "bottom", "left", "right"):
        assert border_val(bot_cell, side) == "nil"


def test_members_table_column_widths_match_reference(tmp_path: Path) -> None:
    """Column widths must match the maintainer's reference .docx exactly."""
    out = tmp_path / "resenja"
    result = generate_documents(_make_persons(), _make_oik(2), _make_election(), out)
    doc = Document(result.files_written[0])
    members_table = doc.tables[1]  # 0: BM-name, 1: members, 2: stamp spacer
    expected = [324, 1736, 2193, 1417, 3968]
    actual = [int(col.width / 635) for col in members_table.columns]
    assert actual == expected, f"got widths {actual}, expected {expected}"


def test_prosireni_slot_number_column_fits_two_digits(tmp_path: Path, monkeypatch) -> None:
    """Field report: a client's generated прошиени решења showed "12)"
    wrapped across two lines ("12" then ")" below it) for every
    double-digit slot. Stalni's slot number is always single-digit
    (1/2/3) so its 324-twip column 0 is fine as-is (and must keep
    matching the reference .docx exactly -- see the test above); but
    прошиени's slot number is a list's permanent ballot-order position
    and routinely reaches two digits, which didn't fit in that width at
    TABLE_PT. Column 0 must be widened for прошиени specifically, with
    the row's total width unchanged (9638 twips) so the right margin
    stays aligned."""
    from biracki_odbor.core.prosireni_allocation import (
        MEMBER_LABEL, SUBSTITUTE_LABEL, slot_function,
    )
    monkeypatch.setenv("BIRACKI_ODBOR_HOME", str(tmp_path / "home"))
    persons = [
        Person(
            bm_number=1, function=slot_function(label, 12), party="СНС",
            first_name=first, last_name="Презименић",
            jmbg=VALID_SAMPLES[i % len(VALID_SAMPLES)],
            mesto="Уб", adresa="Адреса", telefon="0641234567",
        )
        for i, (label, first) in enumerate((
            (MEMBER_LABEL, "Члан12"), (SUBSTITUTE_LABEL, "Заменик12"),
        ))
    ]
    out = tmp_path / "resenja"
    result = generate_documents(
        persons, _make_oik(1), _make_election(), out, spec=PROSIRENI_SPEC,
    )
    doc = Document(result.files_written[0])
    members_table = doc.tables[1]
    widths = [int(col.width / 635) for col in members_table.columns]
    assert sum(widths) == 9638, "row total must still fill the page text width"
    assert widths[0] >= 500, (
        f"col 0 is {widths[0]} twips -- too narrow to fit a two-digit"
        ' slot number + ")" on one line at TABLE_PT'
    )


def test_election_type_line_is_a_single_paragraph(tmp_path: Path) -> None:
    """Each election type renders as ONE paragraph (prefix + bold type
    name on the same line) -- confirmed against the official RIK
    reference form, which shows the prefix and type name as a single
    line, not two separate paragraphs (the earlier, wrong behavior)."""
    out = tmp_path / "resenja"
    result = generate_documents(_make_persons(), _make_oik(2), _make_election(), out)
    doc = Document(result.files_written[0])
    texts = [p.text.strip() for p in doc.paragraphs]
    # Single election type → no "КООРДИНИРАНО" prefix; standalone prefix is
    # "ЗА СПРОВОЂЕЊЕ ИЗБОРА ЗА" and the type name sits on the SAME line.
    # "ЗА ИЗБОРА ЗА" (missing СПРОВОЂЕЊЕ) isn't grammatical Serbian — a
    # real bug caught via live testing, not the intended standalone form.
    assert "ЗА СПРОВОЂЕЊЕ ИЗБОРА ЗА НАРОДНЕ ПОСЛАНИКЕ НАРОДНЕ СКУПШТИНЕ" in texts
    assert "ЗА СПРОВОЂЕЊЕ ИЗБОРА ЗА" not in texts  # no longer its own paragraph
    assert "ЗА ИЗБОРА ЗА" not in texts
    assert "ЗА КООРДИНИРАНО СПРОВОЂЕЊЕ ИЗБОРА ЗА" not in texts


def test_parallel_and_sequential_produce_identical_files(tmp_path: Path) -> None:
    """Determinism: max_workers=1 vs >1 must yield equivalent .docx output.

    Compare extracted ``word/document.xml`` rather than raw ZIP bytes — the
    ZIP container carries filesystem mtime in its per-entry headers which
    differs between processes even when the content is identical.
    """
    import hashlib
    import zipfile

    persons = _make_persons()
    oik = _make_oik(2)
    election = _make_election()

    seq_out = tmp_path / "seq"
    par_out = tmp_path / "par"

    seq_result = generate_documents(persons, oik, election, seq_out, max_workers=1)
    par_result = generate_documents(persons, oik, election, par_out)

    assert sorted(p.name for p in seq_result.files_written) == \
           sorted(p.name for p in par_result.files_written)

    def _doc_xml_hash(p: Path) -> str:
        with zipfile.ZipFile(p, "r") as zf:
            data = zf.read("word/document.xml")
        return hashlib.sha256(data).hexdigest()

    for name in (p.name for p in seq_result.files_written):
        assert _doc_xml_hash(seq_out / name) == _doc_xml_hash(par_out / name), \
            f"{name} document.xml differs between sequential and default-workers paths"
