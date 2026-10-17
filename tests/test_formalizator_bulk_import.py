"""Tests for formalizator.core.bulk_import — filling many voters'
declarations from a .xlsx or .csv file instead of one at a time.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from formalizator.core import bulk_import as bi
from formalizator.core.models import CalibratedTemplate, CustomFieldDef, FieldPosition

VALID_JMBG = "2308985774805"  # 23.08.1985


def _template(custom_fields=()) -> CalibratedTemplate:
    return CalibratedTemplate(
        name="t", background_pdf_path="",
        custom_fields=list(custom_fields),
        positions=[
            FieldPosition(field_key=cf.key, page_index=0, x=1.0, y=500.0 - i * 10)
            for i, cf in enumerate(custom_fields) if cf.kind == "individual"
        ],
    )


# ── headers_for ──────────────────────────────────────────────────────


def test_headers_match_the_labels_already_on_the_fillout_screen() -> None:
    """One vocabulary, not a second one invented for files — these are
    the exact strings fillout_screen.py's form.addRow(...) calls use."""
    headers = bi.headers_for(None)

    assert headers == [
        "Име", "Презиме", "ЈМБГ", "Датум рођења", "Место", "Адреса",
        "Врста документа", "Број документа", "Датум издавања", "Издавалац",
    ]


def test_headers_include_only_individual_custom_fields() -> None:
    """Universal fields are session-wide, typed once on the setup
    screen — never a per-row column here."""
    tpl = _template([
        CustomFieldDef(key="overitelj", label="Оверитељ", kind="universal"),
        CustomFieldDef(key="napomena", label="Напомена", kind="individual"),
    ])

    headers = bi.headers_for(tpl)

    assert "Напомена" in headers
    assert "Оверитељ" not in headers


# ── Template writing ─────────────────────────────────────────────────


def test_write_template_xlsx_has_the_right_header_row(tmp_path: Path) -> None:
    import openpyxl

    tpl = _template([CustomFieldDef(key="napomena", label="Напомена", kind="individual")])
    path = tmp_path / "sablon.xlsx"

    bi.write_template(tpl, path)

    wb = openpyxl.load_workbook(path)
    header = [c.value for c in wb.active[1]]
    assert header == bi.headers_for(tpl)


def test_write_template_csv_is_utf8_sig_and_excel_friendly(tmp_path: Path) -> None:
    path = tmp_path / "sablon.csv"

    bi.write_template(None, path)

    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf"), "must carry a BOM so Excel opens Cyrillic correctly"
    assert path.read_text(encoding="utf-8-sig").strip() == ",".join(bi.headers_for(None))


def test_write_template_rejects_an_unknown_extension(tmp_path: Path) -> None:
    with pytest.raises(ValueError):
        bi.write_template(None, tmp_path / "sablon.docx")


# ── Reading — shared behaviour across both formats ───────────────────


def _write_and_fill_xlsx(tmp_path: Path, tpl, rows: list[dict]) -> Path:
    import openpyxl

    path = tmp_path / "podaci.xlsx"
    bi.write_template(tpl, path)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    for r, row in enumerate(rows, start=2):
        for label, value in row.items():
            ws.cell(row=r, column=headers.index(label) + 1, value=value)
    wb.save(path)
    return path


@pytest.mark.parametrize("fmt", ["xlsx", "csv"])
def test_a_valid_row_produces_a_correct_entry(tmp_path: Path, fmt: str) -> None:
    row = {"Име": "Ана", "Презиме": "Анић", "ЈМБГ": VALID_JMBG}
    if fmt == "xlsx":
        path = _write_and_fill_xlsx(tmp_path, None, [row])
    else:
        path = tmp_path / "podaci.csv"
        path.write_text("Име,Презиме,ЈМБГ\nАна,Анић,{}\n".format(VALID_JMBG), encoding="utf-8-sig")

    result = bi.read_entries(path, None)

    assert result.file_error == ""
    assert len(result.rows) == 1
    row_result = result.rows[0]
    assert row_result.ok, row_result.errors
    assert row_result.entry.full_name == "Ана Анић"
    assert row_result.entry.jmbg == VALID_JMBG
    assert row_result.entry.source == "bulk"


def test_date_of_birth_is_derived_from_jmbg_when_left_blank(tmp_path: Path) -> None:
    path = _write_and_fill_xlsx(
        tmp_path, None, [{"Име": "Ана", "Презиме": "Анић", "ЈМБГ": VALID_JMBG}],
    )

    result = bi.read_entries(path, None)

    assert result.rows[0].entry.date_of_birth == "23.08.1985"


def test_an_explicit_date_of_birth_is_not_overwritten(tmp_path: Path) -> None:
    path = _write_and_fill_xlsx(tmp_path, None, [{
        "Име": "Ана", "Презиме": "Анић", "ЈМБГ": VALID_JMBG, "Датум рођења": "01.01.1999",
    }])

    result = bi.read_entries(path, None)

    assert result.rows[0].entry.date_of_birth == "01.01.1999"


def test_a_numeric_jmbg_cell_is_read_correctly(tmp_path: Path) -> None:
    """Excel stores a plain-typed 13-digit number as a number, not
    text, unless the operator formats the column first — a real
    operator's sheet will have some of these."""
    path = _write_and_fill_xlsx(tmp_path, None, [
        {"Име": "Ана", "Презиме": "Анић", "ЈМБГ": int(VALID_JMBG)},
    ])

    result = bi.read_entries(path, None)

    assert result.rows[0].entry.jmbg == VALID_JMBG
    assert result.rows[0].ok


def test_missing_name_is_reported_per_row(tmp_path: Path) -> None:
    path = _write_and_fill_xlsx(tmp_path, None, [{"Име": "", "Презиме": "", "ЈМБГ": VALID_JMBG}])

    result = bi.read_entries(path, None)

    assert not result.rows[0].ok
    assert any("Име и презиме" in e for e in result.rows[0].errors)


def test_a_bad_jmbg_is_reported_and_does_not_block_other_rows(tmp_path: Path) -> None:
    path = _write_and_fill_xlsx(tmp_path, None, [
        {"Име": "Ана", "Презиме": "Анић", "ЈМБГ": VALID_JMBG},
        {"Име": "Пера", "Презиме": "Перић", "ЈМБГ": "1234567890123"},
        {"Име": "Мика", "Презиме": "Микић", "ЈМБГ": ""},
    ])

    result = bi.read_entries(path, None)

    assert [r.ok for r in result.rows] == [True, False, False]
    assert len(result.valid_rows) == 1
    assert len(result.invalid_rows) == 2


def test_a_blank_row_is_skipped_not_reported_as_an_error(tmp_path: Path) -> None:
    """Trailing blank Excel rows, a stray Enter — not the operator's
    mistake to be told about."""
    path = _write_and_fill_xlsx(tmp_path, None, [
        {"Име": "Ана", "Презиме": "Анић", "ЈМБГ": VALID_JMBG},
        {},
    ])

    result = bi.read_entries(path, None)

    assert len(result.rows) == 1


def test_custom_field_values_round_trip(tmp_path: Path) -> None:
    tpl = _template([CustomFieldDef(key="napomena", label="Напомена", kind="individual")])
    path = _write_and_fill_xlsx(tmp_path, tpl, [{
        "Име": "Ана", "Презиме": "Анић", "ЈМБГ": VALID_JMBG, "Напомена": "прва посета",
    }])

    result = bi.read_entries(path, tpl)

    assert result.rows[0].entry.custom_values == {"napomena": "прва посета"}


def test_reading_without_a_template_does_not_crash_on_a_custom_column(tmp_path: Path) -> None:
    """The template used to write a file and the one active at import
    time can differ in principle -- must degrade gracefully, not raise."""
    tpl = _template([CustomFieldDef(key="napomena", label="Напомена", kind="individual")])
    path = _write_and_fill_xlsx(tmp_path, tpl, [{
        "Име": "Ана", "Презиме": "Анић", "ЈМБГ": VALID_JMBG, "Напомена": "x",
    }])

    result = bi.read_entries(path, None)  # noqa: -- no template this time

    assert result.file_error == ""
    assert result.rows[0].ok


def test_an_unrecognised_extra_column_is_ignored_not_flagged(tmp_path: Path) -> None:
    import openpyxl

    path = tmp_path / "podaci.xlsx"
    bi.write_template(None, path)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws.cell(row=1, column=len(bi.headers_for(None)) + 1, value="Напомена оператера")
    ws.cell(row=2, column=1, value="Ана")
    ws.cell(row=2, column=2, value="Анић")
    ws.cell(row=2, column=3, value=VALID_JMBG)
    ws.cell(row=2, column=len(bi.headers_for(None)) + 1, value="нешто лично")
    wb.save(path)

    result = bi.read_entries(path, None)

    assert result.rows[0].ok


# ── File-level errors ─────────────────────────────────────────────────


def test_a_missing_file_is_a_clean_error_not_a_crash(tmp_path: Path) -> None:
    result = bi.read_entries(tmp_path / "does_not_exist.csv", None)

    assert result.file_error
    assert result.rows == []


def test_an_empty_file_is_a_clean_error(tmp_path: Path) -> None:
    path = tmp_path / "empty.csv"
    path.write_text("", encoding="utf-8-sig")

    result = bi.read_entries(path, None)

    assert result.file_error


def test_an_unsupported_extension_is_a_clean_error(tmp_path: Path) -> None:
    path = tmp_path / "data.txt"
    path.write_text("Име,Презиме,ЈМБГ\n", encoding="utf-8")

    result = bi.read_entries(path, None)

    assert result.file_error
