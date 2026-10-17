"""Tests for BulkImportDialog and its wiring into FilloutScreen.

Print/fill are stubbed exactly like test_formalizator_fillout_screen.py
does for the single-entry path -- these tests are about the batch
plumbing (which rows get processed, in what order, what happens to a
failure mid-batch), not about print_pdf/fill_declaration themselves.
"""

from __future__ import annotations

from pathlib import Path

import pytest

VALID_JMBG_1 = "2308985774805"
VALID_JMBG_2 = "1605991774818"
VALID_JMBG_3 = "0311990774809"


@pytest.fixture
def app():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def _armed_screen(app, monkeypatch):
    import formalizator.gui.fillout_screen as mod
    from formalizator.core.models import CalibratedTemplate, FieldPosition
    from formalizator.gui.fillout_screen import FilloutScreen

    monkeypatch.setattr(mod, "print_pdf", lambda *a, **kw: None)
    monkeypatch.setattr(mod, "fill_declaration", lambda *a, **kw: None)
    screen = FilloutScreen()
    screen.set_template(CalibratedTemplate(
        name="t", background_pdf_path="",
        positions=[FieldPosition(field_key="full_name", page_index=0, x=1.0, y=1.0)],
    ))
    recorded = []
    screen.entry_printed.connect(recorded.append)
    return screen, recorded


def _xlsx_with_rows(tmp_path: Path, rows: list[dict]) -> Path:
    from formalizator.core import bulk_import as bi

    path = tmp_path / "podaci.xlsx"
    bi.write_template(None, path)
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    for r, row in enumerate(rows, start=2):
        for label, value in row.items():
            ws.cell(row=r, column=headers.index(label) + 1, value=value)
    wb.save(path)
    return path


def _open_dialog(screen):
    from formalizator.gui.dialogs.bulk_import_dialog import BulkImportDialog

    return BulkImportDialog(screen._template, screen._process_entry)  # noqa: SLF001


def test_the_button_opens_the_dialog(app, monkeypatch) -> None:
    screen, _recorded = _armed_screen(app, monkeypatch)
    opened = []
    monkeypatch.setattr(
        "formalizator.gui.dialogs.bulk_import_dialog.BulkImportDialog.exec",
        lambda self: opened.append(self) or 0,
    )
    try:
        screen.btn_bulk_import.click()
        assert len(opened) == 1
    finally:
        screen.deleteLater()


def test_the_button_refuses_without_a_template(app, monkeypatch) -> None:
    import formalizator.gui.fillout_screen as mod
    from formalizator.gui.fillout_screen import FilloutScreen

    warned = []
    monkeypatch.setattr(mod.QMessageBox, "warning", lambda *a, **kw: warned.append(a))
    screen = FilloutScreen()
    try:
        screen.btn_bulk_import.click()
        assert warned, "must warn instead of opening a dialog with no template"
    finally:
        screen.deleteLater()


def test_picking_a_file_renders_the_preview_table(app, monkeypatch, tmp_path) -> None:
    screen, _recorded = _armed_screen(app, monkeypatch)
    try:
        dialog = _open_dialog(screen)
        path = _xlsx_with_rows(tmp_path, [
            {"Име": "Ана", "Презиме": "Анић", "ЈМБГ": VALID_JMBG_1},
            {"Име": "Пера", "Презиме": "Перић", "ЈМБГ": "bad"},
        ])
        monkeypatch.setattr(
            "formalizator.gui.dialogs.bulk_import_dialog.QFileDialog.getOpenFileName",
            lambda *a, **kw: (str(path), ""),
        )

        dialog.btn_pick.click()

        assert dialog.table.rowCount() == 2
        assert dialog.table.item(0, 1).text() == "Ана Анић"
        assert dialog.table.item(0, 3).text() == "У реду"
        assert dialog.table.item(1, 3).text() != "У реду"
        assert dialog.btn_run.isEnabled(), "one valid row is enough to enable running"
    finally:
        screen.deleteLater()


def test_running_processes_only_valid_rows_and_records_each_one(
    app, monkeypatch, tmp_path,
) -> None:
    screen, recorded = _armed_screen(app, monkeypatch)
    try:
        dialog = _open_dialog(screen)
        path = _xlsx_with_rows(tmp_path, [
            {"Име": "Ана", "Презиме": "Анић", "ЈМБГ": VALID_JMBG_1},
            {"Име": "Пера", "Презиме": "Перић", "ЈМБГ": "bad"},  # excluded
            {"Име": "Мика", "Презиме": "Микић", "ЈМБГ": VALID_JMBG_2},
        ])
        monkeypatch.setattr(
            "formalizator.gui.dialogs.bulk_import_dialog.QFileDialog.getOpenFileName",
            lambda *a, **kw: (str(path), ""),
        )
        monkeypatch.setattr(
            "formalizator.gui.dialogs.bulk_import_dialog.QMessageBox.question",
            lambda *a, **kw: __import__("PySide6.QtWidgets", fromlist=["QMessageBox"]).QMessageBox.Yes,
        )
        monkeypatch.setattr(
            "formalizator.gui.dialogs.bulk_import_dialog.QMessageBox.information",
            lambda *a, **kw: None,
        )
        dialog.btn_pick.click()

        dialog.btn_run.click()

        names = {e.full_name for e in recorded}
        assert names == {"Ана Анић", "Мика Микић"}, (
            "exactly the two valid rows must be processed -- not the invalid one, "
            "not fewer than both valid ones"
        )
    finally:
        screen.deleteLater()


def test_a_fill_failure_on_one_row_does_not_stop_the_rest(app, monkeypatch, tmp_path) -> None:
    """A batch of many people must survive one row's data breaking
    fill_declaration -- losing everyone after it to one bad row would be
    far worse than losing just that one."""
    import formalizator.gui.fillout_screen as mod

    screen, recorded = _armed_screen(app, monkeypatch)
    try:
        calls = []

        def flaky_fill(template, entry, out_path, session_info):
            calls.append(entry.full_name)
            if entry.full_name == "Пера Перић":
                raise RuntimeError("overlay error")

        monkeypatch.setattr(mod, "fill_declaration", flaky_fill)

        dialog = _open_dialog(screen)
        path = _xlsx_with_rows(tmp_path, [
            {"Име": "Ана", "Презиме": "Анић", "ЈМБГ": VALID_JMBG_1},
            {"Име": "Пера", "Презиме": "Перић", "ЈМБГ": VALID_JMBG_2},
            {"Име": "Мика", "Презиме": "Микић", "ЈМБГ": VALID_JMBG_3},
        ])
        monkeypatch.setattr(
            "formalizator.gui.dialogs.bulk_import_dialog.QFileDialog.getOpenFileName",
            lambda *a, **kw: (str(path), ""),
        )
        from PySide6.QtWidgets import QMessageBox
        monkeypatch.setattr(
            "formalizator.gui.dialogs.bulk_import_dialog.QMessageBox.question",
            lambda *a, **kw: QMessageBox.Yes,
        )
        monkeypatch.setattr(
            "formalizator.gui.dialogs.bulk_import_dialog.QMessageBox.warning",
            lambda *a, **kw: None,
        )
        dialog.btn_pick.click()

        dialog.btn_run.click()

        assert len(calls) == 3, "all three rows must be attempted"
        names = {e.full_name for e in recorded}
        assert names == {"Ана Анић", "Мика Микић"}, "only the row that failed to fill is unrecorded"
    finally:
        screen.deleteLater()


def test_the_template_download_writes_a_real_file(app, monkeypatch, tmp_path) -> None:
    screen, _recorded = _armed_screen(app, monkeypatch)
    try:
        dialog = _open_dialog(screen)
        target = tmp_path / "moj_sablon.xlsx"
        monkeypatch.setattr(
            "formalizator.gui.dialogs.bulk_import_dialog.QFileDialog.getSaveFileName",
            lambda *a, **kw: (str(target), ""),
        )
        monkeypatch.setattr(
            "formalizator.gui.dialogs.bulk_import_dialog.QMessageBox.information",
            lambda *a, **kw: None,
        )

        dialog.btn_template_xlsx.click()

        assert target.is_file()
    finally:
        screen.deleteLater()
