"""Tests for formalizator.core.generic_export and export_template_catalog
— a cumulative export layout the operator defines and reuses, instead
of a hardcoded roster shape per document."""

from __future__ import annotations

from pathlib import Path

import openpyxl

from openpyxl.styles import Border, Side

from formalizator.core import export_template_catalog
from formalizator.core.generic_export import export_roster
from formalizator.core.models import (
    ROW_NUMBER_FIELD_KEY, ExportColumnDef, ExportTemplateDef, SessionInfo, VoterEntry,
)


def _nprs4_like_source(path: Path) -> Path:
    """A miniature of the real NPRS-4 file's actual shape: a title
    block, a bordered header row, ONE ready data row already carrying
    its own "Редни број" of 1, and merged footer blocks below it."""
    wb = openpyxl.Workbook()
    ws = wb.active
    edge = Side(style="thin")
    box = Border(left=edge, right=edge, top=edge, bottom=edge)

    ws.cell(row=1, column=4, value="Образац НПРС-4/23")
    ws.cell(row=2, column=1, value="С П И С А К БИРАЧА…")
    ws.merge_cells("A2:D2")
    for col, text in enumerate(
        ("Редни\nброј", "Презиме и име бирача", "ЈМБГ", "Овлашћени оверитељ"), start=1,
    ):
        cell = ws.cell(row=3, column=col, value=text)
        cell.border = box
    ws.cell(row=4, column=1, value="1").border = box  # the form's own ready data row
    for col in range(2, 5):
        ws.cell(row=4, column=col).border = box
    ws.cell(row=5, column=1, value="(навести све бираче…)")
    ws.merge_cells("A5:D6")
    ws.cell(row=7, column=1, value="Лице које подноси изборну листу ___ (потпис)")
    ws.merge_cells("A7:D8")
    wb.save(path)
    return path


def _nprs4_like_template(source: Path) -> ExportTemplateDef:
    return ExportTemplateDef(
        name="НПРС-4", source_file=str(source), header_row=3,
        columns=[
            ExportColumnDef(header="Редни број", field_keys=[ROW_NUMBER_FIELD_KEY], column_index=1),
            ExportColumnDef(
                header="Презиме и име бирача",
                field_keys=["surname", "given_name"], separators=[" "], column_index=2,
            ),
            ExportColumnDef(header="ЈМБГ", field_keys=["jmbg"], column_index=3),
            ExportColumnDef(header="Овлашћени оверитељ", field_keys=["overitelj"], column_index=4),
        ],
    )


def test_export_roster_writes_header_and_rows(tmp_path: Path) -> None:
    template = ExportTemplateDef(
        name="Тест извоз",
        columns=[
            ExportColumnDef(header="Име и презиме", field_keys=["full_name"]),
            ExportColumnDef(header="ЈМБГ", field_keys=["jmbg"]),
        ],
    )
    entries = [
        VoterEntry(given_name="Марко", surname="Марковић", jmbg="1" * 13),
        VoterEntry(given_name="Ана", surname="Анић", jmbg="2" * 13),
    ]

    out = export_roster(entries, SessionInfo(), template, tmp_path / "out.xlsx")

    wb = openpyxl.load_workbook(out)
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "Име и презиме"
    assert ws.cell(row=1, column=2).value == "ЈМБГ"
    assert ws.cell(row=2, column=1).value == "Марко Марковић"
    assert ws.cell(row=2, column=2).value == "1" * 13
    assert ws.cell(row=3, column=1).value == "Ана Анић"


def test_export_roster_resolves_custom_universal_field(tmp_path: Path) -> None:
    template = ExportTemplateDef(
        name="t", columns=[ExportColumnDef(header="Оверитељ", field_keys=["overitelj_ime"])],
    )
    session = SessionInfo(custom_values={"overitelj_ime": "Петар Петровић"})

    out = export_roster([VoterEntry()], session, template, tmp_path / "out.xlsx")

    ws = openpyxl.load_workbook(out).active
    assert ws.cell(row=2, column=1).value == "Петар Петровић"


def test_export_roster_unknown_field_key_writes_blank(tmp_path: Path) -> None:
    template = ExportTemplateDef(
        name="t", columns=[ExportColumnDef(header="?", field_keys=["no_such_field"])],
    )

    out = export_roster([VoterEntry()], SessionInfo(), template, tmp_path / "out.xlsx")

    # openpyxl doesn't round-trip an empty-string cell value through a
    # save/reload — it comes back as None, same observable "blank" cell.
    ws = openpyxl.load_workbook(out).active
    assert ws.cell(row=2, column=1).value in (None, "")


def test_export_roster_joins_multiple_fields_with_own_separators(tmp_path: Path) -> None:
    """The operator's own join mechanism: any number of fields, each
    with its own separator text — e.g. municipality - certifier,
    certifier address, built entirely from the export template, not
    hardcoded here."""
    template = ExportTemplateDef(
        name="t",
        columns=[
            ExportColumnDef(
                header="Оверитељ",
                field_keys=["opstina", "overitelj_ime", "overitelj_adresa"],
                separators=[" - ", ", "],
            ),
        ],
    )
    session = SessionInfo(custom_values={
        "opstina": "Мионица", "overitelj_ime": "Петар Петровић", "overitelj_adresa": "Ул. 1",
    })

    out = export_roster([VoterEntry()], session, template, tmp_path / "out.xlsx")

    ws = openpyxl.load_workbook(out).active
    assert ws.cell(row=2, column=1).value == "Мионица - Петар Петровић, Ул. 1"


def test_export_roster_sorts_by_configured_columns(tmp_path: Path) -> None:
    template = ExportTemplateDef(
        name="t",
        columns=[
            ExportColumnDef(header="Оверитељ", field_keys=["overitelj"]),
            ExportColumnDef(header="Име", field_keys=["full_name"]),
        ],
        sort_column_1="Оверитељ", sort_column_2="Име",
    )
    entries = [
        VoterEntry(given_name="Зоран", surname="Зорић", custom_values={"overitelj": "Б"}),
        VoterEntry(given_name="Ана", surname="Анић", custom_values={"overitelj": "А"}),
        VoterEntry(given_name="Бранко", surname="Бранковић", custom_values={"overitelj": "А"}),
    ]

    out = export_roster(entries, SessionInfo(), template, tmp_path / "out.xlsx")

    ws = openpyxl.load_workbook(out).active
    names = [ws.cell(row=r, column=2).value for r in (2, 3, 4)]
    assert names == ["Ана Анић", "Бранко Бранковић", "Зоран Зорић"]


def test_export_roster_with_source_file_inserts_rows_below_header_and_keeps_footer(tmp_path: Path) -> None:
    """The core "same physical shape every time" requirement: many
    machines produce files with identical structure so they can be
    merged by hand later — the pristine source is never mutated, and
    whatever sits below the data (footer legal text, signature blocks)
    survives untouched, just pushed down."""
    source = tmp_path / "source.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="Образац НПРС-4/23")
    ws.cell(row=3, column=1, value="Име и презиме")
    ws.cell(row=3, column=2, value="ЈМБГ")
    ws.cell(row=4, column=1, value="НАПОМЕНА: правно обавезујуће.")
    wb.save(source)
    original_bytes = source.read_bytes()

    template = ExportTemplateDef(
        name="t",
        columns=[
            ExportColumnDef(header="Име и презиме", field_keys=["full_name"], column_index=1),
            ExportColumnDef(header="ЈМБГ", field_keys=["jmbg"], column_index=2),
        ],
        source_file=str(source), header_row=3,
    )
    entries = [
        VoterEntry(given_name="Марко", surname="Марковић", jmbg="1" * 13),
        VoterEntry(given_name="Ана", surname="Анић", jmbg="2" * 13),
    ]

    out_path = tmp_path / "out.xlsx"
    export_roster(entries, SessionInfo(), template, out_path)

    assert source.read_bytes() == original_bytes  # pristine source never touched

    ws_out = openpyxl.load_workbook(out_path).active
    assert ws_out.cell(row=1, column=1).value == "Образац НПРС-4/23"
    assert ws_out.cell(row=3, column=1).value == "Име и презиме"
    assert ws_out.cell(row=4, column=1).value == "Марко Марковић"
    assert ws_out.cell(row=4, column=2).value == "1" * 13
    assert ws_out.cell(row=5, column=1).value == "Ана Анић"
    # footer pushed down by 2 (one row per new entry), not overwritten
    assert ws_out.cell(row=6, column=1).value == "НАПОМЕНА: правно обавезујуће."


def test_export_roster_writes_skipped_column_to_its_own_position_not_compacted(tmp_path: Path) -> None:
    """Reported directly: a real NPRS-4-shaped file has four columns
    (Редни број / Презиме и име бирача / ЈМБГ / Овлашћени оверитељ);
    leaving "Редни број" unmapped must not compact the other three
    columns leftward — "Презиме и име бирача" belongs in column B and
    "Овлашћени оверитељ" in column D, exactly where their own headers
    sit, regardless of how many rows the operator's mapping table has
    or what order they're in."""
    source = tmp_path / "source.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="Редни број")
    ws.cell(row=1, column=2, value="Презиме и име бирача")
    ws.cell(row=1, column=3, value="ЈМБГ")
    ws.cell(row=1, column=4, value="Овлашћени оверитељ")
    wb.save(source)

    template = ExportTemplateDef(
        name="t",
        columns=[
            ExportColumnDef(header="Презиме и име бирача", field_keys=["full_name"], column_index=2),
            ExportColumnDef(header="Овлашћени оверитељ", field_keys=["overitelj"], column_index=4),
        ],
        source_file=str(source), header_row=1,
    )
    session = SessionInfo(custom_values={"overitelj": "Јелена Јовановић"})
    entries = [VoterEntry(given_name="Стеван", surname="Стевић")]

    out_path = tmp_path / "out.xlsx"
    export_roster(entries, session, template, out_path)

    ws_out = openpyxl.load_workbook(out_path).active
    assert ws_out.cell(row=2, column=1).value is None  # "Редни број" untouched, not written
    assert ws_out.cell(row=2, column=2).value == "Стеван Стевић"
    assert ws_out.cell(row=2, column=3).value is None  # "ЈМБГ" untouched, not written
    assert ws_out.cell(row=2, column=4).value == "Јелена Јовановић"


def test_export_roster_appends_unpositioned_columns_after_the_known_ones(tmp_path: Path) -> None:
    """A column the operator adds by hand (no file loaded, so no real
    position known) has column_index == 0 — when inserting into a real
    source file, it must land after every positioned column, not
    collide with a real header's own column."""
    source = tmp_path / "source.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="ЈМБГ")
    wb.save(source)

    template = ExportTemplateDef(
        name="t",
        columns=[
            ExportColumnDef(header="ЈМБГ", field_keys=["jmbg"], column_index=1),
            ExportColumnDef(header="Напомена", field_keys=["napomena"]),  # column_index=0
        ],
        source_file=str(source), header_row=1,
    )
    session = SessionInfo(custom_values={"napomena": "нова колона"})

    out = export_roster([VoterEntry(jmbg="1" * 13)], session, template, tmp_path / "out.xlsx")

    ws_out = openpyxl.load_workbook(out).active
    assert ws_out.cell(row=2, column=1).value == "1" * 13
    assert ws_out.cell(row=2, column=2).value == "нова колона"


def test_export_reuses_the_forms_own_ready_data_row_for_the_first_entry(tmp_path: Path) -> None:
    """The operator's description of the real file: "it only has row 1
    for entry, the rest of the file needs to remain intact with new
    rows just be inserted below 1". So entry 1 fills the form's
    existing row 4; only entries 2..N add rows underneath it."""
    source = _nprs4_like_source(tmp_path / "source.xlsx")
    template = _nprs4_like_template(source)
    session = SessionInfo(custom_values={"overitelj": "Наташа Пауновић"})
    entries = [
        VoterEntry(given_name="Стеван", surname="Стевић", jmbg="1" * 13),
        VoterEntry(given_name="Ана", surname="Анић", jmbg="2" * 13),
    ]

    out = export_roster(entries, session, template, tmp_path / "out.xlsx")

    ws = openpyxl.load_workbook(out).active
    assert ws.cell(row=4, column=2).value == "Стевић Стеван"
    assert ws.cell(row=5, column=2).value == "Анић Ана"
    # One extra row added for the second entry, not two.
    assert ws.cell(row=6, column=1).value == "(навести све бираче…)"


def test_export_never_overwrites_footer_when_there_is_no_spare_data_row(tmp_path: Path) -> None:
    """A form whose header runs straight into footer text has no ready
    data row to reuse — writing into it would destroy legal text, so
    every entry gets an inserted row instead."""
    source = tmp_path / "source.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="Име")
    ws.cell(row=2, column=1, value="НАПОМЕНА: правно обавезујуће.")  # footer, not a data row
    wb.save(source)

    template = ExportTemplateDef(
        name="t", source_file=str(source), header_row=1,
        columns=[ExportColumnDef(header="Име", field_keys=["full_name"], column_index=1)],
    )

    out = export_roster([VoterEntry(given_name="Марко")], SessionInfo(), template, tmp_path / "out.xlsx")

    ws_out = openpyxl.load_workbook(out).active
    assert ws_out.cell(row=2, column=1).value == "Марко"
    assert ws_out.cell(row=3, column=1).value == "НАПОМЕНА: правно обавезујуће."


def test_export_shifts_merged_footer_blocks_instead_of_stranding_them(tmp_path: Path) -> None:
    """openpyxl's insert_rows moves cell values but leaves merged
    ranges behind — confirmed directly against the real NPRS-4 file.
    Without re-pinning them, the note/signature blocks end up merged
    over the wrong rows and the form is visibly wrecked."""
    source = _nprs4_like_source(tmp_path / "source.xlsx")
    template = _nprs4_like_template(source)
    entries = [VoterEntry(given_name=str(i), surname="X") for i in range(3)]

    out = export_roster(entries, SessionInfo(), template, tmp_path / "out.xlsx")

    ws = openpyxl.load_workbook(out).active
    merges = {str(m) for m in ws.merged_cells.ranges}
    assert "A2:D2" in merges  # title, above the insert point — unmoved
    # 3 entries into a form with 1 ready row = 2 inserted rows, so every
    # footer block moves down by exactly 2.
    assert "A7:D8" in merges
    assert "A9:D10" in merges
    assert ws.cell(row=7, column=1).value == "(навести све бираче…)"
    assert ws.cell(row=9, column=1).value == "Лице које подноси изборну листу ___ (потпис)"


def test_export_never_merges_cells_inside_the_data_table(tmp_path: Path) -> None:
    """RIK's automatic validation rejects merged cells in the table, so
    added rows must carry formatting only — never a merge."""
    source = _nprs4_like_source(tmp_path / "source.xlsx")
    template = _nprs4_like_template(source)
    entries = [VoterEntry(given_name=str(i), surname="X") for i in range(4)]

    out = export_roster(entries, SessionInfo(), template, tmp_path / "out.xlsx")

    ws = openpyxl.load_workbook(out).active
    data_rows = range(4, 8)
    inside = [
        str(m) for m in ws.merged_cells.ranges
        if any(m.min_row <= r <= m.max_row for r in data_rows)
    ]
    assert inside == []


def test_inserted_rows_inherit_the_forms_table_borders(tmp_path: Path) -> None:
    """Rows openpyxl inserts carry no formatting at all — every entry
    after the first would otherwise print outside the table's grid."""
    source = _nprs4_like_source(tmp_path / "source.xlsx")
    template = _nprs4_like_template(source)
    entries = [VoterEntry(given_name=str(i), surname="X") for i in range(3)]

    out = export_roster(entries, SessionInfo(), template, tmp_path / "out.xlsx")

    ws = openpyxl.load_workbook(out).active
    for row in (4, 5, 6):
        border = ws.cell(row=row, column=2).border
        assert border.left.style and border.right.style
        assert border.top.style and border.bottom.style


def test_row_number_field_numbers_entries_after_sorting(tmp_path: Path) -> None:
    """"Редни број" has to read 1..N straight down the finished page,
    which means numbering the rows only once their order is settled."""
    template = ExportTemplateDef(
        name="t",
        columns=[
            ExportColumnDef(header="Рб", field_keys=[ROW_NUMBER_FIELD_KEY]),
            ExportColumnDef(header="Име", field_keys=["full_name"]),
        ],
        sort_column_1="Име",
    )
    entries = [
        VoterEntry(given_name="Зоран", surname="Зорић"),
        VoterEntry(given_name="Ана", surname="Анић"),
    ]

    out = export_roster(entries, SessionInfo(), template, tmp_path / "out.xlsx")

    ws = openpyxl.load_workbook(out).active
    assert [ws.cell(row=r, column=1).value for r in (2, 3)] == [1, 2]  # real numbers, not text
    assert [ws.cell(row=r, column=2).value for r in (2, 3)] == ["Ана Анић", "Зоран Зорић"]


def test_column_without_a_recorded_position_falls_back_to_matching_its_header(tmp_path: Path) -> None:
    """Reported directly: an export template defined before column
    positions were recorded put its data in columns E and G, outside
    the form's own table. Matching the column's header text against
    the file's header row rescues it — including headers the form
    splits across lines ("Редни\\nброј")."""
    source = _nprs4_like_source(tmp_path / "source.xlsx")
    template = ExportTemplateDef(
        name="t", source_file=str(source), header_row=3,
        columns=[  # no column_index at all, as an older saved template would be
            ExportColumnDef(header="Редни број", field_keys=[ROW_NUMBER_FIELD_KEY]),
            ExportColumnDef(header="Овлашћени оверитељ", field_keys=["overitelj"]),
        ],
    )
    session = SessionInfo(custom_values={"overitelj": "Наташа Пауновић"})

    out = export_roster([VoterEntry()], session, template, tmp_path / "out.xlsx")

    ws = openpyxl.load_workbook(out).active
    assert ws.cell(row=4, column=1).value == 1  # column A, matched despite the line break
    assert ws.cell(row=4, column=4).value == "Наташа Пауновић"  # column D, not E/G


def test_export_template_catalog_round_trips(tmp_path: Path) -> None:
    template = ExportTemplateDef(
        name="Тест", columns=[ExportColumnDef(header="Име", field_keys=["full_name"])],
        source_file="C:/some/source.xlsx", header_row=3,
        sort_column_1="Име", sort_column_2="ЈМБГ",
    )
    export_template_catalog.save_template(tmp_path, template)

    loaded = export_template_catalog.find_by_name(tmp_path, "Тест")
    assert loaded is not None
    assert loaded.columns[0].header == "Име"
    assert loaded.columns[0].field_keys == ["full_name"]
    assert loaded.source_file == "C:/some/source.xlsx"
    assert loaded.header_row == 3
    assert loaded.sort_column_1 == "Име"
    assert loaded.sort_column_2 == "ЈМБГ"


def test_catalog_migrates_the_print_only_jmbg_comb_key_to_the_real_one(tmp_path: Path) -> None:
    """A template saved while the picker still offered the calibration
    wizard's comb key exports a blank JMBG column, because that key
    describes a printing layout and resolves to nothing. Loading
    rewrites it so an already-saved template fixes itself."""
    import json

    (tmp_path / "export_templates.json").write_text(json.dumps([{
        "name": "НПРС-4",
        "columns": [{
            "header": "ЈМБГ", "field_keys": ["__jmbg_comb__"],
            "separators": [], "column_index": 3,
        }],
    }]), encoding="utf-8")

    loaded = export_template_catalog.find_by_name(tmp_path, "НПРС-4")

    assert loaded is not None
    assert loaded.columns[0].field_keys == ["jmbg"]
    assert loaded.columns[0].column_index == 3  # everything else untouched


def test_migrated_jmbg_column_actually_exports_the_number(tmp_path: Path) -> None:
    import json

    (tmp_path / "export_templates.json").write_text(json.dumps([{
        "name": "t",
        "columns": [{"header": "ЈМБГ", "field_keys": ["__jmbg_comb__"], "separators": []}],
    }]), encoding="utf-8")
    template = export_template_catalog.find_by_name(tmp_path, "t")

    out = export_roster(
        [VoterEntry(jmbg="2308985774805")], SessionInfo(), template, tmp_path / "out.xlsx",
    )

    assert openpyxl.load_workbook(out).active.cell(row=2, column=1).value == "2308985774805"


def test_export_template_catalog_empty_when_no_index(tmp_path: Path) -> None:
    assert export_template_catalog.load_catalog(tmp_path) == []


def test_export_template_catalog_save_replaces_by_name(tmp_path: Path) -> None:
    export_template_catalog.save_template(
        tmp_path, ExportTemplateDef(name="t", columns=[ExportColumnDef(header="A", field_keys=["a"])]),
    )
    export_template_catalog.save_template(
        tmp_path, ExportTemplateDef(name="t", columns=[ExportColumnDef(header="B", field_keys=["b"])]),
    )

    loaded = export_template_catalog.load_catalog(tmp_path)
    assert len(loaded) == 1
    assert loaded[0].columns[0].header == "B"


def test_export_template_catalog_remove(tmp_path: Path) -> None:
    export_template_catalog.save_template(
        tmp_path, ExportTemplateDef(name="A", columns=[ExportColumnDef(header="x", field_keys=["x"])]),
    )
    export_template_catalog.save_template(
        tmp_path, ExportTemplateDef(name="B", columns=[ExportColumnDef(header="x", field_keys=["x"])]),
    )

    export_template_catalog.remove_template(tmp_path, "A")

    names = {t.name for t in export_template_catalog.load_catalog(tmp_path)}
    assert names == {"B"}
