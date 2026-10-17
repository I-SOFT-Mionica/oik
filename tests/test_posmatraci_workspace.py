"""Tests for _AkreditacijaTab (posmatraci_workspace.py) — CollapsibleSection
conversion (2026-07-12): org/observer boxes gained count badges, an Import
quick-action visible when collapsed, and Excel template export. Organization
Excel import is new (observer import already existed; orgs were previously
only ever created implicitly via observer import).
"""

from __future__ import annotations

from pathlib import Path

import pytest


@pytest.fixture
def app():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


@pytest.fixture
def home(tmp_path, monkeypatch):
    h = tmp_path / "home"
    monkeypatch.setenv("BIRACKI_ODBOR_HOME", str(h))
    monkeypatch.setattr("biracki_odbor._release_info.GITHUB_REPO", "", raising=False)
    return h


def _make_tab(app, home):
    from biracki_odbor.gui.app_state import AppState
    from biracki_odbor.gui.tabs.posmatraci_workspace import _AkreditacijaTab
    state = AppState()
    tab = _AkreditacijaTab(state)
    # isVisible() reflects EFFECTIVE (composited) visibility, which
    # requires the whole ancestor chain to be shown — without this,
    # QLabel.setVisible(True) deep inside CollapsibleSection would
    # still report isVisible() == False regardless of the call.
    tab.show()
    app.processEvents()
    return tab


def test_tab_constructs_with_empty_registry(app, home) -> None:
    tab = _make_tab(app, home)
    try:
        assert tab.orgs_table is not None
        assert tab.obs_table is not None
    finally:
        tab.deleteLater()


def test_reload_sets_count_badges(app, home) -> None:
    from biracki_odbor.core.observers import AccreditingOrg, Observer

    tab = _make_tab(app, home)
    try:
        registry = tab.state.observer_registry
        # Empty registry -> count hidden (None), not "0".
        assert tab._orgs_section._count_lbl.isVisible() is False

        registry.orgs.append(AccreditingOrg(name="НВО Око", contact="oko@example.rs"))
        registry.observers.append(Observer(
            first_name="Ана", last_name="Анић", organization="НВО Око",
            bm_assignment="Сва БМ", jmbg="0101990710006",
        ))
        tab._reload()

        assert tab._orgs_section._count_lbl.isVisible() is True
        assert tab._orgs_section._count_lbl.text() == "(1)"
        assert tab._obs_section._count_lbl.isVisible() is True
        assert tab._obs_section._count_lbl.text() == "(1)"
    finally:
        tab.deleteLater()


def test_import_orgs_excel_end_to_end(app, home, tmp_path, monkeypatch) -> None:
    import openpyxl

    src = tmp_path / "orgs.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Contact"])
    ws.append(["НВО Око", "oko@example.rs"])
    ws.append(["Центар за слободне изборе", ""])
    wb.save(src)
    wb.close()

    tab = _make_tab(app, home)
    try:
        monkeypatch.setattr(
            "biracki_odbor.gui.tabs.posmatraci_workspace.QFileDialog.getOpenFileName",
            lambda *a, **kw: (str(src), ""),
        )
        monkeypatch.setattr(
            "biracki_odbor.gui.tabs.posmatraci_workspace.QMessageBox.question",
            lambda *a, **kw: __import__(
                "PySide6.QtWidgets", fromlist=["QMessageBox"]
            ).QMessageBox.StandardButton.Yes,
        )

        tab._import_orgs_excel()

        registry = tab.state.observer_registry
        names = {o.name for o in registry.orgs}
        assert names == {"НВО Око", "Центар за слободне изборе"}
        contacts = {o.name: o.contact for o in registry.orgs}
        assert contacts["НВО Око"] == "oko@example.rs"
    finally:
        tab.deleteLater()


def test_import_orgs_excel_skips_duplicates_and_blanks(app, home, tmp_path, monkeypatch) -> None:
    import openpyxl
    from biracki_odbor.core.observers import AccreditingOrg

    src = tmp_path / "orgs.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Contact"])
    ws.append(["НВО Око", "novi@example.rs"])  # already exists (different contact) -> skipped
    ws.append(["", "нема имена"])              # blank name -> skipped
    ws.append(["Нова организација", ""])
    wb.save(src)
    wb.close()

    tab = _make_tab(app, home)
    try:
        tab.state.observer_registry.orgs.append(
            AccreditingOrg(name="НВО Око", contact="stari@example.rs")
        )

        monkeypatch.setattr(
            "biracki_odbor.gui.tabs.posmatraci_workspace.QFileDialog.getOpenFileName",
            lambda *a, **kw: (str(src), ""),
        )
        monkeypatch.setattr(
            "biracki_odbor.gui.tabs.posmatraci_workspace.QMessageBox.question",
            lambda *a, **kw: __import__(
                "PySide6.QtWidgets", fromlist=["QMessageBox"]
            ).QMessageBox.StandardButton.Yes,
        )

        tab._import_orgs_excel()

        registry = tab.state.observer_registry
        assert len(registry.orgs) == 2  # original + the one genuinely new row
        assert registry.orgs[0].contact == "stari@example.rs"  # untouched, not overwritten
        assert {o.name for o in registry.orgs} == {"НВО Око", "Нова организација"}
    finally:
        tab.deleteLater()


def test_export_orgs_template_writes_expected_headers(app, home, tmp_path, monkeypatch) -> None:
    import openpyxl

    dest = tmp_path / "orgs_template.xlsx"
    tab = _make_tab(app, home)
    try:
        monkeypatch.setattr(
            "biracki_odbor.gui.tabs.posmatraci_workspace.QFileDialog.getSaveFileName",
            lambda *a, **kw: (str(dest), ""),
        )
        tab._export_orgs_template()

        assert dest.exists()
        wb = openpyxl.load_workbook(dest)
        ws = wb.active
        assert [ws.cell(row=1, column=c).value for c in (1, 2)] == ["Name", "Contact"]
    finally:
        tab.deleteLater()


def test_export_observers_template_writes_expected_headers(app, home, tmp_path, monkeypatch) -> None:
    import openpyxl

    dest = tmp_path / "obs_template.xlsx"
    tab = _make_tab(app, home)
    try:
        monkeypatch.setattr(
            "biracki_odbor.gui.tabs.posmatraci_workspace.QFileDialog.getSaveFileName",
            lambda *a, **kw: (str(dest), ""),
        )
        tab._export_observers_template()

        assert dest.exists()
        wb = openpyxl.load_workbook(dest)
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, 8)]
        assert headers == [
            "Last name", "First name", "JMBG", "Organization",
            "PS assignment", "ID no. (opt.)", "Note (opt.)",
        ]
    finally:
        tab.deleteLater()


def test_export_template_cancelled_dialog_writes_nothing(app, home, tmp_path, monkeypatch) -> None:
    dest = tmp_path / "orgs_template.xlsx"
    tab = _make_tab(app, home)
    try:
        monkeypatch.setattr(
            "biracki_odbor.gui.tabs.posmatraci_workspace.QFileDialog.getSaveFileName",
            lambda *a, **kw: ("", ""),
        )
        tab._export_orgs_template()  # cancelled dialog -> must not raise or write
        assert not dest.exists()
    finally:
        tab.deleteLater()
