"""Tests for multi-line text preservation in generated .docx decisions.

Field report: a real coalition election-list full name (see
tests/test_safe_filename.py's REAL_WORLD_LIST_NAME) carries intentional
embedded newlines -- RIK's own official layout for long coalition
names wraps them onto specific lines, and that formatting is not
incidental, it must survive into generated Решења/Измене verbatim. A
raw "\n" character inside a single docx run's text is not a Word line
break -- Word silently collapses/ignores it -- so without converting
it to a real `<w:br/>` element, the original formatting vanishes.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from biracki_odbor.core.models import Person
from biracki_odbor.core.phase3_docx import (
    _add_mixed_para, _add_para, _add_run_lines, generate_documents,
)
from tests.test_jmbg import VALID_SAMPLES
from tests.test_phase1 import _make_oik
from tests.test_phase3 import _make_election


_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _break_count(run) -> int:
    return len(run._element.findall(f".//{_W_NS}br"))


REAL_WORLD_LIST_NAME = (
    "СРБИЈА ПРОТИВ НАСИЉА - МИРОСЛАВ МИКИ АЛЕКСИЋ - МАРИНИКА ТЕПИЋ (СТРАНКА СЛОБОДЕ И ПРАВДЕ, НАРОДНИ\n"
    "ПОКРЕТ СРБИЈЕ, ЗЕЛЕНО-ЛЕВИ ФРОНТ, НЕ ДАВИМО БЕОГРАД, ЕКОЛОШКИ УСТАНАК - ЋУТА, ДЕМОКРАТСКА СТРАНКА,\n"
    "ПОКРЕТ СЛОБОДНИХ ГРАЂАНА, СРБИЈА ЦЕНТАР, ЗАЈЕДНО, ПОКРЕТ ЗА ПРЕОКРЕТ, УДРУЖЕНИ СИНДИКАТИ СРБИЈЕ\n"
    "'СЛОГА'', НОВО ЛИЦЕ СРБИЈЕ)"
)


# ── Unit-level: _add_run_lines / _add_para / _add_mixed_para ─────────


def test_add_run_lines_splits_into_separate_runs_with_real_breaks() -> None:
    # Note: python-docx's Run.text getter appends a synthetic "\n" to
    # represent a trailing <w:br/> element when reading it back -- that's
    # its own round-trip convention, not evidence of a raw newline
    # character in the underlying XML. The actual break count (checked
    # via _break_count) is what proves a real <w:br/> element exists.
    doc = Document()
    p = doc.add_paragraph()
    _add_run_lines(p, "line one\nline two\nline three")
    assert len(p.runs) == 3
    assert p.runs[0].text.rstrip("\n") == "line one"
    assert p.runs[1].text.rstrip("\n") == "line two"
    assert p.runs[2].text.rstrip("\n") == "line three"
    assert _break_count(p.runs[0]) == 1
    assert _break_count(p.runs[1]) == 1
    assert _break_count(p.runs[2]) == 0


def test_add_run_lines_single_line_is_one_run_no_break() -> None:
    doc = Document()
    p = doc.add_paragraph()
    _add_run_lines(p, "no newlines here")
    assert len(p.runs) == 1
    assert _break_count(p.runs[0]) == 0


def test_add_para_preserves_embedded_newlines_as_real_breaks() -> None:
    doc = Document()
    p = _add_para(doc, "first\nsecond")
    assert len(p.runs) == 2
    assert _break_count(p.runs[0]) == 1


def test_add_para_empty_text_still_produces_one_run() -> None:
    """Preserves the original behaviour: an empty paragraph still gets
    a run so its base font/size is set."""
    doc = Document()
    p = _add_para(doc)
    assert len(p.runs) == 1
    assert p.runs[0].text == ""


def test_add_mixed_para_handles_newline_within_a_single_segment() -> None:
    doc = Document()
    p = _add_mixed_para(doc, [
        ("prefix ", False),
        ("multi\nline", True),
        (" suffix", False),
    ])
    texts = [r.text.rstrip("\n") for r in p.runs]
    assert texts == ["prefix ", "multi", "line", " suffix"]
    assert _break_count(p.runs[1]) == 1


# ── End-to-end: real coalition name through generate_documents ───────


def test_decision_doc_generation_with_real_multiline_party_name(tmp_path: Path) -> None:
    """The exact field scenario: a party's full name (as stored in
    ElectionConfig.party_full_names) carries embedded newlines. Must
    not crash, and the generated decision's proposal cell must render
    the name as real line breaks, not a collapsed single line."""
    election = _make_election()
    election.party_full_names["СНС"] = REAL_WORLD_LIST_NAME
    oik = _make_oik(1)
    persons = [Person(
        bm_number=1, function="Председник", party="СНС",
        first_name="Особа", last_name="Презимић",
        jmbg=VALID_SAMPLES[0], mesto="Уб", adresa="Адреса", telefon="0641234567",
    )]

    result = generate_documents(persons, oik, election, tmp_path / "resenja")
    assert result.files_written, [i.message_sr for i in result.report.issues]

    doc = Document(result.files_written[0])
    found_multiline_run = False
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if _break_count(r) > 0:
                            found_multiline_run = True
    assert found_multiline_run, (
        "expected at least one run with a real line break in the "
        "generated decision — the multi-line party name should not "
        "have been collapsed onto a single line"
    )
