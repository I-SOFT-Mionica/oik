"""Phone numbers arrive in every format; they're stored one way, and the
leading 0 is never lost (Excel dropping it is repaired, not copied)."""

from __future__ import annotations

import pytest

from biracki_odbor.core.phone import is_valid, normalize


@pytest.mark.parametrize("raw", [
    "0641234567", " 0641234567 ", "064 123 4567", "064-123-4567", "064/123-45-67",
    "064.123.4567", "(064) 123 4567", "064 123 4567",
    "+381641234567", "+381 64 123 4567", "00381641234567", "381641234567",
    641234567, 641234567.0, "641234567",
])
def test_every_way_of_writing_a_mobile_number_ends_up_the_same(raw) -> None:
    assert normalize(raw) == "0641234567"


@pytest.mark.parametrize("raw,expected", [
    ("011/123-4567", "0111234567"),
    ("014 123 456", "014123456"),
    (11234567, "011234567"),           # Belgrade landline, Excel ate the 0
    ("060 000 0001", "0600000001"),
])
def test_landlines_and_short_numbers_keep_their_leading_zero(raw, expected) -> None:
    assert normalize(raw) == expected


def test_a_leading_zero_is_never_stripped() -> None:
    for raw in ("0641234567", "00641234567", "0", "011234567"):
        assert normalize(raw).startswith("0")


def test_foreign_numbers_keep_their_plus() -> None:
    assert normalize("+49 30 1234567") == "+49301234567"


@pytest.mark.parametrize("raw", [None, "", "   ", True])
def test_empty_stays_empty(raw) -> None:
    assert normalize(raw) == ""
    assert is_valid(raw)


@pytest.mark.parametrize("raw", ["064-abc-1234", "zvati posle 17h", "06412345+67"])
def test_notes_and_letters_are_left_for_the_validator(raw) -> None:
    assert normalize(raw) == raw
    assert not is_valid(raw)


def test_phase2_import_stores_the_normalized_number(tmp_path) -> None:
    """End to end: Phase 1 party file, filled with a number typed with
    separators and one Excel stored as an integer (leading 0 lost), then
    imported in Phase 2 -- both arrive clean, with their 0."""
    import openpyxl

    from biracki_odbor.core.phase1_generate import generate_party_files
    from biracki_odbor.core.phase2_aggregate import aggregate_party_files
    from tests.fixtures.stalni_sastav import build_template
    from tests.test_phase1 import _make_oik
    from tests.test_phase2 import _fill_party_file
    from tests.test_jmbg import VALID_SAMPLES

    template = build_template(tmp_path / "t.xlsx", n_polling_stations=1, parties=("СНС",))
    oik = _make_oik(1)
    party_file = generate_party_files(template, tmp_path / "stranke", oik).files_written[0]
    wb = openpyxl.load_workbook(party_file)
    rows = wb.active.max_row - 3
    wb.close()
    _fill_party_file(party_file, member_count=rows, jmbgs=VALID_SAMPLES)
    wb = openpyxl.load_workbook(party_file)
    wb.active.cell(row=4, column=10, value="064 123-4567")
    wb.active.cell(row=5, column=10, value=651234567)
    wb.save(party_file)
    wb.close()

    popunjeni = tmp_path / "popunjeni"
    popunjeni.mkdir()
    party_file.rename(popunjeni / party_file.name)
    result = aggregate_party_files(popunjeni, template, tmp_path / "m" / "z.xlsx", oik)

    phones = {p.telefon for p in result.persons if p.telefon}
    assert {"0641234567", "0651234567"} <= phones
    assert all(ph.startswith("0") for ph in phones)
