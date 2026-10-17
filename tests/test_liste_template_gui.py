"""GUI-level test for the "Generate template..." (liste sablon) button.

Confirms _generate_liste_template() actually computes available_levels
from the checked election-type checkboxes and passes it through to
build_entities_template() -- not just that the underlying function
works in isolation (covered in tests/test_candidacies.py).
"""

from __future__ import annotations

from pathlib import Path

import pytest


@pytest.fixture(autouse=True)
def _isolated_home(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("BIRACKI_ODBOR_HOME", str(tmp_path))
    from biracki_odbor import _release_info
    monkeypatch.setattr(_release_info, "GITHUB_REPO", "", raising=False)


@pytest.fixture
def app():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def _make_state(tmp_path):
    from biracki_odbor.gui.app_state import AppState
    from tests.test_phase1 import _make_oik
    from tests.test_phase3 import _make_election

    state = AppState()
    state._oik = _make_oik(2)  # noqa: SLF001
    state._election = _make_election()  # noqa: SLF001
    return state


def test_generate_liste_template_narrows_levels_to_checked_types(
    app, tmp_path, monkeypatch,
) -> None:
    import openpyxl
    from PySide6.QtWidgets import QFileDialog
    from biracki_odbor.gui.tabs.settings_election import ElectionSettingsTab

    state = _make_state(tmp_path)
    tab = ElectionSettingsTab(state)
    try:
        # Only "local" checked this cycle -- matches the "we only have
        # local elections, no need to keep selecting it" scenario.
        for key, cb in tab._type_checks.items():  # noqa: SLF001
            cb.setChecked(key == "local")

        out_path = tmp_path / "liste_sablon.xlsx"
        monkeypatch.setattr(
            QFileDialog, "getSaveFileName", lambda *a, **k: (str(out_path), ""),
        )
        monkeypatch.setattr("os.startfile", lambda *_a, **_k: None, raising=False)

        tab._generate_liste_template()  # noqa: SLF001

        wb = openpyxl.load_workbook(out_path)
        lookup = wb["_lookup"]
        levels = [lookup.cell(row=r, column=2).value for r in range(1, 3)]
        assert levels == ["local", None]
        # Single active level -> «Повезано са» dropped entirely.
        ws = wb["Листе"]
        headers = [ws.cell(row=1, column=c).value for c in range(1, 7)]
        assert headers[-1] is None
        wb.close()
    finally:
        tab.deleteLater()


def test_generate_liste_template_keeps_linked_column_for_multi_level(
    app, tmp_path, monkeypatch,
) -> None:
    import openpyxl
    from PySide6.QtWidgets import QFileDialog
    from biracki_odbor.gui.tabs.settings_election import ElectionSettingsTab

    state = _make_state(tmp_path)
    tab = ElectionSettingsTab(state)
    try:
        for key, cb in tab._type_checks.items():  # noqa: SLF001
            cb.setChecked(key in ("local", "parliamentary"))

        out_path = tmp_path / "liste_sablon.xlsx"
        monkeypatch.setattr(
            QFileDialog, "getSaveFileName", lambda *a, **k: (str(out_path), ""),
        )
        monkeypatch.setattr("os.startfile", lambda *_a, **_k: None, raising=False)

        tab._generate_liste_template()  # noqa: SLF001

        wb = openpyxl.load_workbook(out_path)
        ws = wb["Листе"]
        assert ws.cell(row=1, column=6).value == "Повезано са"
        wb.close()
    finally:
        tab.deleteLater()
