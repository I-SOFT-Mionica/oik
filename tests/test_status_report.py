"""Tests for core.status_report — dated status-report .docx."""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from biracki_odbor.config.election import ElectionConfig, ElectionInfo, ElectionDates
from biracki_odbor.core.models import Person
from biracki_odbor.core.status_report import generate_status_report
from biracki_odbor.core.workflow import PROSIRENI_SPEC
from tests.test_jmbg import VALID_SAMPLES
from tests.test_phase1 import _make_oik


@pytest.fixture
def app():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


@pytest.fixture
def home(tmp_path, monkeypatch):
    h = tmp_path / "home"
    monkeypatch.setenv("BIRACKI_ODBOR_HOME", str(h))
    monkeypatch.setattr("biracki_odbor._release_info.GITHUB_REPO", "", raising=False)
    return h


def _make_election() -> ElectionConfig:
    return ElectionConfig(
        election=ElectionInfo(dates=ElectionDates(date_full="25. октобра 2026. године")),
    )


def _persons() -> list[Person]:
    return [
        Person(bm_number=1, function="Председник", party="СНС",
               first_name="Пера", last_name="Перић", jmbg=VALID_SAMPLES[0]),
        Person(bm_number=1, function="Члан 1", party="СПС"),  # empty
        Person(bm_number=2, function="Председник", party="СНС",
               first_name="Мика", last_name="Микић", jmbg=VALID_SAMPLES[1]),
    ]


def test_status_report_contains_totals_and_party_table(tmp_path: Path) -> None:
    out = tmp_path / "status.docx"
    generate_status_report(_persons(), _make_oik(2), _make_election(), out)

    doc = Document(out)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Укупно места: 3" in full_text
    assert "Попуњено: 2" in full_text
    assert "Празно: 1" in full_text
    assert "Бирачких места: 2" in full_text

    assert len(doc.tables) == 1
    table_text = [cell.text for row in doc.tables[0].rows for cell in row.cells]
    assert "СНС" in table_text
    assert "СПС" in table_text


def test_status_report_workflow_label_and_party_column(tmp_path: Path) -> None:
    out_stalni = tmp_path / "s.docx"
    generate_status_report(_persons(), _make_oik(2), _make_election(), out_stalni)
    stalni_text = "\n".join(p.text for p in Document(out_stalni).paragraphs)
    assert "Стални састав" in stalni_text

    out_prosireni = tmp_path / "p.docx"
    generate_status_report(
        _persons(), _make_oik(2), _make_election(), out_prosireni, spec=PROSIRENI_SPEC,
    )
    doc = Document(out_prosireni)
    prosireni_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Проширени састав" in prosireni_text
    header = [c.text for c in doc.tables[0].rows[0].cells]
    assert header[0] == "Листа"


def test_export_button_writes_report(app, home, monkeypatch, tmp_path) -> None:
    """Корекције → Статус → "Export status report…" button, wired
    end-to-end (file-dialog path chosen, no errors, file written)."""
    from biracki_odbor import app_paths
    from biracki_odbor.core.master_io import write_master
    from biracki_odbor.gui.app_state import AppState
    from biracki_odbor.gui.tabs.corrections import CorrectionsTab
    from PySide6.QtWidgets import QFileDialog

    write_master(_persons(), _make_oik(2), app_paths.master_path())
    chosen = tmp_path / "my_report.docx"
    monkeypatch.setattr(
        QFileDialog, "getSaveFileName", lambda *a, **kw: (str(chosen), ""),
    )

    state = AppState()
    tab = CorrectionsTab(state)
    try:
        tab.reload_from_master()
        tab.summary_panel._on_export_report()  # noqa: SLF001
        assert chosen.exists()
        assert Document(chosen).tables
    finally:
        tab.deleteLater()
