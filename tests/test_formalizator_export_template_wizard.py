"""Tests for formalizator.gui.export_template_wizard — the "own screen"
for defining a cumulative export layout, decoupled from any specific
calibrated template."""

from __future__ import annotations

from pathlib import Path

import pytest

from formalizator.core.models import CustomFieldDef


@pytest.fixture
def app():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def test_field_options_include_builtins_and_custom_fields(app) -> None:
    from formalizator.gui.export_template_wizard import _field_options

    custom = [CustomFieldDef(key="overitelj_ime", label="Оверитељ — име и презиме", kind="universal")]
    options = _field_options(custom)
    keys = {k for k, _ in options}

    assert "jmbg" in keys  # whole string, not offered in the calibration wizard itself
    assert "full_name" in keys  # built-in
    assert "overitelj_ime" in keys  # this document's own custom field
    # municipality/certifier_name etc. no longer exist as fixed
    # SessionInfo fields — a document that needs them defines its own
    # "universal" custom field instead (see overitelj_ime above).
    assert "municipality" not in keys


def test_only_one_jmbg_option_is_offered_for_export(app) -> None:
    """The calibration wizard's comb key means "13 positioned digit
    boxes on paper" — meaningless in a spreadsheet cell. Offering it
    here alongside the real one put two identical-looking "ЈМБГ"
    choices in the picker, and picking the comb silently exported a
    blank column (hit on a real export)."""
    from formalizator.gui.export_template_wizard import _JMBG_COMB_KEY, _field_options

    options = _field_options([])
    keys = [k for k, _ in options]

    assert _JMBG_COMB_KEY not in keys
    assert keys.count("jmbg") == 1
    jmbg_labels = [label for key, label in options if "ЈМБГ" in label]
    assert len(jmbg_labels) == 1


def test_row_number_counter_is_offered_for_export(app) -> None:
    from formalizator.core.models import ROW_NUMBER_FIELD_KEY
    from formalizator.gui.export_template_wizard import _field_options

    assert ROW_NUMBER_FIELD_KEY in {k for k, _ in _field_options([])}


def _set_column_fields(dlg, row, field_keys, separators=None):
    """Test helper standing in for the operator clicking a column's
    field button and using the join-builder sub-dialog."""
    btn = dlg.table.cellWidget(row, 1)
    btn._field_keys = list(field_keys)  # noqa: SLF001
    btn._separators = list(separators or [])  # noqa: SLF001
    btn.setText(dlg._column_summary(btn._field_keys, btn._separators))  # noqa: SLF001


def test_add_column_and_save_produces_export_template(app) -> None:
    from formalizator.gui.export_template_wizard import ExportTemplateDialog

    dlg = ExportTemplateDialog()
    try:
        dlg.ed_name.setText("Тест извоз")
        dlg._add_column()  # noqa: SLF001
        ed_header = dlg.table.cellWidget(0, 0)
        ed_header.setText("Име и презиме")
        _set_column_fields(dlg, 0, ["full_name"])

        dlg._on_save()  # noqa: SLF001

        result = dlg.result_template()
        assert result is not None
        assert result.name == "Тест извоз"
        assert result.columns[0].header == "Име и презиме"
        assert result.columns[0].field_keys == ["full_name"]
    finally:
        dlg.deleteLater()


def test_save_without_name_shows_warning(app, monkeypatch) -> None:
    import formalizator.gui.export_template_wizard as mod

    shown = []
    monkeypatch.setattr(mod.QMessageBox, "warning", lambda *a: shown.append(a))

    dlg = mod.ExportTemplateDialog()
    try:
        dlg._add_column("Header")  # noqa: SLF001
        _set_column_fields(dlg, 0, ["full_name"])
        dlg._on_save()  # noqa: SLF001

        assert dlg.result_template() is None
        assert len(shown) == 1
    finally:
        dlg.deleteLater()


def test_save_without_any_complete_column_shows_warning(app, monkeypatch) -> None:
    import formalizator.gui.export_template_wizard as mod

    shown = []
    monkeypatch.setattr(mod.QMessageBox, "warning", lambda *a: shown.append(a))

    dlg = mod.ExportTemplateDialog()
    try:
        dlg.ed_name.setText("Тест")
        dlg._add_column()  # noqa: SLF001  # header blank, no fields picked
        dlg._on_save()  # noqa: SLF001

        assert dlg.result_template() is None
        assert len(shown) == 1
    finally:
        dlg.deleteLater()


def test_incomplete_column_skipped_silently(app) -> None:
    """A row with only a header or only fields (not both) is dropped,
    not treated as an error — same tolerance as the calibration
    wizard's own "Поље" combo."""
    from formalizator.gui.export_template_wizard import ExportTemplateDialog

    dlg = ExportTemplateDialog()
    try:
        dlg.ed_name.setText("Тест")
        dlg._add_column("Само наслов")  # noqa: SLF001  # no field picked
        dlg._add_column("Име")  # noqa: SLF001
        _set_column_fields(dlg, 1, ["full_name"])

        dlg._on_save()  # noqa: SLF001

        result = dlg.result_template()
        assert result is not None
        assert len(result.columns) == 1
        assert result.columns[0].header == "Име"
    finally:
        dlg.deleteLater()


def test_column_can_join_several_fields_with_own_separators(app) -> None:
    """The whole point of the redesign: the operator builds an
    arbitrary N-field join themselves, not something this app
    hardcodes for a specific document."""
    from formalizator.gui.export_template_wizard import ExportTemplateDialog

    dlg = ExportTemplateDialog()
    try:
        dlg.ed_name.setText("Тест")
        dlg._add_column("Оверитељ")  # noqa: SLF001
        _set_column_fields(dlg, 0, ["opstina", "overitelj_ime", "overitelj_adresa"], [" - ", ", "])

        dlg._on_save()  # noqa: SLF001

        result = dlg.result_template()
        assert result is not None
        assert result.columns[0].field_keys == ["opstina", "overitelj_ime", "overitelj_adresa"]
        assert result.columns[0].separators == [" - ", ", "]
    finally:
        dlg.deleteLater()


def test_column_fields_dialog_round_trips_fields_and_separators(app) -> None:
    from formalizator.gui.export_template_wizard import _ColumnFieldsDialog

    dlg = _ColumnFieldsDialog([], ["surname", "given_name", "jmbg"], [" - ", ", "])
    try:
        keys, seps = dlg.result_fields()
        assert keys == ["surname", "given_name", "jmbg"]
        assert seps == [" - ", ", "]
    finally:
        dlg.deleteLater()


def test_column_fields_dialog_skips_incomplete_rows(app) -> None:
    from formalizator.gui.export_template_wizard import _ColumnFieldsDialog

    dlg = _ColumnFieldsDialog([])
    try:
        dlg._add_row("surname")  # noqa: SLF001  # blank starter row + this one
        keys, seps = dlg.result_fields()
        assert keys == ["surname"]
        assert seps == []
    finally:
        dlg.deleteLater()


def test_load_template_restores_name_and_columns(app) -> None:
    from formalizator.core.models import ExportColumnDef, ExportTemplateDef
    from formalizator.gui.export_template_wizard import ExportTemplateDialog

    template = ExportTemplateDef(
        name="Постојећи",
        columns=[ExportColumnDef(header="ЈМБГ", field_keys=["jmbg"])],
    )

    dlg = ExportTemplateDialog()
    try:
        dlg.load_template(template)

        assert dlg.ed_name.text() == "Постојећи"
        assert dlg.table.rowCount() == 1
        ed_header = dlg.table.cellWidget(0, 0)
        btn_field = dlg.table.cellWidget(0, 1)
        assert ed_header.text() == "ЈМБГ"
        assert btn_field._field_keys == ["jmbg"]  # noqa: SLF001
    finally:
        dlg.deleteLater()


def test_load_template_restores_sort_columns(app) -> None:
    from formalizator.core.models import ExportColumnDef, ExportTemplateDef
    from formalizator.gui.export_template_wizard import ExportTemplateDialog

    template = ExportTemplateDef(
        name="Постојећи",
        columns=[
            ExportColumnDef(header="Оверитељ", field_keys=["overitelj"]),
            ExportColumnDef(header="Име", field_keys=["full_name"]),
        ],
        sort_column_1="Оверитељ", sort_column_2="Име",
    )

    dlg = ExportTemplateDialog()
    try:
        dlg.load_template(template)

        assert dlg.cb_sort1.currentData() == "Оверитељ"
        assert dlg.cb_sort2.currentData() == "Име"
    finally:
        dlg.deleteLater()


def test_save_captures_selected_sort_columns(app) -> None:
    from formalizator.gui.export_template_wizard import ExportTemplateDialog

    dlg = ExportTemplateDialog()
    try:
        dlg.ed_name.setText("Тест")
        dlg._add_column("Оверитељ")  # noqa: SLF001
        _set_column_fields(dlg, 0, ["overitelj"])
        dlg._add_column("Име")  # noqa: SLF001
        _set_column_fields(dlg, 1, ["full_name"])

        idx1 = dlg.cb_sort1.findData("Оверитељ")
        dlg.cb_sort1.setCurrentIndex(idx1)
        idx2 = dlg.cb_sort2.findData("Име")
        dlg.cb_sort2.setCurrentIndex(idx2)

        dlg._on_save()  # noqa: SLF001

        result = dlg.result_template()
        assert result is not None
        assert result.sort_column_1 == "Оверитељ"
        assert result.sort_column_2 == "Име"
    finally:
        dlg.deleteLater()


def _write_workbook_with_headers(path, headers: list[str], header_row: int = 1, title_rows: int = 0) -> None:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, title_rows + 1):
        ws.cell(row=r, column=1, value=f"Насловни ред {r}")
    for col, header in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=header)
    wb.save(path)


def test_read_header_row_returns_given_row_values(tmp_path) -> None:
    from formalizator.gui.export_template_wizard import read_header_row

    path = tmp_path / "roster.xlsx"
    _write_workbook_with_headers(path, ["Презиме и име", "ЈМБГ", "Овлашћени оверитељ"])

    assert read_header_row(path) == ["Презиме и име", "ЈМБГ", "Овлашћени оверитељ"]


def test_read_header_row_skips_blank_cells(tmp_path) -> None:
    from formalizator.gui.export_template_wizard import read_header_row

    path = tmp_path / "roster.xlsx"
    _write_workbook_with_headers(path, ["Име", "", "ЈМБГ"])

    assert read_header_row(path) == ["Име", "ЈМБГ"]


def test_read_header_row_at_a_specific_row(tmp_path) -> None:
    """Real official spreadsheets (e.g. NPRS-4) don't have their
    headers on row 1 — a title/tag row comes first."""
    from formalizator.gui.export_template_wizard import read_header_row

    path = tmp_path / "roster.xlsx"
    _write_workbook_with_headers(path, ["Име", "ЈМБГ"], header_row=3, title_rows=2)

    assert read_header_row(path, header_row=3) == ["Име", "ЈМБГ"]
    assert read_header_row(path, header_row=1) == ["Насловни ред 1"]


def test_read_header_positions_keeps_real_column_numbers(tmp_path) -> None:
    """Export must write back to a header's real column, not to its
    position in a possibly-shorter mapping list — this is what makes
    that possible."""
    from formalizator.gui.export_template_wizard import read_header_positions

    path = tmp_path / "roster.xlsx"
    _write_workbook_with_headers(path, ["Име", "", "ЈМБГ"])  # blank column B skipped

    assert read_header_positions(path) == [(1, "Име"), (3, "ЈМБГ")]


def test_guess_header_row_skips_a_lone_title_cell(tmp_path) -> None:
    from formalizator.gui.export_template_wizard import _guess_header_row

    path = tmp_path / "roster.xlsx"
    _write_workbook_with_headers(path, ["Име", "ЈМБГ"], header_row=3, title_rows=2)

    assert _guess_header_row(path) == 3


def test_pick_file_populates_columns_from_detected_headers(app, monkeypatch, tmp_path) -> None:
    """The whole point: an operator with an existing roster file loads
    it and gets its real column headers to map, instead of retyping
    them from memory."""
    import formalizator.gui.export_template_wizard as mod

    path = tmp_path / "roster.xlsx"
    _write_workbook_with_headers(path, ["Презиме и име", "ЈМБГ"])
    monkeypatch.setattr(
        mod.QFileDialog, "getOpenFileName", staticmethod(lambda *a, **kw: (str(path), "")),
    )

    dlg = mod.ExportTemplateDialog()
    try:
        dlg._on_pick_file()  # noqa: SLF001

        assert dlg.table.rowCount() == 2
        assert dlg.table.cellWidget(0, 0).text() == "Презиме и име"
        assert dlg.table.cellWidget(1, 0).text() == "ЈМБГ"
        # Field mapping is left for the operator to choose — a detected
        # header isn't the same as knowing which field feeds it.
        assert dlg.table.cellWidget(0, 1)._field_keys == []  # noqa: SLF001
        assert dlg.ed_name.text() == "roster"
        assert dlg.sb_header_row.value() == 1
        assert dlg.sb_header_row.isEnabled()
        # Real column positions captured, not just the header text —
        # this is what export writes back to.
        assert dlg.table.cellWidget(0, 0)._column_index == 1  # noqa: SLF001
        assert dlg.table.cellWidget(1, 0)._column_index == 2  # noqa: SLF001
    finally:
        dlg.deleteLater()


def test_column_index_round_trips_through_save_and_load(app, monkeypatch, tmp_path) -> None:
    """A skipped file column (e.g. "Редни број" left unmapped) must
    still leave the mapped columns after it pointing at their own real
    positions once reloaded for editing."""
    import formalizator.gui.export_template_wizard as mod

    path = tmp_path / "roster.xlsx"
    _write_workbook_with_headers(path, ["Редни број", "Презиме и име", "ЈМБГ", "Оверитељ"])
    monkeypatch.setattr(
        mod.QFileDialog, "getOpenFileName", staticmethod(lambda *a, **kw: (str(path), "")),
    )

    dlg = mod.ExportTemplateDialog()
    try:
        dlg._on_pick_file()  # noqa: SLF001
        dlg.table.removeRow(0)  # operator skips "Редни број"
        dlg.ed_name.setText("Тест")
        for row in range(dlg.table.rowCount()):
            _set_column_fields(dlg, row, ["full_name"])

        dlg._on_save()  # noqa: SLF001
        result = dlg.result_template()
        assert result is not None
        assert [c.column_index for c in result.columns] == [2, 3, 4]

        dlg2 = mod.ExportTemplateDialog()
        try:
            dlg2.load_template(result)
            assert [
                dlg2.table.cellWidget(r, 0)._column_index  # noqa: SLF001
                for r in range(dlg2.table.rowCount())
            ] == [2, 3, 4]
        finally:
            dlg2.deleteLater()
    finally:
        dlg.deleteLater()


def test_pick_file_with_title_rows_detects_the_real_header_row(app, monkeypatch, tmp_path) -> None:
    import formalizator.gui.export_template_wizard as mod

    path = tmp_path / "roster.xlsx"
    _write_workbook_with_headers(path, ["Презиме и име", "ЈМБГ"], header_row=3, title_rows=2)
    monkeypatch.setattr(
        mod.QFileDialog, "getOpenFileName", staticmethod(lambda *a, **kw: (str(path), "")),
    )

    dlg = mod.ExportTemplateDialog()
    try:
        dlg._on_pick_file()  # noqa: SLF001

        assert dlg.sb_header_row.value() == 3
        assert dlg.table.rowCount() == 2
        assert dlg.table.cellWidget(0, 0).text() == "Презиме и име"
    finally:
        dlg.deleteLater()


def test_pick_file_with_no_headers_shows_warning(app, monkeypatch, tmp_path) -> None:
    import formalizator.gui.export_template_wizard as mod

    path = tmp_path / "empty.xlsx"
    _write_workbook_with_headers(path, [])
    monkeypatch.setattr(
        mod.QFileDialog, "getOpenFileName", staticmethod(lambda *a, **kw: (str(path), "")),
    )
    shown = []
    monkeypatch.setattr(mod.QMessageBox, "warning", lambda *a: shown.append(a))

    dlg = mod.ExportTemplateDialog()
    try:
        dlg._on_pick_file()  # noqa: SLF001

        assert dlg.table.rowCount() == 0
        assert len(shown) == 1
    finally:
        dlg.deleteLater()


def test_reload_headers_uses_corrected_header_row(app, monkeypatch, tmp_path) -> None:
    """Operator corrects the guessed header row in the spin box and
    reloads — without needing to re-pick the file."""
    import formalizator.gui.export_template_wizard as mod

    path = tmp_path / "roster.xlsx"
    # Row 1 looks header-ish too (2 non-empty cells) but is actually a
    # title line — the guess picks it, and needs correcting to row 2,
    # the real header.
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="Изборна листа")
    ws.cell(row=1, column=2, value="Мионица")
    ws.cell(row=2, column=1, value="Име")
    ws.cell(row=2, column=2, value="ЈМБГ")
    wb.save(path)
    monkeypatch.setattr(
        mod.QFileDialog, "getOpenFileName", staticmethod(lambda *a, **kw: (str(path), "")),
    )

    dlg = mod.ExportTemplateDialog()
    try:
        dlg._on_pick_file()  # noqa: SLF001
        assert dlg.sb_header_row.value() == 1  # guess landed on the wrong row

        dlg.sb_header_row.setValue(2)
        dlg._on_reload_headers()  # noqa: SLF001

        assert dlg.table.rowCount() == 2
        assert dlg.table.cellWidget(0, 0).text() == "Име"
        assert dlg.table.cellWidget(1, 0).text() == "ЈМБГ"
    finally:
        dlg.deleteLater()


def test_save_with_source_file_and_home_dir_copies_it_into_the_catalog(app, tmp_path) -> None:
    """The pristine template file is copied into stable app storage on
    save, not referenced at its original (possibly removable/movable)
    path."""
    import formalizator.gui.export_template_wizard as mod

    source = tmp_path / "picked" / "roster.xlsx"
    source.parent.mkdir()
    _write_workbook_with_headers(source, ["Име", "ЈМБГ"])
    home = tmp_path / "home"

    dlg = mod.ExportTemplateDialog(home_dir=home)
    try:
        dlg._source_file = str(source)  # noqa: SLF001  # stand-in for _on_pick_file
        dlg.sb_header_row.setEnabled(True)
        dlg.sb_header_row.setValue(1)
        dlg.ed_name.setText("Тест")
        dlg._add_column("Име")  # noqa: SLF001
        _set_column_fields(dlg, 0, ["full_name"])

        dlg._on_save()  # noqa: SLF001

        result = dlg.result_template()
        assert result is not None
        assert result.source_file != str(source)
        assert Path(result.source_file).exists()
        assert Path(result.source_file).read_bytes() == source.read_bytes()
        assert result.header_row == 1
    finally:
        dlg.deleteLater()


def test_save_with_already_copied_source_file_does_not_error(app, tmp_path) -> None:
    """Re-opening an already-saved template via "Уреди…" sets
    `_source_file` to the path it was already copied to on the earlier
    save — saving again (even under the same name) must not try to
    copy that file onto itself. Reported directly: this raised
    shutil.SameFileError and showed an error dialog instead of saving."""
    import formalizator.gui.export_template_wizard as mod

    home = tmp_path / "home"
    dest_dir = home / "export_sources"
    dest_dir.mkdir(parents=True)
    dest = dest_dir / "Тест.xlsx"
    _write_workbook_with_headers(dest, ["Име", "ЈМБГ"])

    dlg = mod.ExportTemplateDialog(home_dir=home)
    try:
        dlg._source_file = str(dest)  # noqa: SLF001  # stand-in for load_template
        dlg.sb_header_row.setEnabled(True)
        dlg.sb_header_row.setValue(1)
        dlg.ed_name.setText("Тест")
        dlg._add_column("Име")  # noqa: SLF001
        _set_column_fields(dlg, 0, ["full_name"])

        dlg._on_save()  # noqa: SLF001

        result = dlg.result_template()
        assert result is not None
        assert result.source_file == str(dest)
    finally:
        dlg.deleteLater()
