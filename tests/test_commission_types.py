"""Commission wording by unit type (TODO §31): municipality / city / city
municipality. Forms verified against real 2026 decisions of ОИК Мионица,
ГИК Ваљево and ИК ГО Лазаревац."""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from docx import Document

from biracki_odbor.config import opstine as opstine_mod
from biracki_odbor.config.commission import (
    commission_terms, default_full_committee_name, terms_for,
)
from biracki_odbor.config.oik import Municipality
from biracki_odbor.config.opstine import VRSTA_GRAD, VRSTA_GRADSKA_OPSTINA, VRSTA_OPSTINA

VALJEVO = Municipality(
    name="Ваљево", name_genitive="Ваљева", name_locative="Ваљеву", sifra_opstine="70360",
    full_committee_name=default_full_committee_name(VRSTA_GRAD, "Ваљево", "Ваљева"),
)
LAZAREVAC = Municipality(
    name="Лазаревац", name_genitive="Лазаревца", name_locative="Лазаревцу", sifra_opstine="70165",
    full_committee_name=default_full_committee_name(VRSTA_GRADSKA_OPSTINA, "Лазаревац"),
)


def test_registry_classifies_cities_and_city_municipalities() -> None:
    reg = opstine_mod.load_registry()
    by_type = {v: [e for e in reg if e.vrsta == v] for v in (VRSTA_OPSTINA, VRSTA_GRAD, VRSTA_GRADSKA_OPSTINA)}
    assert (len(by_type[VRSTA_GRAD]), len(by_type[VRSTA_GRADSKA_OPSTINA])) == (27, 17)
    for e in by_type[VRSTA_GRAD] + by_type[VRSTA_GRADSKA_OPSTINA]:
        assert not re.search(r" - град$| град$|^Град |^Београд-", e.document_name), e
    assert opstine_mod.find_by_sifra("70360").vrsta == VRSTA_GRAD          # Ваљево
    assert opstine_mod.find_by_sifra("70831").vrsta == VRSTA_OPSTINA       # Мионица


@pytest.mark.parametrize("mun, preamble, in_unit, subject_gen, assembly", [
    (Municipality(name="Мионица", sifra_opstine="70831"),
     "Општинска изборна комисија општине Мионица", "у општини Мионица",
     "Општинске изборне комисије", "Скупштине општине Мионица"),
    (VALJEVO,
     "Градска изборна комисија за спровођење избора за одборнике Скупштине града Ваљева",
     "у граду Ваљеву", "Градске изборне комисије", "Скупштине града Ваљева"),
    (LAZAREVAC,
     "Изборна комисија градске општине Лазаревац", "у градској општини Лазаревац",
     "Изборне комисије градске општине Лазаревац", "Скупштине градске општине Лазаревац"),
])
def test_terms_per_type(mun, preamble, in_unit, subject_gen, assembly) -> None:
    t = commission_terms(mun)
    assert (t.preamble_subject, t.in_unit, t.subject_gen, t.local_assembly) == (
        preamble, in_unit, subject_gen, assembly,
    )


def test_unknown_code_behaves_as_municipality() -> None:
    t = commission_terms(Municipality(name="Уб"))
    assert t.preamble_subject == "Општинска изборна комисија општине Уб"


@pytest.mark.parametrize("mun, expected", [
    (Municipality(name="Мионица", sifra_opstine="70831"), "начелника Општинске управе општине Мионица"),
    (VALJEVO, "начелника Градске управе града Ваљева"),
    (LAZAREVAC, "начелника Управе градске општине Лазаревац"),
])
def test_chief_title_follows_unit_type(mun, expected) -> None:
    from biracki_odbor.config.oik import OikConstants
    assert OikConstants(municipality=mun).nacelnik_proposal_name == expected


def test_custom_chief_title_is_kept() -> None:
    from biracki_odbor.config.oik import OikConstants
    oik = OikConstants(municipality=VALJEVO, nacelnik_full_name="в.д. начелника Градске управе")
    assert oik.nacelnik_proposal_name == "в.д. начелника Градске управе града Ваљева"


def _docx_text(path: Path) -> str:
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


_MUNICIPALITY_WORDING = re.compile(r"Општинск|ОПШТИНСК|у општини|У ОПШТИНИ|(?<!градске )општине ")


@pytest.mark.parametrize("mun, must_contain", [
    (VALJEVO, ["Градска изборна комисија за спровођење избора за одборнике Скупштине града Ваљева, на",
               "У ГРАДУ ВАЉЕВУ", "у граду Ваљеву  именују се",
               "Градској изборној комисији у граду Ваљеву су достављени"]),
    (LAZAREVAC, ["Изборна комисија градске општине Лазаревац, на",
                 "У ГРАДСКОЈ ОПШТИНИ ЛАЗАРЕВАЦ",
                 "Изборној комисији градске општине Лазаревац су достављени"]),
])
def test_generated_decision_has_no_municipality_wording(tmp_path, monkeypatch, mun, must_contain) -> None:
    monkeypatch.setenv("BIRACKI_ODBOR_HOME", str(tmp_path / "home"))  # fresh seed of the default texts
    from biracki_odbor.core.phase3_docx import generate_documents
    from tests.test_phase3 import _make_election, _make_persons
    from tests.test_phase1 import _make_oik
    oik = _make_oik()
    oik.municipality = mun
    election = _make_election()
    election.election.types = ["parliamentary", "local"]
    res = generate_documents(_make_persons(), oik, election, tmp_path / "out", max_workers=1)
    text = _docx_text(res.files_written[0])
    for s in must_contain:
        assert s in text, s
    assert "одборнике Скупштине" in text
    assert not _MUNICIPALITY_WORDING.search(text), _MUNICIPALITY_WORDING.search(text).group(0)


def test_amendment_and_lapse_decisions_for_a_city(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("BIRACKI_ODBOR_HOME", str(tmp_path / "home"))
    from biracki_odbor.core.izmene_docx import generate_izmena
    from biracki_odbor.core.prestanak_docx import generate_prestanak
    from tests.test_izmene import _make_election_with_auth, _make_stalni_spec
    from tests.test_prestanak_docx import _make_election, _make_spec
    from tests.test_phase1 import _make_oik
    oik = _make_oik()
    oik.municipality = VALJEVO

    izm = _docx_text(generate_izmena(_make_stalni_spec(), oik, _make_election_with_auth(), tmp_path / "i"))
    assert "овлашћено лице Градске изборне комисије града Ваљева" in izm
    assert "ГРАДСКА ИЗБОРНА КОМИСИЈА ГРАДА ВАЉЕВА" in izm
    assert not _MUNICIPALITY_WORDING.search(izm), _MUNICIPALITY_WORDING.search(izm).group(0)

    pre = _docx_text(generate_prestanak(_make_spec(municipality="Ваљево", existing_municipality=""),
                                        oik, _make_election(), tmp_path / "p"))
    assert "Градска изборна комисија је утврдила" in pre
    assert not _MUNICIPALITY_WORDING.search(pre), _MUNICIPALITY_WORDING.search(pre).group(0)
