"""Build synthetic RIK RASPODELA/UKUPNO workbooks for tests.

Mirrors the real shape confirmed by inspecting an actual 2023 RIK export
(see `core/rik_raspodela_import.py`'s module docstring):

RASPODELA — one sheet, one block per municipality:
    Управни округ:        | | <district>
    Шифра управног округа:| | <district code>
    Општина:              | | <name>            (one block deliberately
    Шифра општине:        | | <sifra>             omits the colon, below)
    (blank)
    (blank, blank, party abbrevs...)
    Председник            | | <counts per party>
    Заменик               | | <counts per party>
    Чл.+Зам.чл.           | | <counts per party>
    Укупно                | | <counts per party>
    (blank)
    Број бирачког места | Председник | Заменик председника | ... (6 функција)
    <bm rows...>
    (blank separator before the next block)

UKUPNO — one sheet, "Напомена:" marker in column B, then short/long
party name pairs in columns B/C.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from biracki_odbor.core.models import FUNCTIONS


def build_raspodela_workbook(path: Path) -> Path:
    """Two municipality blocks: one normal, one using the real-world
    `Општина` (no colon) variant — pins the tolerance fix."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "test"

    def _write_block(
        start_row: int, *, okrug: str, okrug_sifra: str, name: str, sifra: str,
        opstina_label: str, parties: tuple[str, ...], n_bm: int,
    ) -> int:
        r = start_row
        ws.cell(row=r, column=1, value="Управни округ:")
        ws.cell(row=r, column=3, value=okrug)
        r += 1
        ws.cell(row=r, column=1, value="Шифра управног округа:")
        ws.cell(row=r, column=3, value=okrug_sifra)
        r += 1
        ws.cell(row=r, column=1, value=opstina_label)
        ws.cell(row=r, column=3, value=name)
        r += 1
        ws.cell(row=r, column=1, value="Шифра општине:")
        ws.cell(row=r, column=3, value=sifra)
        r += 2  # blank row

        header_row = r
        for c, party in enumerate(parties, start=3):
            ws.cell(row=header_row, column=c, value=party)
        r += 1
        for label, base in (("Председник", 1), ("Заменик", 1), ("Чл.+Зам.чл.", 4), ("Укупно", 6)):
            ws.cell(row=r, column=1, value=label)
            for c, _party in enumerate(parties, start=3):
                ws.cell(row=r, column=c, value=base)
            r += 1
        r += 1  # blank row

        ws.cell(row=r, column=1, value="Број бирачког места")
        for c, func in enumerate(FUNCTIONS, start=2):
            ws.cell(row=r, column=c, value=func)
        r += 1
        for i in range(n_bm):
            ws.cell(row=r, column=1, value=i + 1)
            for c, _func in enumerate(FUNCTIONS, start=2):
                ws.cell(row=r, column=c, value=parties[(i + c) % len(parties)])
            r += 1
        r += 2  # blank separator before next block
        return r

    parties = ("СНС", "СПС", "ССП")
    next_row = _write_block(
        1, okrug="КОЛУБАРСКИ", okrug_sifra="08", name="Уб", sifra="71218",
        opstina_label="Општина:", parties=parties, n_bm=3,
    )
    _write_block(
        next_row, okrug="ЗАЈЕЧАРСКИ", okrug_sifra="16", name="Сокобања",
        sifra="71129", opstina_label="Општина", parties=parties, n_bm=2,
    )

    wb.save(path)
    wb.close()
    return path


def build_raspodela_workbook_no_spacer(path: Path) -> Path:
    """One municipality block in the 2026-cycle shape, confirmed against
    a real RIK export: no blank spacer column between a marker's label
    and its value (value moved from column C into column B), and the
    quota table's 3rd row spelled out as "Чланови +Заменици чланова"
    instead of the older "Чл.+Зам.чл."."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "test"

    parties = ("СНС", "СПС", "ССП")
    r = 1
    ws.cell(row=r, column=1, value="Управни округ:")
    ws.cell(row=r, column=2, value="ГРАД БЕОГРАД")
    r += 1
    ws.cell(row=r, column=1, value="Шифра управног округа:")
    ws.cell(row=r, column=2, value="00")
    r += 1
    ws.cell(row=r, column=1, value="Општина:")
    ws.cell(row=r, column=2, value="Београд-Барајево")
    r += 1
    ws.cell(row=r, column=1, value="Шифра општине:")
    ws.cell(row=r, column=2, value="70092")
    r += 2  # blank row

    header_row = r
    for c, party in enumerate(parties, start=2):
        ws.cell(row=header_row, column=c, value=party)
    r += 1
    for label, base in (
        ("Председник", 1), ("Заменик", 1),
        ("Чланови +Заменици чланова", 4), ("Укупно", 6),
    ):
        ws.cell(row=r, column=1, value=label)
        for c, _party in enumerate(parties, start=2):
            ws.cell(row=r, column=c, value=base)
        r += 1
    r += 1  # blank row

    ws.cell(row=r, column=1, value="Број бирачког места")
    for c, func in enumerate(FUNCTIONS, start=2):
        ws.cell(row=r, column=c, value=func)
    r += 1
    for i in range(3):
        ws.cell(row=r, column=1, value=i + 1)
        for c, _func in enumerate(FUNCTIONS, start=2):
            ws.cell(row=r, column=c, value=parties[(i + c) % len(parties)])
        r += 1

    wb.save(path)
    wb.close()
    return path


def build_ukupno_workbook(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "test"
    ws.cell(row=1, column=2, value="Напомена:")
    pairs = (
        ("СНС", "Посланичка група СНС — пуни назив"),
        ("СПС", "Посланичка група СПС — пуни назив"),
        ("ССП", "Посланичка група ССП — пуни назив"),
    )
    for i, (short, long_name) in enumerate(pairs, start=2):
        ws.cell(row=i, column=2, value=short)
        ws.cell(row=i, column=3, value=long_name)
    wb.save(path)
    wb.close()
    return path


_LEGEND_PAIRS = (
    ("СНС", "Посланичка група СНС — пуни назив"),
    ("СПС", "Посланичка група СПС — пуни назив"),
    ("ССП", "Посланичка група ССП — пуни назив"),
)


def build_skracenice_workbook(path: Path) -> Path:
    """The 2026-cycle legend, confirmed against a real standalone
    "skracenice.xlsx": marker "Легенда:" in column A (not "Напомена:" in
    column B), short name in column A / long name in column B (not B/C)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Skraćenice"
    ws.cell(row=1, column=1, value="Посланичке групе")
    ws.cell(row=2, column=1, value="Легенда:")
    for i, (short, long_name) in enumerate(_LEGEND_PAIRS, start=4):
        ws.cell(row=i, column=1, value=short)
        ws.cell(row=i, column=2, value=long_name)
    wb.save(path)
    wb.close()
    return path


def build_combined_2026_workbook(path: Path) -> Path:
    """The 2026-cycle combined file, confirmed against a real
    "Raspodela_Parlamentarni_rik.xlsx": one workbook with the RASPODELA
    data as its FIRST sheet and the party-name legend as a LATER sheet
    (2nd of 4 in the real file) — not sheet 0. Pins that an operator
    pointing both the RASPODELA and UKUPNO file pickers at this same
    file must still get the real legend, not a silent empty one from
    only ever checking worksheets[0]."""
    build_raspodela_workbook(path)  # writes sheet "test" with 2 blocks
    wb = openpyxl.load_workbook(path)
    wb["test"].title = "00-24"
    legend_ws = wb.create_sheet("Skraćenice")
    legend_ws.cell(row=1, column=1, value="Посланичке групе")
    legend_ws.cell(row=2, column=1, value="Легенда:")
    for i, (short, long_name) in enumerate(_LEGEND_PAIRS, start=4):
        legend_ws.cell(row=i, column=1, value=short)
        legend_ws.cell(row=i, column=2, value=long_name)
    wb.save(path)
    wb.close()
    return path
