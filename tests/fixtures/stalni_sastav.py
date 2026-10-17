"""Build a synthetic Stalni_sastav.xlsx for tests.

Mirrors the structure expected by core.template:
  - 'bm' sheet: BM list (header row 1, data row 2+).
  - 'Raspodela' sheet: header row with 'БР БМ' + 6 function columns,
    then per-BM rows with party abbreviations.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from biracki_odbor.core.models import FUNCTIONS


def build_template(
    path: Path,
    *,
    n_polling_stations: int = 3,
    parties: tuple[str, ...] = ("СНС", "СПС", "ДС", "СРЦЕ", "НПС", "ЗЛФ"),
) -> Path:
    """Create a synthetic template. Returns the path."""
    wb = openpyxl.Workbook()

    # bm sheet
    ws_bm = wb.active
    ws_bm.title = "bm"
    ws_bm.cell(row=1, column=1, value="БР")
    ws_bm.cell(row=1, column=2, value="Назив")
    ws_bm.cell(row=1, column=3, value="Адреса")
    ws_bm.cell(row=1, column=4, value="Пун назив")
    for i in range(n_polling_stations):
        r = i + 2
        ws_bm.cell(row=r, column=1, value=i + 1)
        ws_bm.cell(row=r, column=2, value=f"БМ {i + 1}")
        ws_bm.cell(row=r, column=3, value=f"Адреса {i + 1}")
        ws_bm.cell(row=r, column=4, value=f"БМ {i + 1} (пуни назив)")

    # Raspodela sheet
    ws_r = wb.create_sheet("Raspodela")
    ws_r.cell(row=1, column=1, value="БР БМ")
    for c, func in enumerate(FUNCTIONS, start=2):
        ws_r.cell(row=1, column=c, value=func)
    for i in range(n_polling_stations):
        r = i + 2
        ws_r.cell(row=r, column=1, value=i + 1)
        for c, _ in enumerate(FUNCTIONS, start=2):
            party = parties[(i + c) % len(parties)]
            ws_r.cell(row=r, column=c, value=party)

    wb.save(path)
    wb.close()
    return path
