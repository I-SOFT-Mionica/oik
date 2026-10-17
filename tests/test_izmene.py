"""Tests for core.izmene_docx — amendment resolution generator."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from biracki_odbor.config.election import (
    DocumentInfo, ElectionConfig, ElectionDates, ElectionInfo,
)
from biracki_odbor.core.izmene_docx import (
    IzmenaRecord,
    IzmenaSpec,
    build_izmena_document,
    format_amendment_broj,
    generate_izmena,
)
from tests.test_phase1 import _make_oik


@pytest.fixture
def app():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


@pytest.fixture
def home(tmp_path, monkeypatch):
    h = tmp_path / "home"
    monkeypatch.setenv("BIRACKI_ODBOR_HOME", str(h))
    monkeypatch.setattr("biracki_odbor._release_info.GITHUB_REPO", "", raising=False)
    return h


def _make_state(tmp_path):
    from biracki_odbor.gui.app_state import AppState
    state = AppState()
    state.app_settings.default_files_dir = str(tmp_path / "export_base")
    return state


# ── Pure-function tests (no docx, no OIK) ────────────────────────────────────

def test_format_amendment_broj_standard() -> None:
    result = format_amendment_broj("013-5-30-9/2023", 1, 1)
    assert result == "Број: 013-5-30-9/1/2023-01"


def test_format_amendment_broj_second_amendment() -> None:
    result = format_amendment_broj("013-5-31-7/2023", 2, 2)
    assert result == "Број: 013-5-31-7/2/2023-02"


def test_format_amendment_broj_no_slash_fallback() -> None:
    result = format_amendment_broj("013-5-30-9", 1, 1)
    assert result == "Број: 013-5-30-9/1-01"


# ── Document structure tests ──────────────────────────────────────────────────

def _make_election_with_auth() -> ElectionConfig:
    return ElectionConfig(
        election=ElectionInfo(
            dates=ElectionDates(
                date_full="17. децембар 2023. године",
                session_date="6. децембра 2023. године",
                session_number="5",
            ),
            types=["parliamentary"],
            rik_decision="Одлуком РИК-а 02 број 013-376/23",
        ),
        document=DocumentInfo(
            number_prefix="013-5-30",
            legal_basis="На основу члана 32. Закона ...",
            authorization_number="013-5-32/2023-01",
            authorization_date="6. децембра 2023. године",
            authorized_name="Ивана Тешић",
            authorized_role="заменик председника",
        ),
        party_full_names={
            "НПС": "Посланичке групе Народни покрет Србије",
        },
    )


def _make_stalni_spec() -> IzmenaSpec:
    return IzmenaSpec(
        bm_number=9,
        bm_name="Кућа Јокић Милована, Врело",
        original_broj="013-5-30-9/2023",
        original_date="6. децембра 2023. године",
        amendment_index=1,
        amendment_counter=1,
        issue_date="12. децембра 2023. године",
        records=[
            IzmenaRecord(
                old_name="Никола Ареснић",
                function="ЧЛАНА 1",
                new_name="Славко Максимовић",
                new_place="Такова",
                party="НПС",
            )
        ],
        slug="stalni",
    )


def test_build_stalni_document_returns_document(tmp_path: Path) -> None:
    doc = build_izmena_document(_make_stalni_spec(), _make_oik(9), _make_election_with_auth())
    assert hasattr(doc, "paragraphs") and hasattr(doc, "save")


def test_stalni_document_contains_bm_header() -> None:
    doc = build_izmena_document(_make_stalni_spec(), _make_oik(9), _make_election_with_auth())
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "ЗА БИРАЧКО МЕСТО БРОЈ 9." in full_text


def test_stalni_document_contains_amendment_broj() -> None:
    doc = build_izmena_document(_make_stalni_spec(), _make_oik(9), _make_election_with_auth())
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "013-5-30-9/1/2023-01" in full_text


def test_stalni_document_references_authorization() -> None:
    doc = build_izmena_document(_make_stalni_spec(), _make_oik(9), _make_election_with_auth())
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "013-5-32/2023-01" in full_text
    assert "Ивана Тешић" in full_text


def test_stalni_document_contains_disposition_change() -> None:
    doc = build_izmena_document(_make_stalni_spec(), _make_oik(9), _make_election_with_auth())
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Никола Ареснић" in full_text
    assert "Славко Максимовић" in full_text
    assert "ЧЛАНА 1" in full_text


def test_stalni_document_obrazlozenje_has_page_break_before() -> None:
    """Field report: Образложење ran on directly after Item 3 on the same
    page instead of starting on page 2, unlike the original resenja
    (phase3_docx), which sets page_break_before on this exact heading."""
    doc = build_izmena_document(_make_stalni_spec(), _make_oik(9), _make_election_with_auth())
    found_pbb = False
    for p in doc.paragraphs:
        if "Образложење" in p.text or "О б р а з л о ж е њ е" in p.text:
            if p.paragraph_format.page_break_before is True:
                found_pbb = True
            break
    assert found_pbb, "Образложење paragraph must have page_break_before=True"


def test_stalni_document_has_blank_signature_line() -> None:
    """Field report: the signer's printed name rendered with nothing to
    sign above it -- must have an actual blank signature line (bordered
    empty table), same technique as the original resenja."""
    doc = build_izmena_document(_make_stalni_spec(), _make_oik(9), _make_election_with_auth())
    assert len(doc.tables) >= 1, "expected a signature-line table"
    sig_tbl = doc.tables[-1]
    assert len(sig_tbl.rows) == 1 and len(sig_tbl.columns) == 1
    cell = sig_tbl.rows[0].cells[0]
    assert cell.text.strip() == "", "signature line cell must be empty (signed by hand)"


def test_stalni_document_legal_article() -> None:
    doc = build_izmena_document(_make_stalni_spec(), _make_oik(9), _make_election_with_auth())
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "члана 35." in full_text


def test_fallback_obrazlozenje_uses_exact_official_text() -> None:
    """When no rendered_obrazlozenje/obraz_path is given (the
    corrections.py auto-generate-on-save path), the образложење must
    come from the same bundled official seed the cards system uses —
    not a separate, independently-drifting hardcoded copy. Field
    report: that copy was missing the чл. 41 ст. 1 citation and the
    3-day deadline entirely."""
    doc = build_izmena_document(_make_stalni_spec(), _make_oik(9), _make_election_with_auth())
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "најкасније три дана пре дана гласања" in full_text
    assert "Чланом 41. став 3." in full_text
    # The two non-official intro sentences from the old hardcoded copy
    # must be gone.
    assert "одмах по истеку рока за достављање кандидатских листа" not in full_text


def test_prosireni_document_uses_different_article() -> None:
    spec = _make_stalni_spec()
    spec.slug = "prosireni"
    spec.bm_number = 7
    spec.records = [
        IzmenaRecord(
            old_name="Драган Недељковић",
            function="ЗАМЕНИК ЧЛАНА 7",
            new_name="Драган Шпановић",
            new_place="Такова",
            party="НПС",
        )
    ]
    doc = build_izmena_document(spec, _make_oik(7), _make_election_with_auth())
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "члана 39 став 1." in full_text
    assert "ПРОШИРЕНОМ" in full_text


def test_prosireni_party_proposal_uses_izborne_liste() -> None:
    spec = _make_stalni_spec()
    spec.slug = "prosireni"
    spec.records = [
        IzmenaRecord(old_name="А", function="ЧЛАНА 1", new_name="Б",
                     new_place="Уб", party="НПС")
    ]
    doc = build_izmena_document(spec, _make_oik(7), _make_election_with_auth())
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Изборне листе" in full_text


def test_izmena_normal_style_uses_single_line_spacing() -> None:
    """See core.phase3_docx's identical fix: python-docx's blank-document
    template defaults Normal to 1.15 line spacing, never overridden here
    before now."""
    from docx.enum.text import WD_LINE_SPACING

    spec = _make_stalni_spec()
    doc = build_izmena_document(spec, _make_oik(7), _make_election_with_auth())
    assert doc.styles["Normal"].paragraph_format.line_spacing_rule == WD_LINE_SPACING.SINGLE


def test_izmena_header_and_footer_distance_is_1cm() -> None:
    spec = _make_stalni_spec()
    doc = build_izmena_document(spec, _make_oik(7), _make_election_with_auth())
    section = doc.sections[0]
    assert section.header_distance.cm == pytest.approx(1.0, abs=0.01)
    assert section.footer_distance.cm == pytest.approx(1.0, abs=0.01)


def test_generate_izmena_writes_file(tmp_path: Path) -> None:
    spec = _make_stalni_spec()
    out = generate_izmena(spec, _make_oik(9), _make_election_with_auth(), tmp_path)
    assert out.exists()
    assert out.suffix == ".docx"
    assert "STALNI" in out.name
    assert "izmena_01" in out.name


def test_generate_izmena_prosireni_filename(tmp_path: Path) -> None:
    spec = _make_stalni_spec()
    spec.slug = "prosireni"
    out = generate_izmena(spec, _make_oik(9), _make_election_with_auth(), tmp_path)
    assert "PROSIRENI" in out.name


# ── DocumentInfo round-trip ───────────────────────────────────────────────────

def test_document_info_roundtrip_with_auth_fields() -> None:
    info = DocumentInfo(
        number_prefix="013-5-30",
        legal_basis="На основу ...",
        authorization_number="013-5-32/2023-01",
        authorization_date="6. децембра 2023. године",
        authorized_name="Ивана Тешић",
        authorized_role="заменик председника",
    )
    restored = DocumentInfo.from_dict(info.to_dict())
    assert restored.authorization_number == "013-5-32/2023-01"
    assert restored.authorization_date == "6. децембра 2023. године"
    assert restored.authorized_name == "Ивана Тешић"
    assert restored.authorized_role == "заменик председника"


def test_document_info_backward_compat_missing_auth_fields() -> None:
    """Old configs without auth fields must deserialize without error."""
    old_dict = {"number_prefix": "013-5-30", "legal_basis": "На основу ..."}
    info = DocumentInfo.from_dict(old_dict)
    assert info.authorization_number == ""
    assert info.authorized_name == ""


def test_document_info_reads_pre_rename_izmena_prefixed_keys() -> None:
    """Delegation fields were renamed from izmena_* to generic names (the
    setting now applies to every resenje type, not just Измене) — an
    already-saved election.json still has the OLD key names on disk, and
    must keep loading real delegation data from them, not silently reset
    to blank just because the field was renamed in code."""
    old_dict = {
        "number_prefix": "013-5-30",
        "izmena_use_authorization": False,
        "izmena_authorization_number": "013-5-32/2023-01",
        "izmena_authorization_date": "6. децембра 2023. године",
        "izmena_authorized_name": "Ивана Тешић",
        "izmena_authorized_role": "заменик председника",
    }
    info = DocumentInfo.from_dict(old_dict)
    assert info.use_authorization is False
    assert info.authorization_number == "013-5-32/2023-01"
    assert info.authorization_date == "6. децембра 2023. године"
    assert info.authorized_name == "Ивана Тешић"
    assert info.authorized_role == "заменик председника"


def test_document_info_new_keys_take_priority_over_old(tmp_path: Path) -> None:
    """If a file somehow has both (e.g. saved once post-rename, then hand-
    edited), the new key wins — it's the one future writes will use."""
    mixed_dict = {
        "authorized_name": "Нови Потписник",
        "izmena_authorized_name": "Стари Потписник",
    }
    info = DocumentInfo.from_dict(mixed_dict)
    assert info.authorized_name == "Нови Потписник"


# ── events_to_izmena_records ────────────────────────────────────────────────
#
# Field report: after "Заустави пријем пријава", a member drops out and the
# party can't provide a replacement — the operator must "Очисти место"
# (clear_member) before "Именуј" (appoint_by_nacelnik) becomes available.
# That produces two separate audit events for one semantic correction, and
# auto-generation was producing nothing for it at all.


def _person_event(bm, function, *, before_first="", before_last="",
                   after_first="", after_last="", party="СНС",
                   appointed_by_nacelnik=False, event_type="edit_member",
                   before_jmbg="", after_jmbg="", before_mesto="", after_mesto=""):
    from biracki_odbor.core.corrections import AuditEvent
    from biracki_odbor.core.models import Person

    before = Person(
        bm_number=bm, function=function, party=party,
        first_name=before_first, last_name=before_last,
        jmbg=before_jmbg, mesto=before_mesto,
    ).to_dict()
    after = Person(
        bm_number=bm, function=function, party=party,
        first_name=after_first, last_name=after_last,
        appointed_by_nacelnik=appointed_by_nacelnik,
        jmbg=after_jmbg, mesto=after_mesto,
    ).to_dict()
    return AuditEvent(
        timestamp="2026-08-01T10:00:00+02:00", operator="test", event_type=event_type,
        slot_bm=bm, slot_function=function, before=before, after=after,
    )


def test_events_to_izmena_records_picks_up_plain_name_change() -> None:
    """Regression: Person.to_dict() has no "name" key (it's a @property,
    not a dataclass field) — a bare before.get("name")/after.get("name")
    always returned "", so events_to_izmena_records() silently skipped
    every real name change and auto-generation produced nothing, ever."""
    from biracki_odbor.core.izmene_docx import events_to_izmena_records

    ev = _person_event(
        1, "Члан 1",
        before_first="Никола", before_last="Ареснић",
        after_first="Славко", after_last="Максимовић",
    )
    by_bm = events_to_izmena_records([ev])
    assert 1 in by_bm
    assert len(by_bm[1]) == 1
    rec = by_bm[1][0]
    assert rec.old_name == "Никола Ареснић"
    assert rec.new_name == "Славко Максимовић"


def test_events_to_izmena_records_merges_clear_then_appoint() -> None:
    """The exact reported scenario: clear_member (name -> blank) followed
    by appoint_by_nacelnik (blank -> new name) for the SAME slot must
    collapse into ONE coherent "уместо X именује Y" record, not two
    incoherent half-records."""
    from biracki_odbor.core.izmene_docx import events_to_izmena_records

    clear_ev = _person_event(
        2, "Заменик члана 2",
        before_first="Драгољуб", before_last="Петковић",
        after_first="", after_last="",
        event_type="remove_member",
    )
    appoint_ev = _person_event(
        2, "Заменик члана 2",
        before_first="", before_last="",
        after_first="Ана", after_last="Анић",
        appointed_by_nacelnik=True,
        event_type="nacelnik_appointment",
    )
    by_bm = events_to_izmena_records([clear_ev, appoint_ev])
    assert 2 in by_bm
    assert len(by_bm[2]) == 1, (
        "Two audit events for one slot must merge into one record, not "
        "produce two garbled half-entries."
    )
    rec = by_bm[2][0]
    assert rec.old_name == "Драгољуб Петковић"
    assert rec.new_name == "Ана Анић"
    assert rec.appointed_by_nacelnik is True


def test_events_to_izmena_records_covers_every_changed_slot_on_a_bm() -> None:
    """A BM with two independently-changed slots must produce two
    records, both attributed to the same BM — the amendment decision
    for that BM must cover all of its replacements, not just one."""
    from biracki_odbor.core.izmene_docx import events_to_izmena_records

    ev1 = _person_event(
        3, "Председник",
        before_first="Пера", before_last="Перић",
        after_first="Мика", after_last="Микић",
    )
    ev2 = _person_event(
        3, "Члан 2",
        before_first="Жика", before_last="Жикић",
        after_first="Тома", after_last="Томић",
    )
    by_bm = events_to_izmena_records([ev1, ev2])
    assert len(by_bm[3]) == 2
    names = {(r.old_name, r.new_name) for r in by_bm[3]}
    assert ("Пера Перић", "Мика Микић") in names
    assert ("Жика Жикић", "Тома Томић") in names


# ── Same-person data corrections ────────────────────────────────────────────
#
# Field report: a residence-place correction (same person, JMBG unchanged)
# rendered as "уместо Стеван Стевић ... именује: Стеван Стевић из Мионица" —
# OIK appearing to un-appoint and re-appoint the same person to themselves.
# A correction must never use "уместо X именује Y" phrasing.


def test_events_to_izmena_records_flags_same_jmbg_as_correction() -> None:
    from biracki_odbor.core.izmene_docx import events_to_izmena_records

    ev = _person_event(
        1, "Председник",
        before_first="Стеван", before_last="Стевић",
        after_first="Стеван", after_last="Стевић",
        before_jmbg="0101990770015", after_jmbg="0101990770015",
        before_mesto="Уб", after_mesto="Мионица",
    )
    by_bm = events_to_izmena_records([ev])
    rec = by_bm[1][0]
    assert rec.is_correction is True
    assert rec.old_place == "Уб"
    assert rec.new_place == "Мионица"


def test_events_to_izmena_records_different_jmbg_is_not_a_correction() -> None:
    """A genuine replacement (different person) must keep the normal
    appointment phrasing even if the two names happen to collide."""
    from biracki_odbor.core.izmene_docx import events_to_izmena_records

    ev = _person_event(
        1, "Председник",
        before_first="Никола", before_last="Ареснић",
        after_first="Славко", after_last="Максимовић",
        before_jmbg="0101990770015", after_jmbg="0202980774807",
    )
    by_bm = events_to_izmena_records([ev])
    assert by_bm[1][0].is_correction is False


def test_events_to_izmena_records_skips_no_op_correction() -> None:
    """A correction event where neither name nor place actually differ in
    the document text (e.g. a phone/email-only edit) must not produce a
    false 'уместо X уписује се X' line."""
    from biracki_odbor.core.izmene_docx import events_to_izmena_records

    ev = _person_event(
        1, "Председник",
        before_first="Стеван", before_last="Стевић",
        after_first="Стеван", after_last="Стевић",
        before_jmbg="0101990770015", after_jmbg="0101990770015",
        before_mesto="Мионица", after_mesto="Мионица",
    )
    by_bm = events_to_izmena_records([ev])
    assert by_bm == {}


def _correction_record(**overrides) -> IzmenaRecord:
    base = dict(
        old_name="Стеван Стевић", function="ПРЕДСЕДНИКА", new_name="Стеван Стевић",
        old_place="Уб", new_place="Мионица", party="СНС", is_correction=True,
    )
    base.update(overrides)
    return IzmenaRecord(**base)


def _disposition_item_text(doc, index: int = 1) -> str:
    """The "{index}) ..." disposition paragraph — isolated from the rest
    of the document, since the Образложење boilerplate legitimately
    contains "именује" in an unrelated sentence about the Law."""
    prefix = f"{index}) "
    for p in doc.paragraphs:
        if p.text.startswith(prefix):
            return p.text
    raise AssertionError(f"no disposition paragraph starting with {prefix!r}")


def test_place_correction_does_not_say_imenuje() -> None:
    spec = _make_stalni_spec()
    spec.records = [_correction_record()]
    doc = build_izmena_document(spec, _make_oik(9), _make_election_with_auth())
    item = _disposition_item_text(doc)
    assert "именује" not in item
    assert "на предлог" not in item
    assert "исправља се податак о пребивалишту" in item
    assert "Уб" in item and "Мионица" in item


def test_name_correction_wording() -> None:
    spec = _make_stalni_spec()
    spec.records = [_correction_record(
        old_name="Стеван Стевич", new_name="Стеван Стевић",
        old_place="Мионица", new_place="Мионица",
    )]
    doc = build_izmena_document(spec, _make_oik(9), _make_election_with_auth())
    item = _disposition_item_text(doc)
    assert "именује" not in item
    assert "исправља се податак о имену и презимену" in item
    assert "Стеван Стевич" in item and "Стеван Стевић" in item


def test_name_and_place_correction_wording() -> None:
    spec = _make_stalni_spec()
    spec.records = [_correction_record(
        old_name="Стеван Стевич", new_name="Стеван Стевић",
    )]
    doc = build_izmena_document(spec, _make_oik(9), _make_election_with_auth())
    item = _disposition_item_text(doc)
    assert "именује" not in item
    assert "исправљају се подаци" in item
    assert "Уб" in item and "Мионица" in item


# ── Chairman-lapse track (чл. 41. ст. 2.) ───────────────────────────────────
#
# RIK's official forms ("new forms/1 3 ... zamena predsednika...") show a
# THIRD track, distinct from both the generic replacement and the
# same-person correction: председник/заменик председника replaced because
# they died, fell ill, or lost their voting right. Different legal basis
# citation, singular "ИЗМЕНИ" title, a reason-naming образложење, a longer
# deadline, and a one-sentence disposition with no numbered items.


def test_events_to_izmena_records_flags_chairman_death_as_lapse() -> None:
    from biracki_odbor.core.izmene_docx import events_to_izmena_records

    ev = _person_event(
        1, "Председник",
        before_first="Стеван", before_last="Стевић",
        after_first="Јован", after_last="Јовановић",
        before_jmbg="0101990770015", after_jmbg="0202980774807",
    )
    ev.note = "Смрт"
    by_bm = events_to_izmena_records([ev])
    rec = by_bm[1][0]
    assert rec.is_chairman_lapse is True
    assert rec.is_correction is False
    assert rec.function == "ПРЕДСЕДНИКА БИРАЧКОГ ОДБОРА"


def test_events_to_izmena_records_ordinary_member_death_stays_generic() -> None:
    """Чл. 41. ст. 2. only covers председник/заменик председника — an
    ordinary member who dies is still replaced under ст. 1 (at the
    proposer's request), the generic track."""
    from biracki_odbor.core.izmene_docx import events_to_izmena_records

    ev = _person_event(
        1, "Члан 1",
        before_first="Стеван", before_last="Стевић",
        after_first="Јован", after_last="Јовановић",
        before_jmbg="0101990770015", after_jmbg="0202980774807",
    )
    ev.note = "Смрт"
    by_bm = events_to_izmena_records([ev])
    assert by_bm[1][0].is_chairman_lapse is False


def test_events_to_izmena_records_chairman_replacement_without_reason_stays_generic() -> None:
    """An ordinary party-requested chairman swap (no Смрт/Неспособност/
    Губитак бирачког права reason recorded) is NOT the special track."""
    from biracki_odbor.core.izmene_docx import events_to_izmena_records

    ev = _person_event(
        1, "Председник",
        before_first="Стеван", before_last="Стевић",
        after_first="Јован", after_last="Јовановић",
        before_jmbg="0101990770015", after_jmbg="0202980774807",
    )
    by_bm = events_to_izmena_records([ev])
    assert by_bm[1][0].is_chairman_lapse is False


def _chairman_lapse_record(**overrides) -> IzmenaRecord:
    base = dict(
        old_name="Стеван Стевић", function="ПРЕДСЕДНИКА БИРАЧКОГ ОДБОРА",
        new_name="Јован Јовановић", new_place="Мионица", party="СНС",
        is_chairman_lapse=True,
    )
    base.update(overrides)
    return IzmenaRecord(**base)


def test_group_records_by_variant_splits_chairman_lapse_from_replacement() -> None:
    from biracki_odbor.core.izmene_docx import group_records_by_variant

    replacement = IzmenaRecord(old_name="А", function="ЧЛАНА 1", new_name="Б", new_place="Уб", party="СНС")
    lapse = _chairman_lapse_record()
    groups = group_records_by_variant([replacement, lapse])
    assert len(groups) == 2
    assert groups[0] == [replacement]
    assert groups[1] == [lapse]


def test_chairman_lapse_document_uses_singular_title() -> None:
    spec = _make_stalni_spec()
    spec.records = [_chairman_lapse_record()]
    doc = build_izmena_document(spec, _make_oik(9), _make_election_with_auth())
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "О ИЗМЕНИ РЕШЕЊА" in full_text
    assert "О ИЗМЕНАМА РЕШЕЊА" not in full_text


def test_chairman_lapse_document_legal_basis_cites_clause_41_stav_2() -> None:
    spec = _make_stalni_spec()
    spec.records = [_chairman_lapse_record()]
    # Even with an operator-edited legal_basis for the generic track, the
    # chairman-lapse document must ignore it and use its own citation.
    doc = build_izmena_document(
        spec, _make_oik(9), _make_election_with_auth(),
        legal_basis="На основу неке друге, погрешне одредбе,",
    )
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "члана 41. став 2." in full_text
    assert "погрешне одредбе" not in full_text


def test_chairman_lapse_disposition_is_one_sentence_no_numbered_items() -> None:
    spec = _make_stalni_spec()
    spec.records = [_chairman_lapse_record()]
    doc = build_izmena_document(spec, _make_oik(9), _make_election_with_auth())
    item = next(p.text for p in doc.paragraphs if p.text.startswith("1. Мења се Решење"))
    assert "1)" not in item
    assert "тако што се у бирачки одбор:" not in item
    assert "уместо" in item and "именује" in item
    assert "Стеван Стевић" in item and "Јован Јовановић" in item
    # No other paragraph introduces a numbered sub-item either.
    assert not any(p.text.startswith("1) ") for p in doc.paragraphs)


def test_chairman_lapse_obrazlozenje_names_the_grounds() -> None:
    spec = _make_stalni_spec()
    spec.records = [_chairman_lapse_record()]
    # Even with an operator-customized образложење for the generic track,
    # the chairman-lapse document must ignore it.
    doc = build_izmena_document(
        spec, _make_oik(9), _make_election_with_auth(),
        rendered_obrazlozenje=[[("Нека сасвим друга образложење реченица.", False)]],
    )
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "ако умру, разболе се или ако изгубе изборно право" in full_text
    assert "Нека сасвим друга" not in full_text


def test_chairman_lapse_remedy_includes_poslanicka_grupa() -> None:
    """RIK corrected the official "1 3" form (2026-10-09 "по исправци"
    revision) to match the generic "1 2" template's remedy clause,
    adding "посланичка група у Народној скупштини" — it used to be the
    one thing that differed between the two tracks' remedy text."""
    spec = _make_stalni_spec()
    spec.records = [_chairman_lapse_record()]
    doc = build_izmena_document(spec, _make_oik(9), _make_election_with_auth())
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "посланичка група у Народној скупштини" in full_text


def test_no_sostava_typo_in_chairman_lapse_document() -> None:
    """Field report: "промени состава бирачким одбором" is a Russian/
    Macedonian spelling that slipped into the authorization clause —
    Serbian is "састава". Covers both tracks since both render it (via
    the shared "ovlascenje" paragraph) whenever auth_ref is set."""
    spec = _make_stalni_spec()
    spec.records = [_chairman_lapse_record()]
    doc = build_izmena_document(spec, _make_oik(9), _make_election_with_auth())
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "состава" not in full_text
    assert "састава" in full_text


def test_no_sostava_typo_in_generic_document() -> None:
    doc = build_izmena_document(_make_stalni_spec(), _make_oik(9), _make_election_with_auth())
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "состава" not in full_text
    assert "састава" in full_text


# ── Bundled JSON seed resources (replacing in-code Python literals) ────────
#
# The обрадложење/legal-basis text is now shipped as JSON under
# biracki_odbor/resources/ (not Python literals) so a wording change
# between election cycles is a data edit, not a code release — same
# pattern already used for config/opstine.py's municipality registry.


def test_seed_izmena_paragraphs_reads_bundled_resource() -> None:
    from biracki_odbor.core.izmene_docx import seed_izmena_paragraphs

    stalni = seed_izmena_paragraphs("stalni")
    prosireni = seed_izmena_paragraphs("prosireni")
    assert len(stalni) > 0 and len(prosireni) > 0
    assert stalni != prosireni


def test_chairman_lapse_defaults_resource_has_no_typo() -> None:
    from biracki_odbor.core.izmene_docx import _chairman_lapse_defaults

    legal_basis, paragraphs = _chairman_lapse_defaults()
    assert "состава" not in legal_basis
    for p in paragraphs:
        for seg in p.segments:
            assert "состава" not in seg.text


def test_default_paragraphs_reads_bundled_resource() -> None:
    from biracki_odbor.config.obrazlozenje import default_paragraphs
    from biracki_odbor import app_paths

    stalni = default_paragraphs()
    prosireni = default_paragraphs(app_paths.obrazlozenje_path_for(
        type("S", (), {"obrazlozenje_filename": "obrazlozenje_prosireni.json"})()
    ))
    assert len(stalni) > 0 and len(prosireni) > 0
    assert [p.id for p in stalni] != [p.id for p in prosireni]


# ── gui/tabs/izmene.py: _evs_to_records (manual Измене tab flow) ───────────
# Same merge fix, applied to the flattened-dict shape this function
# consumes (from load_all_post_rik_changes, already correct on names
# since that function already had the first_name/last_name fallback —
# only the merge-by-function behavior was missing here).


def test_gui_evs_to_records_merges_clear_then_appoint() -> None:
    from biracki_odbor.gui.tabs.izmene import _evs_to_records

    evs = [
        {
            "bm": 2, "function": "Заменик члана 2",
            "old_name": "Драгољуб Петковић", "new_name": "",
            "mesto": "", "party": "СНС", "appointed_by_nacelnik": False,
        },
        {
            "bm": 2, "function": "Заменик члана 2",
            "old_name": "", "new_name": "Ана Анић",
            "mesto": "Уб", "party": "СНС", "appointed_by_nacelnik": True,
        },
    ]
    records = _evs_to_records(evs)
    assert len(records) == 1
    rec = records[0]
    assert rec.old_name == "Драгољуб Петковић"
    assert rec.new_name == "Ана Анић"
    assert rec.appointed_by_nacelnik is True


def test_gui_evs_to_records_covers_every_changed_function() -> None:
    from biracki_odbor.gui.tabs.izmene import _evs_to_records

    evs = [
        {"bm": 3, "function": "Председник", "old_name": "Пера Перић",
         "new_name": "Мика Микић", "mesto": "", "party": "СНС",
         "appointed_by_nacelnik": False},
        {"bm": 3, "function": "Члан 2", "old_name": "Жика Жикић",
         "new_name": "Тома Томић", "mesto": "", "party": "СНС",
         "appointed_by_nacelnik": False},
    ]
    records = _evs_to_records(evs)
    assert len(records) == 2


def test_gui_evs_to_records_flags_place_correction() -> None:
    """Same bug, GUI-tab flow: load_all_post_rik_changes() carries
    before_jmbg/after_jmbg/old_mesto per event, so this path must flag
    it as a correction too, not just the corrections.py auto-generate
    path (events_to_izmena_records)."""
    from biracki_odbor.gui.tabs.izmene import _evs_to_records

    evs = [{
        "bm": 1, "function": "Председник",
        "old_name": "Стеван Стевић", "new_name": "Стеван Стевић",
        "old_mesto": "Уб", "mesto": "Мионица", "party": "СНС",
        "before_jmbg": "0101990770015", "after_jmbg": "0101990770015",
        "appointed_by_nacelnik": False,
    }]
    records = _evs_to_records(evs)
    assert len(records) == 1
    rec = records[0]
    assert rec.is_correction is True
    assert rec.old_place == "Уб"
    assert rec.new_place == "Мионица"


def test_gui_evs_to_records_skips_no_op_correction() -> None:
    from biracki_odbor.gui.tabs.izmene import _evs_to_records

    evs = [{
        "bm": 1, "function": "Председник",
        "old_name": "Стеван Стевић", "new_name": "Стеван Стевић",
        "old_mesto": "Мионица", "mesto": "Мионица", "party": "СНС",
        "before_jmbg": "0101990770015", "after_jmbg": "0101990770015",
        "appointed_by_nacelnik": False,
    }]
    assert _evs_to_records(evs) == []


def test_gui_item_html_correction_does_not_say_imenuje() -> None:
    from biracki_odbor.core.izmene_docx import IzmenaRecord
    from biracki_odbor.gui.tabs.izmene import _izmena_item_html

    rec = IzmenaRecord(
        old_name="Стеван Стевић", function="ПРЕДСЕДНИКА", new_name="Стеван Стевић",
        old_place="Уб", new_place="Мионица", party="СНС", is_correction=True,
    )
    html = _izmena_item_html(1, rec)
    assert "именује" not in html
    assert "Уб" in html and "Мионица" in html


def test_gui_evs_to_records_flags_chairman_death_as_lapse() -> None:
    from biracki_odbor.gui.tabs.izmene import _evs_to_records

    evs = [{
        "bm": 1, "function": "Председник",
        "old_name": "Стеван Стевић", "new_name": "Јован Јовановић",
        "old_mesto": "Мионица", "mesto": "Мионица", "party": "СНС",
        "before_jmbg": "0101990770015", "after_jmbg": "0202980774807",
        "note": "Смрт", "appointed_by_nacelnik": False,
    }]
    records = _evs_to_records(evs)
    assert len(records) == 1
    rec = records[0]
    assert rec.is_chairman_lapse is True
    assert rec.function == "ПРЕДСЕДНИКА БИРАЧКОГ ОДБОРА"


def test_gui_evs_to_records_ordinary_member_death_stays_generic() -> None:
    from biracki_odbor.gui.tabs.izmene import _evs_to_records

    evs = [{
        "bm": 1, "function": "Члан 1",
        "old_name": "Стеван Стевић", "new_name": "Јован Јовановић",
        "old_mesto": "Мионица", "mesto": "Мионица", "party": "СНС",
        "before_jmbg": "0101990770015", "after_jmbg": "0202980774807",
        "note": "Смрт", "appointed_by_nacelnik": False,
    }]
    records = _evs_to_records(evs)
    assert records[0].is_chairman_lapse is False


def test_gui_item_html_chairman_lapse_flags_separate_document() -> None:
    """The merged preview can't show the chairman-lapse document's real
    shape (different title/legal basis/образложење) — it must say so
    rather than render a plain numbered item that implies it belongs to
    the same document as everything else in the batch."""
    from biracki_odbor.core.izmene_docx import IzmenaRecord
    from biracki_odbor.gui.tabs.izmene import _izmena_item_html

    rec = IzmenaRecord(
        old_name="Стеван Стевић", function="ПРЕДСЕДНИКА БИРАЧКОГ ОДБОРА",
        new_name="Јован Јовановић", new_place="Мионица", party="СНС",
        is_chairman_lapse=True,
    )
    html = _izmena_item_html(1, rec)
    assert "засебан документ" in html
    assert "Стеван Стевић" in html and "Јован Јовановић" in html


def test_group_records_by_variant_groups_mixed_bm_records_from_gui_flow() -> None:
    from biracki_odbor.core.izmene_docx import group_records_by_variant
    from biracki_odbor.gui.tabs.izmene import _evs_to_records

    evs = [
        {"bm": 1, "function": "Председник",
         "old_name": "Стеван Стевић", "new_name": "Јован Јовановић",
         "old_mesto": "Мионица", "mesto": "Мионица", "party": "СНС",
         "before_jmbg": "0101990770015", "after_jmbg": "0202980774807",
         "note": "Смрт", "appointed_by_nacelnik": False},
        {"bm": 1, "function": "Члан 1",
         "old_name": "Пера Перић", "new_name": "Мика Микић",
         "old_mesto": "", "mesto": "Уб", "party": "СНС",
         "before_jmbg": "", "after_jmbg": "",
         "appointed_by_nacelnik": False},
    ]
    records = _evs_to_records(evs)
    groups = group_records_by_variant(records)
    assert len(groups) == 2
    variants = {"chairman_lapse" if g[0].is_chairman_lapse else "replacement" for g in groups}
    assert variants == {"chairman_lapse", "replacement"}


# ── Step 5 regrouped by legal track (replacement / correction /          ────
#    chairman-lapse), each sub-grouped by BM, plus чл. 15 Престанак folded
#    in as its own always-available section — replacing the old flat,
#    mixed-variant changes table and the standalone прошиени-only step 6.


def test_group_all_records_by_variant_splits_by_bm_and_track() -> None:
    from biracki_odbor.gui.tabs.izmene import _group_all_records_by_variant

    changes_data = [
        {"bm": 3, "function": "Председник", "old_name": "Пера Перић",
         "new_name": "Мика Микић", "old_mesto": "", "mesto": "",
         "before_jmbg": "0101990770015", "after_jmbg": "0202980774807",
         "party": "СНС", "appointed_by_nacelnik": False},
        {"bm": 1, "function": "Председник", "old_name": "Стеван Стевић",
         "new_name": "Стеван Стевић", "old_mesto": "Уб", "mesto": "Мионица",
         "before_jmbg": "0303970770013", "after_jmbg": "0303970770013",
         "party": "СНС", "appointed_by_nacelnik": False},
    ]
    grouped = _group_all_records_by_variant(changes_data)
    assert len(grouped["replacement"]) == 1
    assert grouped["replacement"][0][0] == 3
    assert len(grouped["correction"]) == 1
    assert grouped["correction"][0][0] == 1
    assert grouped["chairman_lapse"] == []


def test_group_prestanak_drafts_picks_up_cl_14_1_8_flagged_slots() -> None:
    from biracki_odbor.gui.tabs.izmene import _group_prestanak_drafts

    changes_data = [
        {"bm": 1, "function": "Члан 2", "old_name": "Ана Анић", "new_name": "Ана Анић",
         "note": "по члану 14. став 1. тачка 8) — изабран за члана бирачког"
                 " одбора на бирачком месту број 5 у општини Ваљево",
         "other_bm": 5, "other_municipality": "Ваљево"},
        {"bm": 3, "function": "Председник", "old_name": "Пера Перић",
         "new_name": "Мика Микић", "note": "Смрт", "other_bm": None, "other_municipality": ""},
    ]
    drafts = _group_prestanak_drafts(changes_data)
    assert len(drafts) == 1
    assert drafts[0]["bm"] == 1
    assert drafts[0]["person_name"] == "Ана Анић"
    assert drafts[0]["other_bm"] == 5
    assert drafts[0]["other_municipality"] == "Ваљево"


def test_group_prestanak_drafts_carries_other_decision_broj() -> None:
    from biracki_odbor.gui.tabs.izmene import _group_prestanak_drafts

    changes_data = [
        {"bm": 1, "function": "Члан 2", "old_name": "Ана Анић", "new_name": "Ана Анић",
         "note": "по члану 14. став 1. тачка 8) — изабран за члана бирачког"
                 " одбора на бирачком месту број 5 у општини Ваљево",
         "other_bm": 5, "other_municipality": "Ваљево",
         "other_decision_broj": "013-9-1-4/2026"},
    ]
    drafts = _group_prestanak_drafts(changes_data)
    assert drafts[0]["other_decision_broj"] == "013-9-1-4/2026"


def test_group_prestanak_drafts_keeps_flag_through_a_later_unrelated_edit() -> None:
    """The flag must survive a LATER, unrelated edit on the same slot
    (e.g. a name typo fix) that doesn't re-select the Напомена reason —
    checking only the slot's latest event would silently "unflag" it.
    Person name still comes from the latest (corrected) event; the
    other BM/municipality come from the event that actually flagged it."""
    from biracki_odbor.gui.tabs.izmene import _group_prestanak_drafts

    changes_data = [
        {"bm": 1, "function": "Члан 2", "old_name": "Aна Анић", "new_name": "Ана Анић",
         "note": "по члану 14. став 1. тачка 8) — изабран за члана бирачког"
                 " одбора на бирачком месту број 5 у општини Ваљево",
         "other_bm": 5, "other_municipality": "Ваљево"},
        {"bm": 1, "function": "Члан 2", "old_name": "Ана Анић", "new_name": "Ана Анић Поповић",
         "note": "", "other_bm": None, "other_municipality": ""},
    ]
    drafts = _group_prestanak_drafts(changes_data)
    assert len(drafts) == 1
    assert drafts[0]["person_name"] == "Ана Анић Поповић"
    assert drafts[0]["other_bm"] == 5
    assert drafts[0]["other_municipality"] == "Ваљево"


def test_izmena_panel_has_one_table_per_track_plus_prestanak(app, tmp_path) -> None:
    from biracki_odbor.core.workflow import STALNI_SPEC
    from biracki_odbor.gui.tabs.izmene import IzmenaPanel

    state = _make_state(tmp_path)
    panel = IzmenaPanel(STALNI_SPEC, state, lambda jmbg: None)
    try:
        assert hasattr(panel, "table_replacement")
        assert hasattr(panel, "table_correction")
        assert hasattr(panel, "table_chairman_lapse")
        # changes_table (the old single flat table) must be gone, not
        # just unused — its continued presence would signal the old and
        # new selection models are both half-wired.
        assert not hasattr(panel, "changes_table")
    finally:
        panel.deleteLater()


def test_empty_track_sections_are_hidden_until_something_is_pending(
    app, home, tmp_path, monkeypatch,
) -> None:
    """Four always-visible accordions (three of them routinely empty)
    was pure clutter — a track's section must hide entirely when
    nothing's pending for it, and reappear (with its count) once a real
    change lands in the log. чл. 15 Престанак is log-driven exactly like
    the other three (Напомена = "по члану 14. став 1. тачка 8)"), so it
    hides on the same empty-by-default basis."""
    from biracki_odbor.core.corrections import AuditEvent
    from biracki_odbor.core.models import Person, PollingStation
    from biracki_odbor.core.workflow import STALNI_SPEC
    from biracki_odbor.gui.tabs.izmene import IzmenaPanel

    state = _make_state(tmp_path)
    state.oik.polling_stations = [PollingStation(bm_number=1, name="БМ 1")]
    wl = state.generation_log.workflow("stalni")
    wl.submit_to_rik()
    state.save_generation_log()

    panel = IzmenaPanel(STALNI_SPEC, state, lambda jmbg: None)
    try:
        # isVisible() needs a real top-level show(); isVisibleTo(panel) checks
        # widget-tree-relative visibility without one (established pattern —
        # see test_person_edit_note_reasons.py).
        assert not panel._section_replacement.isVisibleTo(panel)  # noqa: SLF001
        assert not panel._section_correction.isVisibleTo(panel)  # noqa: SLF001
        assert not panel._section_chairman_lapse.isVisibleTo(panel)  # noqa: SLF001
        assert not panel._section_prestanak.isVisibleTo(panel)  # noqa: SLF001

        log_path = tmp_path / "corrections.log"
        before = Person(bm_number=1, function="Члан 2", party="СНС",
                         first_name="Радомир", last_name="Радић", jmbg="1605991774818")
        after = Person(bm_number=1, function="Члан 2", party="СНС",
                        first_name="Марко", last_name="Јовичић", jmbg="2706994779807")
        from datetime import datetime, timedelta
        event_ts = (datetime.now().astimezone() + timedelta(minutes=1)).isoformat()
        event = AuditEvent(
            timestamp=event_ts, operator="test", event_type="nacelnik_appointment",
            slot_bm=1, slot_function="Члан 2",
            before=before.to_dict(), after=after.to_dict(),
        )
        log_path.write_text(event.to_jsonl() + "\n", encoding="utf-8")
        import biracki_odbor.app_paths as app_paths_mod
        monkeypatch.setattr(app_paths_mod, "corrections_log_path_for", lambda spec: log_path)

        panel._refresh_changes_table()  # noqa: SLF001

        assert panel._section_replacement.isVisibleTo(panel)  # noqa: SLF001
        assert not panel._section_correction.isVisibleTo(panel)  # noqa: SLF001
        assert not panel._section_chairman_lapse.isVisibleTo(panel)  # noqa: SLF001
        assert not panel._section_prestanak.isVisibleTo(panel)  # noqa: SLF001
    finally:
        panel.deleteLater()


def test_izmena_panel_embeds_prestanak_panel_for_stalni_too(app, tmp_path) -> None:
    """чл. 15 cessation must be reachable from step 5 in BOTH workflows
    now — it used to be a прошиени-only standalone step 6."""
    from biracki_odbor.core.workflow import STALNI_SPEC
    from biracki_odbor.gui.tabs.izmene import IzmenaPanel
    from biracki_odbor.gui.tabs.prestanak_panel import PrestanakPanel

    state = _make_state(tmp_path)
    panel = IzmenaPanel(STALNI_SPEC, state, lambda jmbg: None)
    try:
        assert panel.findChild(PrestanakPanel) is not None
    finally:
        panel.deleteLater()


def test_selecting_a_prestanak_row_switches_mode_and_prefills(
    app, home, tmp_path, monkeypatch,
) -> None:
    """End-to-end: a чл. 14(1)(8)-flagged slot appears in the Престанак
    table; selecting it switches the shared left column to
    PrestanakPanel (hiding the izmena editor), pre-filled from the log,
    and the shared Генериши button then calls generate_prestanak() —
    producing a real решење о престанку, not an izmena document."""
    from biracki_odbor.core.corrections import AuditEvent
    from biracki_odbor.core.models import Person, PollingStation
    from biracki_odbor.core.workflow import STALNI_SPEC
    from biracki_odbor.gui.tabs.izmene import IzmenaPanel

    state = _make_state(tmp_path)
    state.oik.municipality.name = "Мионица"
    state.oik.polling_stations = [PollingStation(bm_number=1, name="БМ 1")]
    wl = state.generation_log.workflow("stalni")
    wl.submit_to_rik()
    state.save_generation_log()

    log_path = tmp_path / "corrections.log"
    person = Person(bm_number=1, function="Члан 2", party="СНС",
                     first_name="Ана", last_name="Анић", jmbg="1605991774818")
    from datetime import datetime, timedelta
    event_ts = (datetime.now().astimezone() + timedelta(minutes=1)).isoformat()
    event = AuditEvent(
        timestamp=event_ts, operator="test", event_type="edit_member",
        slot_bm=1, slot_function="Члан 2",
        before=person.to_dict(), after=person.to_dict(),
        note="по члану 14. став 1. тачка 8)", other_bm=5, other_municipality="Ваљево",
    )
    log_path.write_text(event.to_jsonl() + "\n", encoding="utf-8")
    import biracki_odbor.app_paths as app_paths_mod
    monkeypatch.setattr(app_paths_mod, "corrections_log_path_for", lambda spec: log_path)

    panel = IzmenaPanel(STALNI_SPEC, state, lambda jmbg: None)
    try:
        panel._refresh_changes_table()  # noqa: SLF001
        assert panel.table_prestanak.rowCount() == 1

        panel.table_prestanak.selectRow(0)

        assert panel._mode == "prestanak"  # noqa: SLF001
        assert not panel._izmena_editor.isVisibleTo(panel)  # noqa: SLF001
        assert panel._prestanak_editor.isVisibleTo(panel)  # noqa: SLF001
        assert panel._prestanak_editor.ed_person_name.text() == "Ана Анић"  # noqa: SLF001
        assert panel._prestanak_editor.sb_existing_bm.value() == 5  # noqa: SLF001
        assert panel._prestanak_editor.ed_existing_municipality.text() == "Ваљево"  # noqa: SLF001

        panel._prestanak_editor.ed_ceasing_broj.setText("013-5-30-9/2026")  # noqa: SLF001
        panel._prestanak_editor.ed_existing_broj.setText("013-5-30-1/2026")  # noqa: SLF001
        panel._prestanak_editor.ed_document_broj.setText("Број: 013-5-30-15/2026")  # noqa: SLF001

        panel._on_generate()  # noqa: SLF001

        from biracki_odbor import app_paths
        prestanak_files = list(app_paths.prestanak_dir().glob("*.docx"))
        assert len(prestanak_files) == 1
        izmena_files = list(app_paths.izmene_dir().glob("*.docx"))
        assert not izmena_files
    finally:
        panel.deleteLater()


def test_prestanak_preview_refreshes_as_fields_are_typed(app, tmp_path) -> None:
    """Typing into any PrestanakPanel field (broj, date, name, …) must
    update the shared preview immediately — it used to only refresh on
    row selection, so the preview stayed frozen on the "Fill in: …"
    placeholder even after every required field had been filled in."""
    from biracki_odbor.core.workflow import STALNI_SPEC
    from biracki_odbor.gui.tabs.izmene import IzmenaPanel

    state = _make_state(tmp_path)
    state.oik.municipality.name = "Мионица"

    panel = IzmenaPanel(STALNI_SPEC, state, lambda jmbg: None)
    try:
        panel._set_mode("prestanak")  # noqa: SLF001
        editor = panel._prestanak_editor  # noqa: SLF001

        editor.ed_person_name.setText("Ана Анић")
        editor.ed_ceasing_broj.setText("013-5-30-9/2026")
        editor.ed_existing_broj.setText("013-5-30-1/2026")
        editor.ed_document_broj.setText("Број: 013-5-30-15/2026")

        text = panel._preview.to_plain_text()  # noqa: SLF001
        assert "Ана Анић" in text
        assert "Fill in" not in text
    finally:
        panel.deleteLater()


def test_selecting_a_row_in_one_table_clears_the_others(app, tmp_path) -> None:
    """Exactly one group can be selected at a time across the three
    tracks — otherwise _selected_records would be ambiguous."""
    from biracki_odbor.core.izmene_docx import IzmenaRecord
    from biracki_odbor.core.workflow import STALNI_SPEC
    from biracki_odbor.gui.tabs.izmene import IzmenaPanel
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QTableWidgetItem

    state = _make_state(tmp_path)
    panel = IzmenaPanel(STALNI_SPEC, state, lambda jmbg: None)
    try:
        for variant, table in (
            ("replacement", panel.table_replacement),
            ("correction", panel.table_correction),
        ):
            table.setRowCount(1)
            item = QTableWidgetItem("1")
            item.setData(Qt.UserRole, (1, [IzmenaRecord()], ""))
            table.setItem(0, 0, item)

        panel.table_replacement.selectRow(0)
        assert panel._selected_variant == "replacement"  # noqa: SLF001

        panel.table_correction.selectRow(0)
        assert panel._selected_variant == "correction"  # noqa: SLF001
        assert not panel.table_replacement.selectionModel().hasSelection()
    finally:
        panel.deleteLater()


# ── Овлашћење section: CollapsibleSection + ToggleSwitch ────────────────────
# Replaced a plain QRadioButton pair with the established collapsible-box +
# red/blue-switch pattern (same widgets as the Начелник/RIK sections).


def test_authorization_toggle_defaults_on_and_shows_fields(app, home, tmp_path) -> None:
    from biracki_odbor.core.workflow import STALNI_SPEC
    from biracki_odbor.gui.tabs.izmene import IzmenaPanel

    state = _make_state(tmp_path)
    panel = IzmenaPanel(STALNI_SPEC, state, lambda jmbg: None)
    try:
        assert panel.sw_delegated.isChecked() is True
        assert panel._auth_frame.isVisibleTo(panel)  # noqa: SLF001
    finally:
        panel.deleteLater()


def test_authorization_toggle_off_hides_fields_and_saves(app, home, tmp_path) -> None:
    from biracki_odbor.core.workflow import STALNI_SPEC
    from biracki_odbor.gui.tabs.izmene import IzmenaPanel

    state = _make_state(tmp_path)
    panel = IzmenaPanel(STALNI_SPEC, state, lambda jmbg: None)
    try:
        panel.sw_delegated.setChecked(False)
        assert not panel._auth_frame.isVisibleTo(panel)  # noqa: SLF001
        assert state.election.document.use_authorization is False
    finally:
        panel.deleteLater()


# ── Authorization date — date picker, not free text ─────────────────────────
#
# Field report: "Датум овлашћивања" was a plain QLineEdit, so the operator
# could type anything ("111") and it would be saved verbatim and embedded
# straight into the decision document's legal prose. Replaced with the same
# QDateEdit + calendar-popup widget already used for the amendment date
# (dt_issue_date) and the election date elsewhere in the app — the stored
# value is still the genitive date PHRASE core.izmene_docx expects
# ("6. децембра 2023. године"), just produced/read from a real QDate
# instead of free-typed text.


def test_auth_date_is_a_real_date_picker(app, home, tmp_path) -> None:
    from PySide6.QtWidgets import QDateEdit
    from biracki_odbor.core.workflow import STALNI_SPEC
    from biracki_odbor.gui.tabs.izmene import IzmenaPanel

    state = _make_state(tmp_path)
    panel = IzmenaPanel(STALNI_SPEC, state, lambda jmbg: None)
    try:
        assert isinstance(panel.ed_auth_date, QDateEdit)
    finally:
        panel.deleteLater()


def test_auth_date_save_produces_genitive_phrase(app, home, tmp_path) -> None:
    """Picking a date and saving must store the same "D. month-genitive
    YYYY. године" phrase the free-text field used to require the
    operator to type by hand — now derived, never mistyped."""
    from PySide6.QtCore import QDate
    from biracki_odbor.core.workflow import STALNI_SPEC
    from biracki_odbor.gui.tabs.izmene import IzmenaPanel

    state = _make_state(tmp_path)
    panel = IzmenaPanel(STALNI_SPEC, state, lambda jmbg: None)
    try:
        panel.ed_auth_date.setDate(QDate(2023, 12, 6))
        assert state.election.document.authorization_date == "6. децембра 2023. године"
    finally:
        panel.deleteLater()


def test_auth_date_load_parses_stored_genitive_phrase(app, home, tmp_path) -> None:
    """Reverse of the save test: an existing saved config's genitive
    phrase must populate the picker with the matching real date, not
    leave it stuck on today's date."""
    from PySide6.QtCore import QDate
    from biracki_odbor.core.workflow import STALNI_SPEC
    from biracki_odbor.gui.tabs.izmene import IzmenaPanel

    state = _make_state(tmp_path)
    state.election.document.authorization_date = "6. децембра 2023. године"
    panel = IzmenaPanel(STALNI_SPEC, state, lambda jmbg: None)
    try:
        panel._load_auth_fields()  # noqa: SLF001
        assert panel.ed_auth_date.date() == QDate(2023, 12, 6)
    finally:
        panel.deleteLater()


# ── Manual "Generate" button — izmene_per_bm persistence ───────────────────
#
# Field report: "Генерисано" stayed "—" forever, and every generate click
# for the same BM overwrote the same "..._izmena_01.docx" instead of
# advancing to _02, _03, ... Root cause: _on_generate() called
# `self.state.generation_log.record_izmena(bm_num)` — record_izmena is a
# WorkflowLog method, not a GenerationLog one (GenerationLog only exposes
# .workflow(slug) -> WorkflowLog). This always raised AttributeError,
# silently swallowed by a bare `except Exception: pass`, so the counter
# never advanced and never reached disk — even though the .docx itself
# was written successfully by a separate, unaffected call just above it.


def test_manual_generate_persists_izmene_per_bm(app, home, tmp_path, monkeypatch) -> None:
    from biracki_odbor.core.corrections import AuditEvent
    from biracki_odbor.core.models import Person, PollingStation
    from biracki_odbor.core.workflow import STALNI_SPEC
    from biracki_odbor.gui.tabs.izmene import IzmenaPanel

    state = _make_state(tmp_path)
    state.oik.municipality.name = "Уб"
    state.oik.polling_stations = [PollingStation(bm_number=1, name="БМ 1")]
    state.election.election.dates.date_full = "17. децембар 2026. године"
    state.election.election.dates.session_date = "1. јануара 2026. године"

    wl = state.generation_log.workflow("stalni")
    state.generation_log.record_phase("stalni", "phase1")
    state.generation_log.record_phase("stalni", "phase2")
    state.generation_log.record_phase("stalni", "phase3")
    wl.submit_to_rik()
    state.save_generation_log()

    log_path = tmp_path / "corrections.log"
    before = Person(bm_number=1, function="Члан 2", party="СНС",
                     first_name="Радомир", last_name="Радић", jmbg="1605991774818").to_dict()
    after = Person(bm_number=1, function="Члан 2", party="СНС",
                    first_name="Марко", last_name="Јовичић", jmbg="2706994779807",
                    mesto="Врело", appointed_by_nacelnik=True).to_dict()
    # Must postdate wl.submit_to_rik() above (real wall-clock time) — the
    # panel filters to post-RIK changes only. A fixed clock-time string
    # here would be a time bomb that only passes before that hour.
    from datetime import datetime, timedelta
    event_ts = (datetime.now().astimezone() + timedelta(minutes=1)).isoformat()
    event = AuditEvent(
        timestamp=event_ts, operator="test",
        event_type="nacelnik_appointment", slot_bm=1, slot_function="Члан 2",
        before=before, after=after,
    )
    log_path.write_text(event.to_jsonl() + "\n", encoding="utf-8")
    import biracki_odbor.app_paths as app_paths_mod
    monkeypatch.setattr(app_paths_mod, "corrections_log_path_for", lambda spec: log_path)

    panel = IzmenaPanel(STALNI_SPEC, state, lambda jmbg: None)
    try:
        panel.sw_delegated.setChecked(False)  # skip authorization fields
        panel._refresh_changes_table()  # noqa: SLF001
        assert panel._changes_data, "the seeded event must appear in the changes table"

        # Different JMBG before/after on an ordinary member -> the
        # "replacement" track's own grouped table/section.
        panel.table_replacement.selectRow(0)
        panel._on_generate()  # noqa: SLF001

        assert wl.izmene_per_bm.get("1") == 1, (
            "record_izmena must be called on the WorkflowLog"
            " (wl), not on GenerationLog directly."
        )

        # Second generate for the same BM must advance to amendment 2,
        # not silently reuse amendment 1 forever.
        panel._refresh_bm_info()  # noqa: SLF001
        assert panel._derived_amendment_idx == 2  # noqa: SLF001
    finally:
        panel.deleteLater()
