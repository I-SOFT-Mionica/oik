"""End-to-end verification of a complete simulated election cycle.

Walks every backend module through one realistic scenario:
  Phase 1 → fake party returns → Phase 2 → injected JMBG conflict →
  Phase 2.5 resolution + manual replacement → Phase 3 docx generation.

Also asserts the audit log captured every state change.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from docx import Document

from biracki_odbor.config.election import ElectionConfig, ElectionInfo, ElectionDates, DocumentInfo
from biracki_odbor.config.oik import OikConstants
from biracki_odbor.core.corrections import CorrectionsState, detect_conflicts, read_log
from biracki_odbor.core.master_io import read_master, write_master
from biracki_odbor.core.phase1_generate import generate_party_files
from biracki_odbor.core.phase2_aggregate import aggregate_party_files
from biracki_odbor.core.phase3_docx import generate_documents
from biracki_odbor.validation.reports import Codes
from tests.fixtures.stalni_sastav import build_template
from tests.test_jmbg import VALID_SAMPLES, _build_valid


def _unique_jmbgs(n: int) -> list[str]:
    """Generate N distinct valid JMBGs for tests."""
    out: list[str] = []
    i = 0
    while len(out) < n:
        candidate = _build_valid(f"{i:012d}")
        if candidate not in out:
            out.append(candidate)
        i += 1
    return out
from tests.test_phase1 import _make_oik
from tests.test_phase3 import _make_election


def _fill_party_file(file_path: Path, *, jmbg_per_row: list[str]) -> None:
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    for i, jmbg in enumerate(jmbg_per_row):
        row = i + 4
        # Schema post-split: E=Име, F=Презиме, G=ЈМБГ, H=Место, I=Адреса, J=Тел, K=Е-пошта
        ws.cell(row=row, column=5, value=f"Особа{row}")
        ws.cell(row=row, column=6, value=f"Презименић{row}")
        ws.cell(row=row, column=7, value=jmbg)
        ws.cell(row=row, column=8, value="Уб")
        ws.cell(row=row, column=9, value=f"Адреса {row}")
        ws.cell(row=row, column=10, value="0641234567")
        ws.cell(row=row, column=11, value=f"o{row}@ub.rs")
    wb.save(file_path)
    wb.close()


def test_full_cycle_with_conflict_and_replacement(tmp_path: Path) -> None:
    oik = _make_oik(n_stations=2)
    election = _make_election()
    template = build_template(tmp_path / "Stalni.xlsx",
                              n_polling_stations=2, parties=("СНС", "СПС"))

    # ── Phase 1 ────────────────────────────────────────────
    stranke_dir = tmp_path / "stranka_fajlovi"
    p1 = generate_party_files(template, stranke_dir, oik)
    assert p1.report.ok, [i.message_sr for i in p1.report.issues]
    assert len(p1.files_written) == 2

    # ── Party returns: inject ONE duplicate JMBG across the two parties ──
    # Use enough unique JMBGs that the only conflict is the deliberate one.
    same_jmbg = _build_valid("999988887777")
    pool = _unique_jmbgs(60)  # plenty
    pool_iter = iter(pool)
    for f in p1.files_written:
        wb = openpyxl.load_workbook(f)
        ws = wb.active
        n_data_rows = 0
        for r in range(4, ws.max_row + 1):
            if ws.cell(row=r, column=2).value is None:
                break
            n_data_rows += 1
        wb.close()
        jmbgs = [same_jmbg] + [next(pool_iter) for _ in range(n_data_rows - 1)]
        _fill_party_file(f, jmbg_per_row=jmbgs)

    popunjeni = tmp_path / "popunjeni"
    popunjeni.mkdir()
    for f in p1.files_written:
        f.rename(popunjeni / f.name)

    # ── Phase 2 ────────────────────────────────────────────
    master_path = tmp_path / "master" / "zbirana.xlsx"
    p2 = aggregate_party_files(popunjeni, template, master_path, oik)
    assert master_path.exists()
    codes = {i.code for i in p2.report.issues}
    # Across-party duplicate JMBG must surface
    assert Codes.JMBG_DUPLICATE in codes

    # ── Phase 2.5: resolve the conflict in favor of СНС ────
    persons = read_master(master_path)
    log_path = tmp_path / "corrections.log"
    state = CorrectionsState(persons=persons, log_path=log_path)
    conflicts = detect_conflicts(persons)
    assert len(conflicts) == 1
    # Pick the slot_key of the СНС submission to keep — v0.5.1 switched
    # resolve_conflict from party-name dispatch to slot-key dispatch.
    sns_sub = next(s for s in conflicts[0].submissions if s.party == "СНС")
    state.resolve_conflict(
        conflicts[0], keep_slot_key=sns_sub.slot_key, note="Прва пријава",
    )
    # After resolution: only one occurrence of same_jmbg remains
    after = [p for p in persons if p.jmbg == same_jmbg and p.is_filled]
    assert len(after) == 1
    assert after[0].party == "СНС"

    # Manual replacement on a different slot (simulate illness)
    target = next(p for p in persons if p.party == "СПС" and p.is_filled)
    replacement_jmbg = _build_valid("111122223333")
    state.replace_member(target.bm_number, target.function, {
        "name": "Замена Особа", "jmbg": replacement_jmbg,
    }, note="Болест")
    swapped = next(p for p in persons if p.slot_key == target.slot_key)
    assert swapped.name == "Замена Особа"
    assert swapped.jmbg == replacement_jmbg

    # Audit log captured both events
    events = read_log(log_path)
    types = {e["event_type"] for e in events}
    assert "resolve_conflict" in types
    assert "replace_member" in types

    # Persist corrections back to master
    write_master(persons, oik, master_path)

    # ── Phase 3 ────────────────────────────────────────────
    resenja = tmp_path / "resenja"
    p3 = generate_documents(persons, oik, election, resenja)
    assert len(p3.files_written) == 2
    for fp in p3.files_written:
        assert fp.exists()
        doc = Document(fp)
        text = "\n".join(p.text for p in doc.paragraphs)
        # Title present (compact form per reference .docx, no inter-letter spaces).
        assert "РЕШЕЊЕ" in text


def test_smoke_main_window_constructs(tmp_path: Path, monkeypatch) -> None:
    """GUI sanity: the main window builds without exceptions in a temp home."""
    monkeypatch.setenv("BIRACKI_ODBOR_HOME", str(tmp_path / "home"))
    from PySide6.QtWidgets import QApplication
    from biracki_odbor import _release_info
    from biracki_odbor.gui.main_window import MainWindow
    from biracki_odbor.gui.style import apply_global_style

    # Documented opt-out: empty GITHUB_REPO disables the updater so the
    # smoke test doesn't spawn a network-bound QThread that outlives it.
    monkeypatch.setattr(_release_info, "GITHUB_REPO", "")

    app = QApplication.instance() or QApplication([])
    apply_global_style(app)
    w = MainWindow()
    try:
        # V1 sidebar layout: 6 top-level sections (Изборна, Стални,
        # Прошир (V2), Посматрачи (V2), Поставке ОИК, О апликацији). Each
        # section is one stack item — settings pages are no longer nested
        # under a single Поставке tab.
        from biracki_odbor.gui.sidebar import (
            SECTION_ABOUT, SECTION_ALATI, SECTION_IZBORNA, SECTION_OIK,
            SECTION_POSMATRACI, SECTION_PROSIRENI, SECTION_STALNI, SECTION_ZAKONI,
        )
        assert w.stack.count() == 8  # SECTION_ZAKONI + SECTION_ALATI
        # Workflow tabs: Parties (0) + Consolidation (1) + Corrections (2)
        # + Decisions (3) + Amendments (4). Template moved to Избори → tab 3.
        assert w.tabs.count() == 5
        workflow_labels = [w.tabs.tabText(i) for i in range(w.tabs.count())]
        assert "Parties" in workflow_labels[0]
        assert "Amendments" in workflow_labels[4]
        assert w.izborna_page is not None
        assert w.oik_page is not None
        assert w.about_page is not None
        assert w.prosireni_page is not None
        assert w.posmatraci_page is not None
        assert w.zakoni_page is not None
        assert w.alati_page is not None
        # Prosireni workspace: 1. Templates, 2. Aggregation, 3. Corrections,
        # 4. Decisions, 5. Amendments (чл. 15 Cessation of function now
        # lives inside step 5's own grouped sections, not a separate tab)
        assert w.prosireni_page.tabs.count() == 5

        # Default landing: empty home → onboarding gate redirects to election
        # settings (SECTION_IZBORNA) so the operator fills in the config first.
        assert w.stack.currentIndex() == SECTION_IZBORNA

        # All tabs have a numeric prefix.
        assert all(w.tabs.tabText(i)[0].isdigit() for i in range(w.tabs.count()))

        # Empty config → workflow sections disabled until setup is complete.
        # Посматрачи stays disabled always (still V2 placeholder).
        assert not w.sidebar._items[SECTION_STALNI].isEnabled()
        assert not w.sidebar._items[SECTION_PROSIRENI].isEnabled()
        assert not w.sidebar._items[SECTION_POSMATRACI].isEnabled()
        assert w.sidebar._items[SECTION_IZBORNA].isEnabled()
        assert w.sidebar._items[SECTION_OIK].isEnabled()
        assert w.sidebar._items[SECTION_ABOUT].isEnabled()

        # Sidebar emit → stack switches. Each enabled section maps to its
        # own stack slot in V1 (formerly Изборна + ОИК shared a slot via
        # inner tabs).
        w.sidebar.section_changed.emit(SECTION_IZBORNA)
        assert w.stack.currentIndex() == SECTION_IZBORNA
        assert w.stack.currentWidget() is w.izborna_page
        w.sidebar.section_changed.emit(SECTION_OIK)
        assert w.stack.currentIndex() == SECTION_OIK
        assert w.stack.currentWidget() is w.oik_page
        w.sidebar.section_changed.emit(SECTION_ABOUT)
        assert w.stack.currentIndex() == SECTION_ABOUT
        w.sidebar.section_changed.emit(SECTION_STALNI)
        assert w.stack.currentIndex() == SECTION_STALNI
    finally:
        # MainWindow schedules an immediate auto-backup check
        # (QTimer.singleShot(0, ...)) plus an hourly repeating timer on
        # construction — without stopping/draining these, they stay queued
        # against this test's (about-to-be-deleted) tmp_path home and fire
        # unpredictably later, during whichever future test first pumps
        # the event loop in a sustained way (seen causing a real CI
        # failure in tests/test_updates.py, far later in the suite).
        w._auto_backup_timer.stop()
        w.deleteLater()
        app.processEvents()


def test_smoke_installer_wizard_constructs() -> None:
    """Installer sanity: wizard pages all build."""
    from PySide6.QtWidgets import QApplication
    from biracki_odbor.installer.installer_main import Installer
    from biracki_odbor.gui.style import apply_global_style

    app = QApplication.instance() or QApplication([])
    apply_global_style(app)
    w = Installer()
    # Welcome / Location / Install / Restore / Done.
    assert len(w.pageIds()) == 5


def test_smoke_uninstaller_wizard_constructs() -> None:
    """Uninstaller sanity: all four pages build."""
    from PySide6.QtWidgets import QApplication
    from biracki_odbor.installer.uninstaller_main import Uninstaller
    from biracki_odbor.gui.style import apply_global_style

    app = QApplication.instance() or QApplication([])
    apply_global_style(app)
    w = Uninstaller()
    # Welcome / Options / Execute / Done
    assert len(w.pageIds()) == 4
