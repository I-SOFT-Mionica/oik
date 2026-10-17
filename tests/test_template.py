"""Tests for core.template."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from biracki_odbor.core import template
from biracki_odbor.core.models import FUNCTIONS, PollingStation
from biracki_odbor.validation.reports import Codes
from tests.fixtures.stalni_sastav import build_template


def test_build_blank_template_default_uses_sequential_range(tmp_path: Path) -> None:
    p = template.build_blank_template(tmp_path / "blank.xlsx", n_polling_stations=3)
    wb = openpyxl.load_workbook(p)
    ws_r = wb[template.SHEET_RASPODELA]
    assert [ws_r.cell(row=r, column=1).value for r in (2, 3, 4)] == [1, 2, 3]


def test_build_blank_template_uses_real_stations_when_given(tmp_path: Path) -> None:
    """Regression: a generated matrix must carry the OIK's ACTUAL BM
    numbers (which may be non-sequential), not a synthetic 1..N range —
    otherwise it round-trips straight into BM_UNKNOWN warnings at Phase 1.
    """
    stations = [
        PollingStation(bm_number=7, name="Школа А"),
        PollingStation(bm_number=2, name="Дом културе"),
        PollingStation(bm_number=15, name="Месна заједница"),
    ]
    p = template.build_blank_template(tmp_path / "blank.xlsx", stations=stations)
    wb = openpyxl.load_workbook(p)

    ws_r = wb[template.SHEET_RASPODELA]
    assert [ws_r.cell(row=r, column=1).value for r in (2, 3, 4)] == [2, 7, 15]

    ws_bm = wb[template.SHEET_BM]
    assert [ws_bm.cell(row=r, column=1).value for r in (2, 3, 4)] == [2, 7, 15]
    assert ws_bm.cell(row=2, column=2).value == "Дом културе"


def test_build_blank_template_allocations_fill_matching_cells(tmp_path: Path) -> None:
    """The in-app matrix builder writes cell values via `allocations` —
    each (bm, function) slot must land in the exact cell the header
    contract promises, and slots not covered stay blank.
    """
    from biracki_odbor.core.models import Allocation

    stations = [
        PollingStation(bm_number=1, name="Школа А"),
        PollingStation(bm_number=2, name="Дом културе"),
    ]
    allocations = [
        Allocation(bm_number=1, function="Председник", party="СНС"),
        Allocation(bm_number=1, function="Члан 2", party="СПС"),
        Allocation(bm_number=2, function="Заменик председника", party="ЗЛФ"),
    ]
    p = template.build_blank_template(
        tmp_path / "blank.xlsx", stations=stations, allocations=allocations,
    )
    wb = openpyxl.load_workbook(p)
    ws_r = wb[template.SHEET_RASPODELA]

    assert ws_r.cell(row=2, column=2).value == "СНС"          # BM1, Председник
    assert ws_r.cell(row=2, column=6).value == "СПС"          # BM1, Члан 2
    assert ws_r.cell(row=3, column=3).value == "ЗЛФ"          # BM2, Заменик председника
    # Untouched slots stay blank, not missing.
    assert ws_r.cell(row=2, column=3).value in (None, "")
    assert ws_r.cell(row=3, column=2).value in (None, "")

    data, report = template.read_template(p)
    assert report.ok, [i.message_sr for i in report.issues]
    assert len(data.allocations) == 3


def test_build_blank_template_headers_match_canonical_style_and_reimport(tmp_path: Path) -> None:
    """The exporter must use our canonical spelling (WITH spaces, "БР
    БМ") — NOT the real-world no-space RIK style — because already-
    deployed builds only recognize the canonical form; only the reader's
    tolerance (2026-07-10) understands both. Also confirms whatever it
    produces reads back cleanly through read_template().
    """
    stations = [PollingStation(bm_number=1, name="Школа А")]
    p = template.build_blank_template(tmp_path / "blank.xlsx", stations=stations)
    wb = openpyxl.load_workbook(p)
    ws_r = wb[template.SHEET_RASPODELA]
    assert ws_r.cell(row=1, column=1).value == "БР БМ"
    assert ws_r.cell(row=1, column=3).value == "Заменик председника"
    wb.close()

    # Fill in one row (as an operator would) and confirm it reads back.
    wb = openpyxl.load_workbook(p)
    ws_r = wb[template.SHEET_RASPODELA]
    for c in range(2, 8):
        ws_r.cell(row=2, column=c, value="СНС")
    wb.save(p)
    wb.close()

    data, report = template.read_template(p)
    assert report.ok, [i.message_sr for i in report.issues]
    assert len(data.allocations) == 6


def test_raspodela_header_past_party_quota_preamble(tmp_path: Path) -> None:
    """Regression (2026-07-10, real RIK file): the Raspodela sheet isn't
    just the BM matrix — RIK prepends a party-seat-quota summary table
    (one row per party + a "УКУПНО" total row + a blank separator)
    ABOVE the actual "БР БМ" header. On a real file this pushed the
    real header to row 20. The old 10-row scan window silently
    HEADER_MISMATCHed — this was the actual "BM not in the list" root
    cause reported from the field, not a header-spelling issue.
    """
    p = tmp_path / "real_shape.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = template.SHEET_RASPODELA

    # 17-row party quota preamble + "УКУПНО" row + blank separator = the
    # real header lands on row 20, well past a small fixed scan window.
    ws.cell(row=1, column=1, value="Посланичка група")
    for r in range(2, 19):
        ws.cell(row=r, column=1, value=f"Странка {r}")
    ws.cell(row=18, column=1, value="УКУПНО (За 52 бирачка места)")
    # row 19 intentionally blank (separator)
    ws.cell(row=20, column=1, value="БР БМ")
    for c, func in enumerate(FUNCTIONS, start=2):
        ws.cell(row=20, column=c, value=func)
    ws.cell(row=21, column=1, value=1)
    for c in range(2, 8):
        ws.cell(row=21, column=c, value="СНС")
    wb.save(p)
    wb.close()

    data, report = template.read_template(p)
    assert report.ok, [i.message_sr for i in report.issues]
    assert len(data.allocations) == 6


def test_read_valid_template(tmp_path: Path) -> None:
    p = build_template(tmp_path / "t.xlsx", n_polling_stations=3)
    data, report = template.read_template(p)
    assert report.ok, [i.message_sr for i in report.issues]
    # 3 BMs × 6 functions = 18 allocations
    assert len(data.allocations) == 18
    assert len(data.polling_stations_from_template) == 3
    assert len(data.parties) >= 2


def test_missing_file(tmp_path: Path) -> None:
    data, report = template.read_template(tmp_path / "nope.xlsx")
    codes = {i.code for i in report.issues}
    assert Codes.FILE_MISSING in codes
    assert data.allocations == []


def test_read_bm_sheet_from_stalni_sastav(tmp_path: Path) -> None:
    """read_bm_sheet pulls polling stations out of the bm sheet."""
    p = build_template(tmp_path / "t.xlsx", n_polling_stations=4)
    stations, errors = template.read_bm_sheet(p)
    assert errors == []
    assert len(stations) == 4
    assert stations[0].bm_number == 1
    assert "БМ 1" in stations[0].name


def test_read_bm_sheet_missing_file(tmp_path: Path) -> None:
    stations, errors = template.read_bm_sheet(tmp_path / "nope.xlsx")
    assert stations == []
    assert errors and "не постоји" in errors[0]


def test_read_bm_sheet_empty_workbook(tmp_path: Path) -> None:
    p = tmp_path / "empty.xlsx"
    wb = openpyxl.Workbook()
    wb.save(p)
    wb.close()
    stations, errors = template.read_bm_sheet(p)
    assert stations == []
    assert errors


def test_parties_workbook_round_trip(tmp_path: Path) -> None:
    p = tmp_path / "parties.xlsx"
    src = {
        "СНС": "Посланичке групе Србија не сме да стане",
        "СПС": "Посланичке групе Ивица Дачић",
        "_default": "начелника Општинске управе",
    }
    template.build_parties_workbook(src, p)
    loaded, contacts, errors = template.read_parties_sheet(p)
    assert errors == []
    assert loaded == src
    # No contacts emitted when none were supplied to build_parties_workbook.
    assert contacts == {}


def test_parties_workbook_with_contacts(tmp_path: Path) -> None:
    """v0.11+ schema: parties.xlsx grows to 4 columns. Contacts round-trip
    only for rows where email OR phone is non-empty — fully-empty
    contacts don't pollute the import dict."""
    from biracki_odbor.config.election import PartyContact
    p = tmp_path / "parties.xlsx"
    src = {"СНС": "Посланичка група", "СПС": "Друга странка", "_default": "."}
    contacts = {
        "СНС": PartyContact(email="kontakt@sns.rs", phone="+381 11 1111"),
        "СПС": PartyContact(email="info@sps.rs", phone=""),
    }
    template.build_parties_workbook(src, p, contacts=contacts)
    loaded, loaded_contacts, errors = template.read_parties_sheet(p)
    assert errors == []
    assert loaded == src
    assert loaded_contacts == {
        "СНС": {"email": "kontakt@sns.rs", "phone": "+381 11 1111"},
        "СПС": {"email": "info@sps.rs", "phone": ""},
    }


def test_read_parties_missing_file(tmp_path: Path) -> None:
    loaded, contacts, errors = template.read_parties_sheet(tmp_path / "nope.xlsx")
    assert loaded == {}
    assert contacts == {}
    assert errors and "не постоји" in errors[0]


def test_read_parties_empty_workbook(tmp_path: Path) -> None:
    p = tmp_path / "empty.xlsx"
    wb = openpyxl.Workbook()
    wb.save(p)
    wb.close()
    loaded, contacts, errors = template.read_parties_sheet(p)
    assert loaded == {}
    assert contacts == {}
    assert errors


def test_read_parties_rejects_wrong_headers(tmp_path: Path) -> None:
    """Picking Stalni_sastav.xlsx by mistake must NOT import garbage."""
    p = build_template(tmp_path / "stalni.xlsx", n_polling_stations=3,
                       parties=("СНС", "СПС"))
    loaded, contacts, errors = template.read_parties_sheet(p)
    assert loaded == {}
    assert contacts == {}
    assert errors and "странака" in errors[0].lower() or errors


def test_quotas_workbook_round_trip(tmp_path: Path) -> None:
    from biracki_odbor.core.models import PartyQuota

    p = tmp_path / "quotas.xlsx"
    src = {
        "СНС": PartyQuota(predsednik=23, zamenik=5, clan_i_zamenik=93),
        "СПС": PartyQuota(predsednik=5, zamenik=3, clan_i_zamenik=20),
    }
    template.build_quotas_workbook(src, p)
    loaded, errors = template.read_quotas_sheet(p)
    assert errors == []
    assert loaded == src


def test_quotas_workbook_non_numeric_cell_treated_as_zero(tmp_path: Path) -> None:
    """A manually-edited Excel file is more likely to have a stray blank
    cell than a genuinely corrupt one — don't reject the whole file."""
    p = tmp_path / "quotas.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "quotas"
    ws.append(["Кратки назив", "Председник", "Заменик", "Чл. + Зам. чл."])
    ws.append(["СНС", "не знам", None, 5])
    wb.save(p)
    wb.close()

    loaded, errors = template.read_quotas_sheet(p)
    assert errors == []
    from biracki_odbor.core.models import PartyQuota
    assert loaded == {"СНС": PartyQuota(predsednik=0, zamenik=0, clan_i_zamenik=5)}


def test_read_quotas_missing_file(tmp_path: Path) -> None:
    loaded, errors = template.read_quotas_sheet(tmp_path / "nope.xlsx")
    assert loaded == {}
    assert errors and "не постоји" in errors[0]


def test_read_quotas_rejects_wrong_headers(tmp_path: Path) -> None:
    p = build_template(tmp_path / "stalni.xlsx", n_polling_stations=3,
                       parties=("СНС", "СПС"))
    loaded, errors = template.read_quotas_sheet(p)
    assert loaded == {}
    assert errors


def test_read_bm_rejects_wrong_headers(tmp_path: Path) -> None:
    """A non-bm 2-col workbook must be rejected (e.g. parties.xlsx)."""
    p = tmp_path / "parties.xlsx"
    template.build_parties_workbook({"СНС": "Посланичка група"}, p)
    stations, errors = template.read_bm_sheet(p)
    assert stations == []
    assert errors


def test_read_bm_sheet_accepts_real_rik_header(tmp_path: Path) -> None:
    """Regression (2026-07-10, real RIK file): the 'bm' sheet header was
    "БМ #", not "БР"/"Број" — read_bm_sheet (the standalone "Бирачка
    места → Увоз" import) rejected it outright even though
    read_template()'s internal _read_bm() (no header validation) read
    the exact same file fine. Both paths must now accept it.
    """
    p = tmp_path / "real_bm.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = template.SHEET_BM
    for c, h in enumerate(("БМ #", "Назив", "Адреса", "Назив + Адреса"), start=1):
        ws.cell(row=1, column=c, value=h)
    ws.cell(row=2, column=1, value=1)
    ws.cell(row=2, column=2, value="ОСНОВНА ШКОЛА")
    wb.save(p)
    wb.close()

    stations, errors = template.read_bm_sheet(p)
    assert errors == []
    assert len(stations) == 1


def test_missing_raspodela_sheet(tmp_path: Path) -> None:
    p = tmp_path / "broken.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "bm"
    wb.save(p)
    wb.close()
    data, report = template.read_template(p)
    codes = {i.code for i in report.issues}
    assert Codes.SHEET_MISSING in codes


def test_raspodela_tolerates_real_world_header_variants(tmp_path: Path) -> None:
    """Regression: a real RIK-provided Raspodela sheet was seen using
    "Бројбирача" (no space) for the BM column and "Заменикпредседника" /
    "Заменикчлана 1" / "Заменикчлана 2" (no internal space) for the
    function columns — our canonical text has spaces in all of these.
    Must still be recognized, not HEADER_MISMATCH.
    """
    p = tmp_path / "real_world.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Raspodela"
    ws.cell(row=1, column=1, value="Бројбирача")
    no_space_headers = [
        "Председник", "Заменикпредседника", "Члан 1",
        "Заменикчлана 1", "Члан 2", "Заменикчлана 2",
    ]
    for c, h in enumerate(no_space_headers, start=2):
        ws.cell(row=1, column=c, value=h)
    ws.cell(row=2, column=1, value=1)
    for c in range(2, 8):
        ws.cell(row=2, column=c, value="СНС")
    wb.save(p)
    wb.close()

    data, report = template.read_template(p)
    assert report.ok, [i.message_sr for i in report.issues]
    assert len(data.allocations) == 6
    assert {a.function for a in data.allocations} == set(FUNCTIONS)


def test_missing_function_column(tmp_path: Path) -> None:
    p = tmp_path / "broken2.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Raspodela"
    ws.cell(row=1, column=1, value="БР БМ")
    for c, func in enumerate(FUNCTIONS[:-1], start=2):  # skip last function
        ws.cell(row=1, column=c, value=func)
    wb.save(p)
    wb.close()
    data, report = template.read_template(p)
    codes = {i.code for i in report.issues}
    assert Codes.HEADER_MISMATCH in codes


def test_no_data_rows(tmp_path: Path) -> None:
    p = tmp_path / "empty.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Raspodela"
    ws.cell(row=1, column=1, value="БР БМ")
    for c, func in enumerate(FUNCTIONS, start=2):
        ws.cell(row=1, column=c, value=func)
    wb.save(p)
    wb.close()
    data, report = template.read_template(p)
    codes = {i.code for i in report.issues}
    assert Codes.NO_ALLOCATIONS in codes
