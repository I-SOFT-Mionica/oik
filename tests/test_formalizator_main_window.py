"""Tests for formalizator.gui.main_window.MainWindow — the two-screen
(setup -> fill-out) flow host."""

from __future__ import annotations

from pathlib import Path

import pytest


@pytest.fixture
def app():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def _make_window(app, tmp_path: Path):
    from formalizator.gui.main_window import MainWindow
    return MainWindow(home_dir=tmp_path)


def _blank_background(path: Path) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path), pagesize=A4)
    c.showPage()
    c.save()
    return path


def _seed_template(tmp_path: Path, name: str = "Тест шаблон"):
    from formalizator.core import template_catalog
    from formalizator.core.models import CalibratedTemplate, FieldPosition

    bg = _blank_background(tmp_path / "bg.pdf")
    tpl = CalibratedTemplate(
        name=name, background_pdf_path=str(bg),
        positions=[FieldPosition(field_key="full_name", page_index=0, x=100, y=700)],
    )
    template_catalog.save_template(tmp_path / "templates", tpl)
    return tpl


def _seed_template_with_custom_fields(tmp_path: Path, name: str = "Тест шаблон"):
    from formalizator.core import template_catalog
    from formalizator.core.models import CalibratedTemplate, CustomFieldDef, FieldPosition

    bg = _blank_background(tmp_path / "bg.pdf")
    tpl = CalibratedTemplate(
        name=name, background_pdf_path=str(bg),
        positions=[
            FieldPosition(field_key="full_name", page_index=0, x=100, y=700),
            FieldPosition(field_key="overitelj_ime", page_index=0, x=100, y=600),
            FieldPosition(field_key="napomena", page_index=0, x=100, y=500),
        ],
        custom_fields=[
            CustomFieldDef(key="overitelj_ime", label="Оверитељ — име и презиме", kind="universal"),
            CustomFieldDef(key="napomena", label="Напомена", kind="individual"),
        ],
    )
    template_catalog.save_template(tmp_path / "templates", tpl)
    return tpl


def test_starts_on_setup_screen(app, tmp_path: Path) -> None:
    w = _make_window(app, tmp_path)
    try:
        assert w.currentWidget() is w.setup_screen
    finally:
        w.deleteLater()


def test_setup_lists_seeded_templates(app, tmp_path: Path) -> None:
    _seed_template(tmp_path)
    w = _make_window(app, tmp_path)
    try:
        items = [w.setup_screen.cb_template.itemText(i) for i in range(w.setup_screen.cb_template.count())]
        assert "Тест шаблон" in items
    finally:
        w.deleteLater()


def test_completing_setup_switches_to_fillout(app, tmp_path: Path) -> None:
    _seed_template(tmp_path)
    w = _make_window(app, tmp_path)
    try:
        w.setup_screen.cb_template.setCurrentText("Тест шаблон")
        w.setup_screen.sb_copies.setValue(2)
        w.setup_screen._on_continue()  # noqa: SLF001

        assert w.currentWidget() is w.fillout_screen
        assert w.fillout_screen._copies == 2  # noqa: SLF001
    finally:
        w.deleteLater()


def test_printing_an_entry_updates_count_and_persists(app, tmp_path: Path, monkeypatch) -> None:
    import formalizator.gui.fillout_screen as fillout_mod
    monkeypatch.setattr(fillout_mod, "print_pdf", lambda *a, **kw: None)

    _seed_template(tmp_path)
    w = _make_window(app, tmp_path)
    try:
        w.setup_screen.cb_template.setCurrentText("Тест шаблон")
        w.setup_screen._on_continue()  # noqa: SLF001

        w.fillout_screen.ed_given_name.setText("Марко")
        w.fillout_screen.ed_surname.setText("Марковић")
        # A valid JMBG is required to print — an entry without one
        # cannot legally appear on the NPRS-4 roster.
        w.fillout_screen.ed_jmbg.setText("2308985774805")
        w.fillout_screen._on_print()  # noqa: SLF001

        assert w.fillout_screen.lbl_count.text() == "Унето: 1"

        from formalizator.core import session_store
        entries = session_store.list_entries(w._db_path, w._session_id)  # noqa: SLF001
        assert len(entries) == 1
        assert entries[0].given_name == "Марко"
    finally:
        w.deleteLater()


def _configure_fake_export_template(monkeypatch, setup_mod, export_template) -> None:
    """Stands in for the operator clicking "+ Нови извозни шаблон…" on
    the setup screen and completing the definition dialog — same
    exec()-can't-run-headlessly reasoning as the calibration wizard's
    own custom-field dialog tests."""

    class _FakeDialog:
        DialogCode = setup_mod.ExportTemplateDialog.DialogCode

        def __init__(self, custom_fields=None, parent=None, home_dir=None) -> None:
            pass

        def exec(self):
            return self.DialogCode.Accepted

        def result_template(self):
            return export_template

    monkeypatch.setattr(setup_mod, "ExportTemplateDialog", _FakeDialog)


def test_finish_session_writes_a_roster_using_the_configured_export_template(
    app, tmp_path: Path, monkeypatch,
) -> None:
    """The export step is configured once, on the setup screen, right
    after the template — not picked again later. Ending the session
    is what actually produces the file."""
    import formalizator.gui.fillout_screen as fillout_mod
    import formalizator.gui.main_window as main_window_mod
    import formalizator.gui.setup_screen as setup_mod
    from formalizator.core.models import ExportColumnDef, ExportTemplateDef

    monkeypatch.setattr(fillout_mod, "print_pdf", lambda *a, **kw: None)
    # _on_finish_session ends with a real QMessageBox.information() —
    # left unmocked once already in this session, it popped a real
    # blocking dialog and turned a fast test into a 34s hang.
    monkeypatch.setattr(main_window_mod.QMessageBox, "information", lambda *a, **kw: None)

    export_template = ExportTemplateDef(
        name="Тест извоз", columns=[ExportColumnDef(header="Име и презиме", field_keys=["full_name"])],
    )
    _configure_fake_export_template(monkeypatch, setup_mod, export_template)

    _seed_template(tmp_path)
    w = _make_window(app, tmp_path)
    try:
        w.setup_screen.cb_template.setCurrentText("Тест шаблон")

        idx = w.setup_screen.cb_export_template.findData(setup_mod._NEW_EXPORT_TEMPLATE_SENTINEL)  # noqa: SLF001
        w.setup_screen.cb_export_template.setCurrentIndex(idx)
        assert w.setup_screen.cb_export_template.currentText() == "Тест извоз"

        w.setup_screen._on_continue()  # noqa: SLF001

        w.fillout_screen.ed_given_name.setText("Марко")
        w.fillout_screen.ed_surname.setText("Марковић")
        # A valid JMBG is required to print — an entry without one
        # cannot legally appear on the NPRS-4 roster.
        w.fillout_screen.ed_jmbg.setText("2308985774805")
        w.fillout_screen._on_print()  # noqa: SLF001

        out_path = tmp_path / "roster.xlsx"
        monkeypatch.setattr(
            main_window_mod.QFileDialog, "getSaveFileName",
            staticmethod(lambda *a, **kw: (str(out_path), "")),
        )

        w._on_finish_session()  # noqa: SLF001

        assert out_path.exists()
        import openpyxl
        ws = openpyxl.load_workbook(out_path).active
        assert ws.cell(row=1, column=1).value == "Име и презиме"
        assert ws.cell(row=2, column=1).value == "Марко Марковић"

        # Ending the session closes it — resuming shouldn't find it open.
        from formalizator.core import session_store
        assert session_store.get_open_session(w._db_path) is None  # noqa: SLF001
        assert w.currentWidget() is w.setup_screen
    finally:
        w.deleteLater()


def test_finish_session_without_export_template_just_closes_the_session(
    app, tmp_path: Path, monkeypatch,
) -> None:
    """"(без извоза)" is the default — ending a session shouldn't
    force a save dialog nobody asked for."""
    import formalizator.gui.fillout_screen as fillout_mod
    import formalizator.gui.main_window as main_window_mod
    monkeypatch.setattr(fillout_mod, "print_pdf", lambda *a, **kw: None)

    def _fail_if_called(*a, **kw):
        raise AssertionError("no export template configured — getSaveFileName shouldn't be called")

    monkeypatch.setattr(main_window_mod.QFileDialog, "getSaveFileName", staticmethod(_fail_if_called))

    _seed_template(tmp_path)
    w = _make_window(app, tmp_path)
    try:
        w.setup_screen.cb_template.setCurrentText("Тест шаблон")
        w.setup_screen._on_continue()  # noqa: SLF001

        w.fillout_screen.ed_given_name.setText("Марко")
        w.fillout_screen.ed_surname.setText("Марковић")
        # A valid JMBG is required to print — an entry without one
        # cannot legally appear on the NPRS-4 roster.
        w.fillout_screen.ed_jmbg.setText("2308985774805")
        w.fillout_screen._on_print()  # noqa: SLF001

        w._on_finish_session()  # noqa: SLF001

        from formalizator.core import session_store
        assert session_store.get_open_session(w._db_path) is None  # noqa: SLF001
    finally:
        w.deleteLater()


def test_edit_export_template_button_with_no_export_offers_to_create_one(
    app, tmp_path: Path, monkeypatch,
) -> None:
    """Reported as "I don't see a way to configure export" — root
    cause: clicking "Уреди…" with "(без извоза)" selected silently did
    nothing at all, no matter what the operator tried. It has to open
    the definition dialog instead of being a dead end."""
    import formalizator.gui.setup_screen as setup_mod
    from formalizator.core.models import ExportColumnDef, ExportTemplateDef

    export_template = ExportTemplateDef(
        name="Нови извоз", columns=[ExportColumnDef(header="Име", field_keys=["full_name"])],
    )
    _configure_fake_export_template(monkeypatch, setup_mod, export_template)

    _seed_template(tmp_path)
    w = _make_window(app, tmp_path)
    try:
        w.setup_screen.cb_template.setCurrentText("Тест шаблон")
        assert w.setup_screen.cb_export_template.currentText() == "(без извоза)"

        w.setup_screen._edit_export_template()  # noqa: SLF001

        assert w.setup_screen.cb_export_template.currentText() == "Нови извоз"
    finally:
        w.deleteLater()


def test_permanent_data_hidden_until_template_chosen(app, tmp_path: Path) -> None:
    # isVisible() also depends on the whole ancestor chain actually
    # being shown on screen — irrelevant here and false for every
    # widget in an unshown test window — so isHidden() (this widget's
    # own explicit shown/hidden state) is what we actually want.
    _seed_template(tmp_path)
    w = _make_window(app, tmp_path)
    try:
        assert w.setup_screen.copies_box.isHidden() is True
        assert w.setup_screen.btn_continue.isHidden() is True

        w.setup_screen.cb_template.setCurrentText("Тест шаблон")

        assert w.setup_screen.copies_box.isHidden() is False
        assert w.setup_screen.btn_continue.isHidden() is False
    finally:
        w.deleteLater()


def test_edit_template_button_enabled_only_once_a_template_is_selected(app, tmp_path: Path) -> None:
    _seed_template(tmp_path)
    w = _make_window(app, tmp_path)
    try:
        assert not w.setup_screen.btn_edit_template.isEnabled()

        w.setup_screen.cb_template.setCurrentText("Тест шаблон")

        assert w.setup_screen.btn_edit_template.isEnabled()
    finally:
        w.deleteLater()


def test_edit_template_button_loads_and_resaves_the_template(app, tmp_path: Path, monkeypatch) -> None:
    """Calibrating is a lot of fiddly clicking — this is what lets an
    already-saved template be reopened and added to instead of forcing
    the whole thing to be redone in one sitting."""
    import formalizator.gui.setup_screen as mod
    from formalizator.core import template_catalog
    from formalizator.core.models import CalibratedTemplate, FieldPosition

    _seed_template(tmp_path)
    w = _make_window(app, tmp_path)
    try:
        w.setup_screen.cb_template.setCurrentText("Тест шаблон")

        calls = {}

        class _FakeDialog:
            DialogCode = mod.NewTemplateDialog.DialogCode

            def __init__(self, parent=None) -> None:
                pass

            def load_template(self, tpl) -> None:
                calls["loaded"] = tpl

            def exec(self):
                return self.DialogCode.Accepted

            def result_template(self):
                return CalibratedTemplate(
                    name="Тест шаблон",
                    background_pdf_path=calls["loaded"].background_pdf_path,
                    positions=[FieldPosition(field_key="surname", page_index=0, x=1.0, y=2.0)],
                )

        monkeypatch.setattr(mod, "NewTemplateDialog", _FakeDialog)

        w.setup_screen._edit_selected_template()  # noqa: SLF001

        assert calls["loaded"].name == "Тест шаблон"
        reloaded = template_catalog.find_by_name(tmp_path / "templates", "Тест шаблон")
        assert reloaded.positions[0].field_key == "surname"
    finally:
        w.deleteLater()


def test_universal_custom_field_shown_on_setup_and_saved_to_session(app, tmp_path: Path) -> None:
    _seed_template_with_custom_fields(tmp_path)
    w = _make_window(app, tmp_path)
    try:
        w.setup_screen.cb_template.setCurrentText("Тест шаблон")
        assert "overitelj_ime" in w.setup_screen._custom_universal_edits  # noqa: SLF001
        # The individual-kind custom field belongs on the fill-out
        # screen, not here.
        assert "napomena" not in w.setup_screen._custom_universal_edits  # noqa: SLF001

        w.setup_screen._custom_universal_edits["overitelj_ime"].setText("Петар Петровић")  # noqa: SLF001
        w.setup_screen._on_continue()  # noqa: SLF001

        from formalizator.core import session_store
        _, info = session_store.get_open_session(w._db_path)  # noqa: SLF001
        assert info.custom_values["overitelj_ime"] == "Петар Петровић"
    finally:
        w.deleteLater()


def test_individual_custom_field_shown_on_fillout_and_saved_to_entry(app, tmp_path: Path, monkeypatch) -> None:
    import formalizator.gui.fillout_screen as fillout_mod
    monkeypatch.setattr(fillout_mod, "print_pdf", lambda *a, **kw: None)

    _seed_template_with_custom_fields(tmp_path)
    w = _make_window(app, tmp_path)
    try:
        w.setup_screen.cb_template.setCurrentText("Тест шаблон")
        w.setup_screen._on_continue()  # noqa: SLF001

        assert "napomena" in w.fillout_screen._custom_edits  # noqa: SLF001
        assert "overitelj_ime" not in w.fillout_screen._custom_edits  # noqa: SLF001

        w.fillout_screen.ed_given_name.setText("Марко")
        w.fillout_screen.ed_surname.setText("Марковић")
        # A valid JMBG is required to print — an entry without one
        # cannot legally appear on the NPRS-4 roster.
        w.fillout_screen.ed_jmbg.setText("2308985774805")
        w.fillout_screen._custom_edits["napomena"].setText("нешто битно")  # noqa: SLF001
        w.fillout_screen._on_print()  # noqa: SLF001

        from formalizator.core import session_store
        entries = session_store.list_entries(w._db_path, w._session_id)  # noqa: SLF001
        assert entries[0].custom_values["napomena"] == "нешто битно"
    finally:
        w.deleteLater()


def test_back_button_returns_to_setup(app, tmp_path: Path) -> None:
    _seed_template(tmp_path)
    w = _make_window(app, tmp_path)
    try:
        w.setup_screen.cb_template.setCurrentText("Тест шаблон")
        w.setup_screen._on_continue()  # noqa: SLF001

        w.fillout_screen.back_requested.emit()

        assert w.currentWidget() is w.setup_screen
    finally:
        w.deleteLater()


# The smallest usable viewport this app actually has to survive: a
# 1920x1080 laptop at 150% Windows scaling is 1280x720 logical pixels,
# and the taskbar takes roughly another 48 from the bottom. The 125%
# case (864, or 816 under the taskbar) is the machine the field reports
# came from.
_SMALLEST_LAPTOP_VIEWPORT_HEIGHT = 672


def test_both_screens_fit_inside_a_scaled_laptop_screen(app, tmp_path: Path) -> None:
    """Neither screen may demand more height than the screen has.

    This has now caused three separate field failures — clipped text in
    entry fields, action buttons sitting under the taskbar, and the
    symptom easing when the taskbar was set to auto-hide (which is what
    finally identified it: auto-hide changes nothing except available
    height). A screen that asks for more vertical space than exists
    does not compress gracefully, it gets cut off at the bottom."""
    w = _make_window(app, tmp_path)
    try:
        for screen in (w.setup_screen, w.fillout_screen):
            required = screen.minimumSizeHint().height()
            assert required <= _SMALLEST_LAPTOP_VIEWPORT_HEIGHT, (
                f"{type(screen).__name__} needs {required}px of height but a "
                f"scaled laptop offers {_SMALLEST_LAPTOP_VIEWPORT_HEIGHT}px"
            )
        assert w.minimumSize().height() <= _SMALLEST_LAPTOP_VIEWPORT_HEIGHT
    finally:
        w.deleteLater()


@pytest.mark.parametrize("window_height", [950, 864, 816, 720, 672])
def test_the_primary_action_buttons_stay_on_screen_at_every_height(
    app, tmp_path: Path, window_height: int,
) -> None:
    """Scrolling fixes the clipping, but the buttons used on every
    single voter must never be something the operator has to scroll to
    find — so they are pinned outside the scroll areas."""
    w = _make_window(app, tmp_path)
    try:
        w.resize(1500, window_height)
        w.show()
        app.processEvents()

        w.setCurrentWidget(w.fillout_screen)
        app.processEvents()
        button = w.fillout_screen.btn_print
        bottom = button.mapTo(w, button.rect().bottomLeft()).y()
        assert 0 < bottom <= window_height, f"Штампај cut off at height {window_height}"

        w.setCurrentWidget(w.setup_screen)
        w.setup_screen.btn_continue.setVisible(True)
        app.processEvents()
        button = w.setup_screen.btn_continue
        bottom = button.mapTo(w, button.rect().bottomLeft()).y()
        assert 0 < bottom <= window_height, f"Настави cut off at height {window_height}"
    finally:
        w.hide()
        w.deleteLater()


def test_entry_fields_are_never_squeezed_below_their_natural_height(
    app, tmp_path: Path,
) -> None:
    """The stylesheet gives QLineEdit 8px of vertical padding, so a
    field shorter than its size hint has its text pushed out of view —
    the same defect already fixed once for table-cell editors. Scrolling
    means the fields keep their full height instead of absorbing the
    shortfall."""
    w = _make_window(app, tmp_path)
    try:
        w.resize(1500, _SMALLEST_LAPTOP_VIEWPORT_HEIGHT)
        w.show()
        w.setCurrentWidget(w.fillout_screen)
        app.processEvents()

        field = w.fillout_screen.ed_given_name
        assert field.height() >= field.sizeHint().height()
    finally:
        w.hide()
        w.deleteLater()
