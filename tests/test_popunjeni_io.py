"""Tests for biracki_odbor.core.popunjeni_io — safety + identity + import."""

from __future__ import annotations

import zipfile
from pathlib import Path

import openpyxl
import pytest

from biracki_odbor.config.election import (
    DocumentInfo, ElectionConfig, ElectionDates, ElectionInfo,
)
from biracki_odbor.core import popunjeni_io
from biracki_odbor.core.phase1_generate import (
    MARKER_CELL, MARKER_COLUMN_LETTER, build_marker, generate_party_files,
)
from biracki_odbor.validation.reports import Codes
from tests.fixtures.stalni_sastav import build_template
from tests.test_phase1 import _make_oik


def _make_election(uuid: str = "test-uuid-1234") -> ElectionConfig:
    return ElectionConfig(
        election=ElectionInfo(
            dates=ElectionDates(
                date_full="17. децембар 2026. године",
                session_date="6. децембра 2026. године",
                session_number="5",
            ),
            types=["parliamentary"],
            rik_decision="Одлуком РИК-а",
        ),
        document=DocumentInfo(
            number_prefix="013-5-30",
            legal_basis="На основу члана 32. Закона",
        ),
        party_full_names={"СНС": "Посланичка група X"},
        election_uuid=uuid,
    )


def _phase1_file(tmp_path: Path, *, election: ElectionConfig) -> Path:
    """Generate one real Phase 1 .xlsx that carries the marker."""
    template = build_template(
        tmp_path / "tpl.xlsx",
        n_polling_stations=1, parties=("СНС",),
    )
    out = tmp_path / "stranke"
    r = generate_party_files(template, out, _make_oik(1), election=election)
    assert r.files_written, "Phase 1 should produce at least one file"
    return r.files_written[0]


# ── Safety pipeline ────────────────────────────────────────


def test_safety_rejects_oversize(tmp_path: Path) -> None:
    p = tmp_path / "big.xlsx"
    # Just over 10 MB cap.
    p.write_bytes(b"0" * (10 * 1024 * 1024 + 1))
    r = popunjeni_io.check_file_safety(p)
    assert not r.ok
    assert r.code == Codes.FILE_TOO_LARGE


def test_safety_rejects_undersize(tmp_path: Path) -> None:
    p = tmp_path / "tiny.xlsx"
    p.write_bytes(b"x" * 200)
    r = popunjeni_io.check_file_safety(p)
    assert not r.ok
    assert r.code == Codes.FILE_TOO_SMALL


def test_safety_rejects_non_xlsx(tmp_path: Path) -> None:
    p = tmp_path / "fake.xlsx"
    p.write_bytes(b"this is not a zip" * 100)
    r = popunjeni_io.check_file_safety(p)
    assert not r.ok
    assert r.code == Codes.NOT_XLSX


def test_safety_rejects_macros(tmp_path: Path) -> None:
    p = tmp_path / "macro.xlsx"
    with zipfile.ZipFile(p, "w") as zf:
        zf.writestr("xl/vbaProject.bin", b"\x00" * 2000)
        zf.writestr("xl/workbook.xml", b"<workbook/>")
        zf.writestr("[Content_Types].xml", b"<Types/>")
    r = popunjeni_io.check_file_safety(p)
    assert not r.ok
    assert r.code == Codes.MACROS_PRESENT


def test_safety_rejects_ole_embeddings(tmp_path: Path) -> None:
    p = tmp_path / "ole.xlsx"
    with zipfile.ZipFile(p, "w") as zf:
        zf.writestr("xl/embeddings/oleObject1.bin", b"\x00" * 2000)
        zf.writestr("xl/workbook.xml", b"<workbook/>")
        zf.writestr("[Content_Types].xml", b"<Types/>")
    r = popunjeni_io.check_file_safety(p)
    assert not r.ok
    assert r.code == Codes.EMBEDDED_OBJECTS


def test_safety_rejects_external_links(tmp_path: Path) -> None:
    p = tmp_path / "external.xlsx"
    with zipfile.ZipFile(p, "w") as zf:
        zf.writestr("xl/externalLinks/externalLink1.xml", b"<link/>")
        zf.writestr("xl/workbook.xml", b"<workbook/>")
        zf.writestr("[Content_Types].xml", b"<Types/>")
        # Pad above the 1 KB FILE_TOO_SMALL floor so the safety pipeline gets
        # past the size check and actually evaluates the external-links rule.
        zf.writestr("padding.bin", b"\x00" * 2000)
    r = popunjeni_io.check_file_safety(p)
    assert not r.ok
    assert r.code == Codes.EXTERNAL_LINKS


def test_safety_accepts_legitimate_phase1_file(tmp_path: Path) -> None:
    file = _phase1_file(tmp_path, election=_make_election())
    r = popunjeni_io.check_file_safety(file)
    assert r.ok, f"Phase 1 output should pass safety; got {r.code}: {r.detail}"


# ── Marker reading ────────────────────────────────────────


def test_read_marker_from_phase1_file(tmp_path: Path) -> None:
    election = _make_election("abcd1234")
    file = _phase1_file(tmp_path, election=election)
    marker = popunjeni_io.read_marker(file)
    assert marker is not None
    assert marker.election_uuid == "abcd1234"
    assert marker.party  # safe-filename derived


def test_read_marker_missing_returns_none(tmp_path: Path) -> None:
    # Build a fresh workbook with no marker in M1.
    p = tmp_path / "no_marker.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    wb.save(p)
    assert popunjeni_io.read_marker(p) is None


def test_read_marker_malformed_returns_none(tmp_path: Path) -> None:
    p = tmp_path / "bad_marker.xlsx"
    wb = openpyxl.Workbook()
    wb.active[MARKER_CELL] = "not-the-right-shape"
    wb.save(p)
    assert popunjeni_io.read_marker(p) is None


# ── Full import pipeline ──────────────────────────────────


def test_summary_row_does_not_trip_marker_row_mismatch(tmp_path: Path) -> None:
    """Regression for v0.4.0 field bug.

    Phase 1 writes a summary row at `last_data_row + 2` with formulas in
    cols B (count), E (COUNTA), L (COUNTIF). When the file has been
    opened in Excel, those formulas' cached values get returned by
    `data_only=True`. Without filtering, the parser saw the summary
    row as a fake data row with party = COUNTIF result ("0" or "1"),
    polluting `row_parties` and tripping MARKER_ROW_MISMATCH.

    The fix is `if func not in FUNCTIONS: continue` inside
    parse_party_file. This test simulates Excel's caching by writing a
    number directly into the summary row's col L on a generated file,
    then importing — must accept, not reject.
    """
    import openpyxl
    election = _make_election()
    src = _phase1_file(tmp_path, election=election)
    # Open the just-generated file and write a numeric value into the
    # summary row's column L — same shape Excel would write after evaluating
    # the COUNTIF. last_data_row was len(slots)+3; summary_row = +2.
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    # The party file for ("СНС",) with 1 polling station has 6 slots →
    # data rows 4-9, summary row at 11. Col L = column 12.
    ws.cell(row=11, column=12, value=0)  # simulate cached COUNTIF result
    wb.save(src)
    wb.close()
    # Now import — it must NOT be rejected.
    popunjeni = tmp_path / "popunjeni"
    outcome = popunjeni_io.import_file(src, election, popunjeni)
    assert outcome.ok, (
        f"summary-row leak: expected accept, got {outcome.status} "
        f"({outcome.issue.code if outcome.issue else 'no issue'})"
    )


def test_parse_party_file_reads_prosireni_slot_functions(tmp_path: Path) -> None:
    """Field report: every real prosireni submission imported "successfully"
    but produced zero records, so the master stayed blank no matter what
    was filled in. Root cause: the same `if func not in FUNCTIONS: continue`
    guard above (added for the summary-row leak, see the previous test)
    checked ONLY the fixed stalni FUNCTIONS tuple -- prosireni's slot
    functions ("Члан [1]", "Заменик члана [1]", ...) never appear there,
    so every prosireni data row silently hit the same `continue` as the
    summary row it was meant to filter out. Confirmed via direct
    reproduction of the exact import_file -> parse_for_cache pipeline
    the GUI uses, not just this parser in isolation."""
    from biracki_odbor.config.candidacies import CandidaciesRegistry, Entity
    from biracki_odbor.core.phase1_generate import generate_entity_files

    election = _make_election()
    registry = CandidaciesRegistry(entities=[
        Entity(entity_id="ent1", name="Ентитет Један", entity_type="party",
               level="local", submission_order=1, linked_to_entity_id=""),
    ], candidates=[])
    result = generate_entity_files(
        registry=registry, bm_numbers=[1], output_dir=tmp_path / "sablon_prosireni",
        oik=_make_oik(1), election=election, run_ts=1,
    )
    assert result.files_written, "prosireni template generation should produce a file"
    tpl_file = result.files_written[0]

    wb = openpyxl.load_workbook(tpl_file)
    ws = wb.active
    ws.cell(row=4, column=5, value="Марко")
    ws.cell(row=4, column=6, value="Марковић")
    ws.cell(row=4, column=7, value="1605991774818")
    wb.save(tpl_file)
    wb.close()

    records, issues = popunjeni_io.parse_for_cache(tpl_file)
    assert len(records) == 2, "expected both Члан and Заменик члана slots as records"
    by_func = {func: person for (_, func), person in records}
    assert by_func["Члан [1]"].first_name == "Марко"
    assert by_func["Члан [1]"].last_name == "Марковић"
    assert by_func["Заменик члана [1]"].first_name == ""


# ── combined_name_header round-trip ───────────────────────────────


def _combined_phase1_file(tmp_path: Path) -> Path:
    election = _make_election()
    election.combined_name_header = True
    return _phase1_file(tmp_path, election=election)


def test_combined_name_header_auto_splits_two_word_name(tmp_path: Path) -> None:
    """Explicitly requested over the old cosmetic-only header grouping:
    the operator types one full name into the merged E:F box; a clean
    2-word name splits automatically with no issue raised."""
    path = _combined_phase1_file(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws.cell(row=4, column=5, value="Марко Марковић")
    ws.cell(row=4, column=7, value="1605991774818")
    wb.save(path)
    wb.close()

    records, issues = popunjeni_io.parse_party_file(path)
    assert not any(i.code == Codes.NAME_AMBIGUOUS_SPLIT for i in issues)
    person = records[0][1]
    assert person.first_name == "Марко"
    assert person.last_name == "Марковић"


def test_combined_name_header_single_word_has_no_surname(tmp_path: Path) -> None:
    path = _combined_phase1_file(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws.cell(row=4, column=5, value="Марко")
    wb.save(path)
    wb.close()

    records, issues = popunjeni_io.parse_party_file(path)
    person = records[0][1]
    assert person.first_name == "Марко"
    assert person.last_name == ""
    assert not any(i.code == Codes.NAME_AMBIGUOUS_SPLIT for i in issues)


def test_combined_name_header_three_plus_words_split_and_flagged(tmp_path: Path) -> None:
    """User-specified behavior: first word is the given name, every
    remaining word joins as the surname — never leave the surname
    blank, even past the unambiguous two-word case — but still flag it
    for a quick operator check in Корекције, since a compound GIVEN
    name ("Ана Марија Петровић") would split wrong the same way."""
    path = _combined_phase1_file(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws.cell(row=4, column=5, value="Петар Петар Петровић")
    wb.save(path)
    wb.close()

    records, issues = popunjeni_io.parse_party_file(path)
    person = records[0][1]
    assert person.first_name == "Петар"
    assert person.last_name == "Петар Петровић"
    ambiguous = [i for i in issues if i.code == Codes.NAME_AMBIGUOUS_SPLIT]
    assert len(ambiguous) == 1
    assert ambiguous[0].severity == "warning"
    assert "Петар Петар Петровић" in ambiguous[0].message_sr


def test_non_combined_file_still_parses_as_two_real_columns(tmp_path: Path) -> None:
    """Detection must be per-file (from the header row), not from
    whatever the election's current toggle happens to be — an ordinary
    split-column file must never be misread as a single merged name."""
    election = _make_election()
    election.combined_name_header = False
    path = _phase1_file(tmp_path, election=election)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws.cell(row=4, column=5, value="Марко")
    ws.cell(row=4, column=6, value="Марковић")
    wb.save(path)
    wb.close()

    records, issues = popunjeni_io.parse_party_file(path)
    person = records[0][1]
    assert person.first_name == "Марко"
    assert person.last_name == "Марковић"
    assert not any(i.code == Codes.NAME_AMBIGUOUS_SPLIT for i in issues)


def test_import_happy_path(tmp_path: Path) -> None:
    """v2 (2026-05-30): filenames now embed the received_at timestamp
    so multiple submissions per party can coexist on disk. Canonical
    name is `<party>__<iso8601-no-colons>.xlsx`."""
    election = _make_election()
    src = _phase1_file(tmp_path, election=election)
    popunjeni = tmp_path / "popunjeni"
    outcome = popunjeni_io.import_file(src, election, popunjeni)
    assert outcome.ok
    assert outcome.status == "imported"
    assert outcome.canonical_path is not None
    assert outcome.canonical_path.exists()
    assert outcome.canonical_path.parent == popunjeni
    # Timestamped filename: starts with the party + the `__` separator.
    assert outcome.canonical_path.name.startswith("СНС__")
    assert outcome.canonical_path.suffix == ".xlsx"
    assert outcome.party == "СНС"
    # Metadata travels back on the outcome so the cache update can
    # persist it without re-reading anything.
    assert outcome.received_time_iso  # non-empty ISO timestamp
    assert outcome.import_source == "file_picker"


def test_import_re_import_keeps_both(tmp_path: Path) -> None:
    """A second submission no longer OVERWRITES the first — both files
    coexist on disk so the accountability log can show the full
    submission history, and phase2_aggregate picks the latest by
    received_time_iso."""
    from datetime import datetime, timezone
    election = _make_election()
    src = _phase1_file(tmp_path, election=election)
    popunjeni = tmp_path / "popunjeni"
    # Pin distinct received_at values so the two imports get distinct
    # filenames (the default `now()` could land in the same second and
    # would force the microsecond suffix).
    t1 = datetime(2026, 5, 1, 10, 0, 0, tzinfo=timezone.utc)
    t2 = datetime(2026, 5, 2, 11, 30, 0, tzinfo=timezone.utc)
    o1 = popunjeni_io.import_file(src, election, popunjeni, received_at=t1)
    o2 = popunjeni_io.import_file(src, election, popunjeni, received_at=t2)
    assert o1.status == "imported" and o2.status == "imported"
    # Both files on disk; distinct names.
    assert o1.canonical_path != o2.canonical_path
    assert o1.canonical_path.exists()
    assert o2.canonical_path.exists()
    # The later import's received_time_iso sorts after the earlier one
    # (lexicographic ↔ chronological is the contract phase2 relies on).
    assert o1.received_time_iso < o2.received_time_iso


def test_import_rejects_no_marker(tmp_path: Path) -> None:
    p = tmp_path / "no_marker.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    wb.save(p)
    election = _make_election()
    popunjeni = tmp_path / "popunjeni"
    outcome = popunjeni_io.import_file(p, election, popunjeni)
    assert not outcome.ok
    assert outcome.status == "rejected_identity"
    assert outcome.issue is not None
    assert outcome.issue.code == Codes.NO_MARKER
    # Critically: no file copied into popunjeni/ on rejection.
    assert not popunjeni.exists() or not any(popunjeni.iterdir())


def test_import_rejects_wrong_election_uuid(tmp_path: Path) -> None:
    # Phase 1 file built for election A.
    election_a = _make_election("election-A")
    src = _phase1_file(tmp_path, election=election_a)
    # Import attempted against election B.
    election_b = _make_election("election-B")
    popunjeni = tmp_path / "popunjeni"
    outcome = popunjeni_io.import_file(src, election_b, popunjeni)
    assert not outcome.ok
    assert outcome.issue.code == Codes.WRONG_ELECTION


def test_import_rejects_safety_first(tmp_path: Path) -> None:
    """Safety pipeline must fire before identity validation —
    a malformed file should never reach the marker reader."""
    p = tmp_path / "fake.xlsx"
    p.write_bytes(b"not a zip" * 1000)
    election = _make_election()
    popunjeni = tmp_path / "popunjeni"
    outcome = popunjeni_io.import_file(p, election, popunjeni)
    assert outcome.status == "rejected_safety"
    assert outcome.issue.code == Codes.NOT_XLSX
    # And nothing was copied.
    assert not popunjeni.exists() or not any(popunjeni.iterdir())


def test_import_rejects_macro_file_without_copying(tmp_path: Path) -> None:
    p = tmp_path / "macro.xlsx"
    with zipfile.ZipFile(p, "w") as zf:
        zf.writestr("xl/vbaProject.bin", b"\x00" * 2000)
        zf.writestr("xl/workbook.xml", b"<workbook/>")
        zf.writestr("[Content_Types].xml", b"<Types/>")
    election = _make_election()
    popunjeni = tmp_path / "popunjeni"
    outcome = popunjeni_io.import_file(p, election, popunjeni)
    assert outcome.status == "rejected_safety"
    assert outcome.issue.code == Codes.MACROS_PRESENT
    assert not popunjeni.exists() or not any(popunjeni.iterdir())


# ── Phase 1 marker shape (regression on the writer side) ──


def test_phase1_writes_marker_to_M1(tmp_path: Path) -> None:
    election = _make_election("phase1-uuid")
    file = _phase1_file(tmp_path, election=election)
    wb = openpyxl.load_workbook(file)
    ws = wb.active
    val = ws[MARKER_CELL].value
    assert val is not None
    assert val.startswith("BO/")
    assert "phase1-uuid" in val
    assert ws[MARKER_CELL].protection.locked is True
    assert ws.column_dimensions[MARKER_COLUMN_LETTER].hidden is True
    wb.close()


def test_phase1_without_election_omits_marker(tmp_path: Path) -> None:
    """Legacy callers (no election arg) should not get a marker — keeps
    test fixtures and any pre-feature code paths working."""
    template = build_template(
        tmp_path / "tpl.xlsx",
        n_polling_stations=1, parties=("СНС",),
    )
    out = tmp_path / "stranke"
    r = generate_party_files(template, out, _make_oik(1))
    file = r.files_written[0]
    wb = openpyxl.load_workbook(file)
    ws = wb.active
    assert ws[MARKER_CELL].value is None
    wb.close()


def test_build_marker_shape() -> None:
    s = build_marker("uuid-x", "СНС")
    assert s == "BO/uuid-x/СНС"
