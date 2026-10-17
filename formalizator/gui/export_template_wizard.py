"""Define a cumulative export layout — its own screen, decoupled from
any specific calibrated template, so a column can be mapped to
whichever field(s) (built-in, or a custom field already used for that
document's own print positions) a document actually needs, joined
together however the operator wants, instead of new export code being
written per document. Opened from the setup screen's optional "Извоз у
Excel" step, right after a template is picked/created (see
gui.setup_screen).
"""

from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
from PySide6.QtWidgets import (
    QComboBox, QDialog, QDialogButtonBox, QFileDialog, QHBoxLayout,
    QHeaderView, QLabel, QLineEdit, QMessageBox, QPushButton, QSpinBox,
    QTableWidget, QVBoxLayout,
)

from biracki_odbor.gui.widgets.spinbox import inset_spinbox_text
from formalizator.core.models import (
    ROW_NUMBER_FIELD_KEY, CustomFieldDef, ExportColumnDef, ExportTemplateDef,
)
from formalizator.gui.template_wizard import (
    _FIELD_KEY_LABELS, _FIELD_KEYS, _JMBG_COMB_KEY, _size_combo_popup,
)

_BLANK_FIELD_LABEL = "— изаберите поље —"
_NO_SORT_LABEL = "(без сортирања)"
_MAX_HEADER_ROW_SCAN = 20


def read_header_row(path: Path | str, header_row: int = 1) -> list[str]:
    """`header_row`'s non-empty cell values, in column order — e.g. from
    an existing roster file the operator already has, so its real
    column headers seed the mapping table instead of being retyped by
    hand. Real official spreadsheets often don't have their headers on
    row 1 (title/boilerplate rows come first) — the caller is
    responsible for picking the right row (see `_guess_header_row`)."""
    return [text for _col, text in read_header_positions(path, header_row)]


def read_header_positions(path: Path | str, header_row: int = 1) -> list[tuple[int, str]]:
    """Like `read_header_row`, but keeps each header's real (1-based)
    column number alongside its text — export needs to write back to
    that exact column, not just the Nth mapped column, or skipping/
    reordering a row in the mapping table silently shifts every later
    column's data sideways relative to the file's own header row."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        row = next(ws.iter_rows(min_row=header_row, max_row=header_row), ())
        return [
            (cell.column, str(cell.value).strip())
            for cell in row if cell.value is not None and str(cell.value).strip()
        ]
    finally:
        wb.close()


def _guess_header_row(path: Path | str, max_scan: int = _MAX_HEADER_ROW_SCAN) -> int:
    """A real official spreadsheet's headers often sit a few rows down
    (title/boilerplate above them) — guess the first row with at least
    two non-empty cells, since a lone tag/title cell isn't a header
    row. The operator confirms/corrects this via the header-row spin
    box before the mapping is built, so a wrong guess here is just a
    starting point, not a silent mistake."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        for row_index, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1):
            non_empty = sum(1 for cell in row if cell is not None and str(cell).strip())
            if non_empty >= 2:
                return row_index
        return 1
    finally:
        wb.close()


def _field_options(custom_fields: list[CustomFieldDef]) -> list[tuple[str, str]]:
    """Every field a column could reasonably pull from: the built-in
    ID-card fields and whichever custom fields the active document's
    template already defines (its "универзално" ones are exactly the
    session-fixed values — a list name, a certifier, whatever a
    document needs — that used to be a handful of concepts hardcoded
    onto SessionInfo), plus the automatic row counter an official
    roster's "Редни број" column needs — the one column that isn't
    voter data at all.

    Exactly one JMBG entry, deliberately. The calibration wizard's
    `_JMBG_COMB_KEY` describes a *printing* layout — 13 separately
    positioned digit boxes on paper — which has no meaning in a
    spreadsheet, where one cell holds the whole number. Offering both
    put two identical-looking "ЈМБГ" choices in this picker, and
    picking the comb one silently produced a blank column (reported
    directly: a real export lost its JMBG that way)."""
    options = [
        (ROW_NUMBER_FIELD_KEY, "Редни број (аутоматски: 1, 2, 3…)"),
        ("jmbg", "ЈМБГ"),
    ]
    options += [
        (k, _FIELD_KEY_LABELS.get(k, k)) for k in _FIELD_KEYS if k != _JMBG_COMB_KEY
    ]
    options += [(cf.key, cf.label) for cf in custom_fields]
    return options


class _ColumnFieldsDialog(QDialog):
    """Builds one column's field join: an ordered list of fields, each
    with its own separator text placed before it (except the first).
    Entirely operator-driven — any number of fields, any separators —
    so a column like "municipality - certifier, certifier address" is
    just four rows here, not something this app hardcodes."""

    def __init__(
        self, custom_fields: list[CustomFieldDef],
        field_keys: list[str] | None = None, separators: list[str] | None = None,
        parent=None,
    ) -> None:
        super().__init__(parent)
        self.setWindowTitle("Састав колоне")
        self.resize(480, 360)
        self._custom_fields = custom_fields

        root = QVBoxLayout(self)
        root.addWidget(QLabel(
            "Изаберите једно или више поља. За свако поље после првог,"
            " можете уписати текст који га раздваја од претходног (нпр."
            " размак, зарез, цртица)."
        ))

        self.table = QTableWidget(0, 2)
        self.table.setHorizontalHeaderLabels(("Поље", "Раздвајач испред (осим за прво поље)"))
        header = self.table.horizontalHeader()
        header.setSectionResizeMode(0, QHeaderView.Stretch)
        header.setSectionResizeMode(1, QHeaderView.Stretch)
        root.addWidget(self.table, stretch=1)

        row_buttons = QHBoxLayout()
        btn_add = QPushButton("Додај поље")
        btn_add.clicked.connect(lambda: self._add_row())
        row_buttons.addWidget(btn_add)
        btn_remove = QPushButton("Уклони изабрано")
        btn_remove.clicked.connect(self._on_remove_row)
        row_buttons.addWidget(btn_remove)
        root.addLayout(row_buttons)

        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Ok).setText("У реду")
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        root.addWidget(buttons)

        field_keys = field_keys or []
        separators = separators or []
        if field_keys:
            for i, key in enumerate(field_keys):
                sep = separators[i - 1] if i > 0 and i - 1 < len(separators) else ""
                self._add_row(key, sep)
        else:
            self._add_row()

    def _add_row(self, field_key: str = "", separator: str = "") -> None:
        row = self.table.rowCount()
        self.table.insertRow(row)

        combo = QComboBox()
        combo.addItem(_BLANK_FIELD_LABEL, "")
        for key, label in _field_options(self._custom_fields):
            combo.addItem(label, key)
        idx = combo.findData(field_key)
        combo.setCurrentIndex(idx if idx >= 0 else 0)
        _size_combo_popup(combo)
        self.table.setCellWidget(row, 0, combo)

        ed_sep = QLineEdit(separator)
        ed_sep.setEnabled(row > 0)
        self.table.setCellWidget(row, 1, ed_sep)

    def _on_remove_row(self) -> None:
        row = self.table.currentRow()
        if row >= 0:
            self.table.removeRow(row)
            # Row 0's separator is meaningless (nothing comes before
            # the first field) — disable it if a new row landed at 0.
            if self.table.rowCount() > 0:
                self.table.cellWidget(0, 1).setEnabled(False)

    def result_fields(self) -> tuple[list[str], list[str]]:
        field_keys: list[str] = []
        separators: list[str] = []
        for row in range(self.table.rowCount()):
            combo: QComboBox = self.table.cellWidget(row, 0)
            key = combo.currentData()
            if not key:
                continue  # incomplete row — skip silently
            if field_keys:
                ed_sep: QLineEdit = self.table.cellWidget(row, 1)
                separators.append(ed_sep.text())
            field_keys.append(key)
        return field_keys, separators


class ExportTemplateDialog(QDialog):
    """One export layout: a name plus a list of columns, each built
    from one or more joined fields — each row of the resulting
    spreadsheet fills one column per voter using the exact same field
    resolution a print position uses."""

    def __init__(
        self, custom_fields: list[CustomFieldDef] | None = None,
        parent=None, home_dir: Path | str | None = None,
    ) -> None:
        super().__init__(parent)
        self.setWindowTitle("Извозни шаблон")
        self.resize(760, 560)
        self._custom_fields = custom_fields or []
        self._field_labels = dict(_field_options(self._custom_fields))
        self._home_dir = Path(home_dir) if home_dir is not None else None
        self._source_file: str = ""
        self._result: ExportTemplateDef | None = None

        root = QVBoxLayout(self)

        top = QHBoxLayout()
        top.addWidget(QLabel("Назив извозног шаблона:"))
        self.ed_name = QLineEdit()
        self.ed_name.setPlaceholderText("нпр. НПРС-4 (прилагођено)")
        top.addWidget(self.ed_name, stretch=1)
        root.addLayout(top)

        self.lbl_hint = QLabel(
            "Учитајте постојећи Excel фајл да преузмете његове наслове"
            " колона и да сваки извоз убацује нове редове испод заглавља"
            " у тај исти фајл (остатак фајла — нпр. напомене на дну —"
            " остаје нетакнут). Или додајте колоне ручно испод."
        )
        self.lbl_hint.setWordWrap(True)
        root.addWidget(self.lbl_hint)

        pick_row = QHBoxLayout()
        self.btn_pick_file = QPushButton("Учитај Excel документ…")
        self.btn_pick_file.clicked.connect(self._on_pick_file)
        pick_row.addWidget(self.btn_pick_file)
        pick_row.addWidget(QLabel("Ред заглавља:"))
        self.sb_header_row = QSpinBox()
        self.sb_header_row.setRange(1, 500)
        inset_spinbox_text(self.sb_header_row)
        self.sb_header_row.setEnabled(False)
        pick_row.addWidget(self.sb_header_row)
        self.btn_reload_headers = QPushButton("Поново учитај наслове")
        self.btn_reload_headers.setEnabled(False)
        self.btn_reload_headers.clicked.connect(self._on_reload_headers)
        pick_row.addWidget(self.btn_reload_headers)
        pick_row.addStretch(1)
        root.addLayout(pick_row)

        self.table = QTableWidget(0, 2)
        self.table.setHorizontalHeaderLabels(("Наслов колоне", "Поље"))
        header = self.table.horizontalHeader()
        header.setSectionResizeMode(0, QHeaderView.Stretch)
        header.setSectionResizeMode(1, QHeaderView.Stretch)
        root.addWidget(self.table, stretch=1)

        row_buttons = QHBoxLayout()
        self.btn_add = QPushButton("Додај колону")
        self.btn_add.clicked.connect(lambda: self._add_column())
        row_buttons.addWidget(self.btn_add)
        self.btn_remove = QPushButton("Уклони изабрано")
        self.btn_remove.clicked.connect(self._on_remove_column)
        row_buttons.addWidget(self.btn_remove)
        root.addLayout(row_buttons)

        sort_row = QHBoxLayout()
        sort_row.addWidget(QLabel("Сортирај прво по:"))
        self.cb_sort1 = QComboBox()
        sort_row.addWidget(self.cb_sort1, stretch=1)
        sort_row.addWidget(QLabel("затим по:"))
        self.cb_sort2 = QComboBox()
        sort_row.addWidget(self.cb_sort2, stretch=1)
        root.addLayout(sort_row)
        self._refresh_sort_columns()

        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Save).setText("Сачувај")
        buttons.accepted.connect(self._on_save)
        buttons.rejected.connect(self.reject)
        root.addWidget(buttons)

    def _on_pick_file(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self, "Изабери Excel документ", str(Path.home()),
            "Excel фајлови (*.xlsx)",
        )
        if not path:
            return
        guess = _guess_header_row(path)
        try:
            headers = read_header_positions(path, guess)
        except Exception as e:  # noqa: BLE001
            QMessageBox.critical(self, "Грешка", f"Не могу да прочитам фајл: {e}")
            return
        if not headers:
            QMessageBox.warning(
                self, "Нема наслова",
                "У изабраном реду фајла нису пронађени наслови колона.",
            )
            return
        self._source_file = path
        self.sb_header_row.blockSignals(True)
        self.sb_header_row.setValue(guess)
        self.sb_header_row.blockSignals(False)
        self.sb_header_row.setEnabled(True)
        self.btn_reload_headers.setEnabled(True)
        if not self.ed_name.text().strip():
            self.ed_name.setText(Path(path).stem)
        self.table.setRowCount(0)
        for col_index, header_text in headers:
            self._add_column(header_text, column_index=col_index)

    def _on_reload_headers(self) -> None:
        if not self._source_file:
            return
        try:
            headers = read_header_positions(self._source_file, self.sb_header_row.value())
        except Exception as e:  # noqa: BLE001
            QMessageBox.critical(self, "Грешка", f"Не могу да прочитам фајл: {e}")
            return
        if not headers:
            QMessageBox.warning(
                self, "Нема наслова",
                "У изабраном реду фајла нису пронађени наслови колона.",
            )
            return
        self.table.setRowCount(0)
        for col_index, header_text in headers:
            self._add_column(header_text, column_index=col_index)

    def _column_summary(self, field_keys: list[str], separators: list[str]) -> str:
        if not field_keys:
            return _BLANK_FIELD_LABEL
        labels = [self._field_labels.get(k, k) for k in field_keys]
        if len(labels) == 1:
            return labels[0]
        parts = [labels[0]]
        for i, label in enumerate(labels[1:]):
            sep = separators[i] if i < len(separators) else " "
            parts.append(f"[{sep}]" if sep else "[ ]")
            parts.append(label)
        return " ".join(parts)

    def _add_column(
        self, header_text: str = "",
        field_keys: list[str] | None = None, separators: list[str] | None = None,
        column_index: int = 0,
    ) -> None:
        row = self.table.rowCount()
        self.table.insertRow(row)

        ed_header = QLineEdit(header_text)
        ed_header._column_index = column_index  # noqa: SLF001
        ed_header.textChanged.connect(self._refresh_sort_columns)
        self.table.setCellWidget(row, 0, ed_header)

        btn_field = QPushButton()
        btn_field._field_keys = list(field_keys or [])  # noqa: SLF001
        btn_field._separators = list(separators or [])  # noqa: SLF001
        btn_field.setText(self._column_summary(btn_field._field_keys, btn_field._separators))
        btn_field.clicked.connect(lambda: self._on_edit_column_fields(btn_field))
        self.table.setCellWidget(row, 1, btn_field)

        self._refresh_sort_columns()

    def _on_edit_column_fields(self, btn_field: QPushButton) -> None:
        dlg = _ColumnFieldsDialog(
            self._custom_fields, btn_field._field_keys, btn_field._separators, self,  # noqa: SLF001
        )
        if dlg.exec() == dlg.DialogCode.Accepted:
            field_keys, separators = dlg.result_fields()
            btn_field._field_keys = field_keys  # noqa: SLF001
            btn_field._separators = separators  # noqa: SLF001
            btn_field.setText(self._column_summary(field_keys, separators))

    def _on_remove_column(self) -> None:
        row = self.table.currentRow()
        if row >= 0:
            self.table.removeRow(row)
        self._refresh_sort_columns()

    def _refresh_sort_columns(self) -> None:
        current1 = self.cb_sort1.currentData() if self.cb_sort1.count() else ""
        current2 = self.cb_sort2.currentData() if self.cb_sort2.count() else ""
        headers = []
        for row in range(self.table.rowCount()):
            ed_header: QLineEdit = self.table.cellWidget(row, 0)
            text = ed_header.text().strip()
            if text:
                headers.append(text)
        for combo, current in ((self.cb_sort1, current1), (self.cb_sort2, current2)):
            combo.blockSignals(True)
            combo.clear()
            combo.addItem(_NO_SORT_LABEL, "")
            for h in headers:
                combo.addItem(h, h)
            idx = combo.findData(current)
            combo.setCurrentIndex(idx if idx >= 0 else 0)
            combo.blockSignals(False)

    def _on_save(self) -> None:
        name = self.ed_name.text().strip()
        if not name:
            QMessageBox.warning(self, "Недостаје назив", "Унесите назив извозног шаблона.")
            return
        columns: list[ExportColumnDef] = []
        for row in range(self.table.rowCount()):
            ed_header: QLineEdit = self.table.cellWidget(row, 0)
            btn_field: QPushButton = self.table.cellWidget(row, 1)
            header_text = ed_header.text().strip()
            field_keys = btn_field._field_keys  # noqa: SLF001
            separators = btn_field._separators  # noqa: SLF001
            column_index = getattr(ed_header, "_column_index", 0)
            if not header_text or not field_keys:
                continue  # incomplete row — skip silently, same as the calibration wizard
            columns.append(ExportColumnDef(
                header=header_text, field_keys=field_keys, separators=separators,
                column_index=column_index,
            ))
        if not columns:
            QMessageBox.warning(self, "Нема колона", "Додајте бар једну колону пре чувања.")
            return

        source_file = self._source_file
        if source_file and self._home_dir is not None:
            try:
                dest_dir = self._home_dir / "export_sources"
                dest_dir.mkdir(parents=True, exist_ok=True)
                dest = dest_dir / f"{name}.xlsx"
                if Path(source_file).resolve() != dest.resolve():
                    shutil.copyfile(source_file, dest)
                source_file = str(dest)
            except OSError as e:
                QMessageBox.critical(self, "Грешка", f"Не могу да сачувам полазни фајл: {e}")
                return

        self._result = ExportTemplateDef(
            name=name, columns=columns,
            source_file=source_file, header_row=self.sb_header_row.value() if source_file else 1,
            sort_column_1=self.cb_sort1.currentData() or "",
            sort_column_2=self.cb_sort2.currentData() or "",
        )
        self.accept()

    def load_template(self, template: ExportTemplateDef) -> None:
        self.ed_name.setText(template.name)
        for column in template.columns:
            self._add_column(column.header, column.field_keys, column.separators, column.column_index)
        if template.source_file:
            self._source_file = template.source_file
            self.sb_header_row.setEnabled(True)
            self.sb_header_row.setValue(template.header_row)
            self.btn_reload_headers.setEnabled(True)
        self._refresh_sort_columns()
        if template.sort_column_1:
            idx = self.cb_sort1.findData(template.sort_column_1)
            if idx >= 0:
                self.cb_sort1.setCurrentIndex(idx)
        if template.sort_column_2:
            idx = self.cb_sort2.findData(template.sort_column_2)
            if idx >= 0:
                self.cb_sort2.setCurrentIndex(idx)

    def result_template(self) -> ExportTemplateDef | None:
        return self._result
