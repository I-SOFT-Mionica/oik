"""Export a session's entries using an operator-defined `ExportTemplateDef`
instead of a hardcoded roster shape.

Replaces an earlier, NPRS-4-specific hardcoded exporter: the operator
defines column headers and maps each one to a field (built-in,
session-level, or a custom field already used for that document's own
print positions) on its own screen (`gui.setup_screen`'s optional
"Извоз у Excel" step, which opens `gui.export_template_wizard`) —
producing an NPRS-4-shaped roster (or whatever a completely different
document's own roster needs) is just a matter of defining the matching
columns, not new Python code per document.

When the export template carries a `source_file` (a pristine copy of a
real official spreadsheet), every export starts from that untouched
file and only adds rows to its existing table, so each machine's
output keeps the exact same physical shape and a pile of them can be
merged by hand afterward.
"""

from __future__ import annotations

from copy import copy
from pathlib import Path
from typing import Any

import openpyxl

from formalizator.core.field_values import resolve_values
from formalizator.core.models import (
    ROW_NUMBER_FIELD_KEY, ExportColumnDef, ExportTemplateDef, SessionInfo, VoterEntry,
)


def _column_value(column: ExportColumnDef, values: dict) -> str:
    """A column's own resolved value — one field, or several joined in
    order with the operator's own separator text between each pair
    (`separators[i]` between `field_keys[i]` and `field_keys[i+1]`).
    Pure string concatenation of already-resolved values; the column
    has no idea what any of them "mean"."""
    parts = [str(values.get(k, "") or "") for k in column.field_keys]
    if not parts:
        return ""
    result = parts[0]
    for i in range(1, len(parts)):
        sep = column.separators[i - 1] if i - 1 < len(column.separators) else " "
        result += sep + parts[i]
    return result


def _normalized(text: Any) -> str:
    """Header text compared without caring about the line breaks real
    forms put inside a header cell (NPRS-4's own "Редни\\nброј")."""
    return " ".join(str(text or "").split())


def _column_positions(ws, export_template: ExportTemplateDef, header_row: int) -> list[int]:
    """Which real spreadsheet column each mapped column writes to.

    A recorded `column_index` (captured when the operator loaded the
    file) wins. Failing that the column's header text is matched
    against the file's own header row — which is what rescues export
    templates defined before positions were recorded at all. Only a
    column matching neither gets appended after the table."""
    header_columns: dict[str, int] = {}
    for cell in ws[header_row]:
        if cell.value is not None and str(cell.value).strip():
            header_columns.setdefault(_normalized(cell.value), cell.column)

    known = [c.column_index for c in export_template.columns if c.column_index > 0]
    next_free = max([*known, *header_columns.values()], default=0) + 1

    positions: list[int] = []
    for column in export_template.columns:
        if column.column_index > 0:
            positions.append(column.column_index)
            continue
        matched = header_columns.get(_normalized(column.header))
        if matched:
            positions.append(matched)
        else:
            positions.append(next_free)
            next_free += 1
    return positions


def _row_is_merged(ws, row: int) -> bool:
    return any(rng.min_row <= row <= rng.max_row for rng in ws.merged_cells.ranges)


def _is_spare_data_row(ws, row: int, max_col: int) -> bool:
    """Whether the row right under the header is the form's own ready
    (empty) data row rather than footer content.

    NPRS-4 supplies exactly one such row, already carrying its "Редни
    број" of 1 — the operator described it as "it only has row 1 for
    entry" — so the first entry belongs *in* it. Anything holding real
    text (a note, a signature line) is footer and must never be
    written over; a lone number is the form's own row counter, not
    content."""
    if _row_is_merged(ws, row):
        return False
    for col in range(1, max_col + 1):
        value = ws.cell(row=row, column=col).value
        if value is None or not str(value).strip():
            continue
        if isinstance(value, (int, float)) or str(value).strip().isdigit():
            continue  # the form's own pre-filled row number
        return False
    return True


def _shift_merges_down(ws, at_row: int, amount: int) -> None:
    """openpyxl's `insert_rows` moves cell values but leaves merged
    ranges exactly where they were (confirmed directly against the real
    NPRS-4 file: the note text moved from A5 to A8 while its A5:D6
    merge stayed put) — so every merged block at or below the
    insertion point has to be re-pinned by hand, or the whole footer
    (note, place/date, signature block, НАПОМЕНА) ends up merged over
    the wrong rows."""
    for rng in ws.merged_cells.ranges:
        if rng.min_row >= at_row:
            rng.shift(row_shift=amount)
        elif rng.max_row >= at_row:  # straddles the insert point — grows instead of moving
            rng.expand(down=amount)


def _insert_rows_styled_like(ws, at_row: int, amount: int, style_row: int, max_col: int) -> None:
    """Insert blank rows that look like `style_row`. Rows openpyxl
    inserts carry no formatting at all, so without this every entry
    after the first would print outside the form's own table borders."""
    ws.insert_rows(at_row, amount=amount)
    _shift_merges_down(ws, at_row, amount)
    height = ws.row_dimensions[style_row].height
    for offset in range(amount):
        target = at_row + offset
        if height is not None:
            ws.row_dimensions[target].height = height
        for col in range(1, max_col + 1):
            ws.cell(row=target, column=col)._style = copy(ws.cell(row=style_row, column=col)._style)  # noqa: SLF001


def export_roster(
    entries: list[VoterEntry], session_info: SessionInfo,
    export_template: ExportTemplateDef, out_path: Path | str,
) -> Path:
    columns = export_template.columns
    resolved = [resolve_values(entry, session_info) for entry in entries]

    sort_columns = []
    for header in (export_template.sort_column_1, export_template.sort_column_2):
        if not header:
            continue
        sort_columns += [c for c in columns if c.header == header][:1]
    if sort_columns:
        resolved.sort(key=lambda values: tuple(_column_value(c, values) for c in sort_columns))

    # Numbered only once the final order is settled, so "Редни број"
    # always reads 1..N straight down the finished page.
    for position, values in enumerate(resolved, start=1):
        values[ROW_NUMBER_FIELD_KEY] = str(position)

    rows = [[_column_value(c, values) for c in columns] for values in resolved]
    # A pure counter column is written as a real number, not text —
    # otherwise Excel flags every one of them as "number stored as text".
    counters = [c.field_keys == [ROW_NUMBER_FIELD_KEY] for c in columns]

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    source_file = Path(export_template.source_file) if export_template.source_file else None
    if source_file and source_file.exists():
        wb = openpyxl.load_workbook(source_file)
        ws = wb.active
        header_row = max(1, export_template.header_row)
        first_data_row = header_row + 1
        positions = _column_positions(ws, export_template, header_row)
        max_col = max([*positions, ws.max_column])

        if rows:
            if _is_spare_data_row(ws, first_data_row, max_col):
                # The form supplies one ready data row (NPRS-4's is even
                # pre-numbered "1"); the first entry goes into it and
                # only the rest are added underneath.
                if len(rows) > 1:
                    _insert_rows_styled_like(ws, first_data_row + 1, len(rows) - 1, first_data_row, max_col)
            else:
                # No spare data row — the header runs straight into
                # footer content, so nothing there may be overwritten.
                ws.insert_rows(first_data_row, amount=len(rows))
                _shift_merges_down(ws, first_data_row, len(rows))

        for offset, row_values in enumerate(rows):
            for col_index, value, is_counter in zip(positions, row_values, counters):
                if is_counter and value.isdigit():
                    value = int(value)
                ws.cell(row=first_data_row + offset, column=col_index, value=value)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = export_template.name[:31] or "Извоз"  # Excel sheet-name length limit
        for col_index, column in enumerate(columns, start=1):
            ws.cell(row=1, column=col_index, value=column.header)
        for row_offset, row_values in enumerate(rows, start=1):
            for col_index, (value, is_counter) in enumerate(zip(row_values, counters), start=1):
                if is_counter and value.isdigit():
                    value = int(value)
                ws.cell(row=1 + row_offset, column=col_index, value=value)

    wb.save(out_path)
    wb.close()
    return out_path
