"""Решење о престанку функције — the чл. 15 "automatic lapse" decision.

Confirmed against the official RIK-distributed form (biracki_odbor/new
forms/2 3 resenje o prestanku funkcije u prosirenom birackom odboru -
Srbija.docx) — structurally and legally distinct from every other
decision this app generates:
  - Different legal basis: чл. 14 став 1. тачка 8), чл. 15. и чл. 39.
    став 1. — not the чл. 32/35/39/41 family every other решење cites.
  - Different disposition order: "ступа на снагу" comes BEFORE
    "објављује се" — the opposite of every other decision here.
  - A 72-hour appeal window, not 48.

The scenario (confirmed with the user, not assumed from the text
alone): this is a NATIONAL election, and the same person can appear on
a party's list submission to MORE THAN ONE municipality at once — not
just on прошиени lists; nothing in чл. 15 restricts this to one
composition. Each OIK only sees its own data — there is no cross-
municipality duplicate detection possible, unlike the in-app JMBG
cross-check that already guards against the same person holding two
seats WITHIN one municipality's own стали/прошиени workflows. So
whichever of the person's two appointments came LATER automatically
lapses the moment it's noticed — чл. 15: "по сили закона му престаје
функција у оном органу у којем је касније именовано."

RIK's own template (form 2.3) only covers one direction — прошиени as
the lapsing side — but that's a reflection of NORMAL election
scheduling (стали is normally finalized days before прошиени, so
прошиени is structurally almost always the later of the two), not a
legal restriction: нothing in чл. 15 prevents a стали seat from being
the one that lapses. Confirmed with the user: this particular cycle
collapses that normal gap to zero (both deadlines land the same day),
so either direction is genuinely live, not just a hypothetical. Both
`ceasing_sastav` and `existing_sastav` are therefore free operator
choices — which one is "ceasing" is a fact RIK (or a direct comparison
of the two decisions) tells the operator, not something this module
determines; it only renders whichever two facts it's given correctly.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from biracki_odbor.config.commission import terms_for, vrsta_for
from biracki_odbor.config import declension
from biracki_odbor.config.election import ElectionConfig
from biracki_odbor.config.oik import OikConstants
from biracki_odbor.core.phase3_docx import (
    BODY_PT, FONT, TITLE_PT, _add_mixed_para, _add_para, _set_cell_borders, _twips,
    apply_document_format,
)


# "stalni" | "prosireni" -> the Serbian locative used in "у ... саставу".
SASTAV_LOCATIVE: dict[str, str] = {"stalni": "сталном", "prosireni": "проширеном"}

ROLE_CHOICES = ("члана", "заменика члана")

DEFAULT_LEGAL_BASIS = (
    'На основу члана 14. став 1. тачка 8), члана 15. и члана 39. став 1.'
    ' Закона о избору народних посланика ("Службени гласник РС", бр.'
    ' 14/22 и 47/26),'
)


@dataclass
class PrestanakSpec:
    """All facts needed for one "Решење о престанку функције" document.

    `ceasing_*` — the seat that lapses by operation of law (the LATER
    of the two appointments). `existing_*` — the earlier, continuing
    appointment that triggers the lapse. Both `*_sastav` are free
    choices ("stalni" | "prosireni") — see this module's docstring for
    why: RIK's own template only shows прошиени lapsing because that's
    normally the later of the two, not because the law requires it.
    """

    person_name: str  # nominative, e.g. "Петар Петровић"

    ceasing_role: str  # one of ROLE_CHOICES
    ceasing_sastav: str  # "stalni" | "prosireni"
    ceasing_bm: int
    ceasing_decision_broj: str
    ceasing_decision_date: str  # genitive phrase, e.g. "6. децембра 2026. године"

    existing_role: str  # one of ROLE_CHOICES
    existing_sastav: str  # "stalni" | "prosireni"
    existing_bm: int
    existing_decision_broj: str
    existing_decision_date: str

    municipality: str  # this decision's own municipality (the one issuing it)
    session_date: str  # genitive phrase for THIS decision's own session
    document_broj: str  # "Број: ..." for this decision itself

    # The existing (earlier) seat's OWN municipality — the whole reason
    # this document exists is that it's normally a DIFFERENT municipality
    # from the one issuing this decision (that's exactly why neither OIK
    # could detect the double appointment on its own — see module
    # docstring). Blank falls back to `municipality`, for the rarer
    # same-municipality case this module also supports (e.g. a стални/
    # проширени collision the JMBG cross-check already would have
    # caught some other way, but recorded through this same flow anyway).
    existing_municipality: str = ""
    legal_basis: str = ""  # empty -> DEFAULT_LEGAL_BASIS


def _set_normal_style(doc: Document, oik: OikConstants) -> None:
    section = doc.sections[0]
    section.page_width = _twips(11906)
    section.page_height = _twips(16838)
    section.top_margin = _twips(1134)
    section.right_margin = _twips(1134)
    section.bottom_margin = _twips(1701)
    section.left_margin = _twips(1134)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(BODY_PT)
    apply_document_format(doc, oik)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(attr), FONT)


def build_prestanak_document(
    spec: PrestanakSpec, oik: OikConstants, election: ElectionConfig,
) -> Document:
    mun_name = spec.municipality or oik.municipality.name
    own = oik.municipality
    terms = terms_for(
        vrsta_for(own.sifra_opstine), mun_name,
        own.name_genitive if mun_name == own.name else "",
        own.name_locative if mun_name == own.name else "",
        own.full_committee_name,
    )
    # The earlier appointment may be in another municipality whose type
    # we don't know -- keep the municipality wording for it then.
    other = spec.existing_municipality.strip()
    if other and other != mun_name:
        existing_in_unit = f"у општини {other}"
        existing_commission_gen = "Општинске изборне комисије"
    else:
        existing_in_unit = terms.in_unit
        existing_commission_gen = terms.subject_gen

    primary_type = election.election.types[0] if election.election.types else "parliamentary"
    election_phrase = declension.expand_type(primary_type, mun_name, terms.local_assembly)
    election_date = election.election.dates.date_full

    ceasing_sastav_loc = SASTAV_LOCATIVE.get(spec.ceasing_sastav, spec.ceasing_sastav)
    existing_sastav_loc = SASTAV_LOCATIVE.get(spec.existing_sastav, spec.existing_sastav)

    doc = Document()
    _set_normal_style(doc, oik)

    JUSTIFY = WD_ALIGN_PARAGRAPH.JUSTIFY
    CENTER = WD_ALIGN_PARAGRAPH.CENTER

    # ── Legal basis + session line ──────────────────────────────────
    legal_text = spec.legal_basis.strip() or DEFAULT_LEGAL_BASIS
    for line in legal_text.splitlines():
        if line.strip():
            _add_para(doc, line.strip(), align=JUSTIFY)

    _add_para(
        doc,
        f"{terms.preamble_subject}, на седници"
        f" одржаној {spec.session_date}, донела је",
        align=JUSTIFY, space_after_pt=10,
    )

    _add_para(doc, "Р Е Ш Е Њ Е", bold=True, size_pt=TITLE_PT,
              align=CENTER, space_before_pt=6, space_after_pt=12)

    # ── Disposition ──────────────────────────────────────────────────
    # Item 1 states the conclusion first (which seat lapses), then why
    # — mirrors the official form's own order exactly.
    _add_mixed_para(doc, [
        ("1. \tУТВРЂУЈЕ СЕ да ", False),
        (spec.person_name, True),
        (" по сили закона престаје функција ", False),
        (spec.ceasing_role, True),
        (f" бирачког одбора у {ceasing_sastav_loc} саставу за спровођење"
         f" {election_phrase} расписаних за {election_date} на бирачком"
         f" месту број ", False),
        (str(spec.ceasing_bm), True),
        (f" {terms.in_unit}, јер је именован за ", False),
        (spec.existing_role, True),
        (f" бирачког одбора у {existing_sastav_loc} саставу за спровођење"
         f" {election_phrase} расписаних за {election_date} на бирачком"
         f" месту број ", False),
        (str(spec.existing_bm), True),
        (f" {existing_in_unit}.", False),
    ], align=JUSTIFY)

    # Order matches the original exactly: ступа на снагу BEFORE
    # објављује се — the opposite of every other decision this app
    # generates. Not a typo to "fix" — this document is simply final
    # and effective immediately, independent of publication.
    _add_para(doc, "2. Ово решење ступа на снагу даном доношења.", align=JUSTIFY)
    _add_para(doc, (
        "3. Ово решење објављује се на веб-презентацији Републичке"
        " изборне комисије."
    ), align=JUSTIFY)

    # ── Образложење ──────────────────────────────────────────────────
    obr = _add_para(doc, "О б р а з л о ж е њ е", bold=True, align=CENTER, space_before_pt=12)
    obr.paragraph_format.page_break_before = True

    # Paragraphs 1-2 and the closing two paragraphs are the law's own
    # fixed wording — verbatim from the official form, not paraphrased.
    _add_para(doc, (
        "У складу са чланом 14. став 1. тачка 8) Закона о избору народних"
        " посланика, члану и заменику члана органа за спровођење избора"
        " функција престаје по сили закона, а орган надлежан за његово"
        " именовање по службеној дужности утврђује престанак његове"
        " функције у случајевима предвиђеним законом."
    ), align=JUSTIFY)
    _add_para(doc, (
        "У складу са чланом 15. Закона о избору народних посланика, исто"
        " лице не може бити члан двају органа за спровођење избора, те,"
        " ако је једно лице именовано у више органа за спровођење"
        " избора, по сили закона му престаје функција у оном органу у"
        " којем је касније именовано."
    ), align=JUSTIFY)

    # Chronological narrative — the EARLIER (existing) appointment
    # first, then "а да је, након тога" the LATER (ceasing) one. Same
    # facts as item 1, opposite order, matching the original.
    _add_mixed_para(doc, [
        (f"{terms.subject} је утврдила да је ", False),
        (spec.person_name, True),
        (f" Решењем {existing_commission_gen} ", False),
        (f"број {spec.existing_decision_broj} од {spec.existing_decision_date}", True),
        (" именован за ", False),
        (spec.existing_role, True),
        (f" бирачког одбора у {existing_sastav_loc} саставу на бирачком"
         f" месту број ", False),
        (str(spec.existing_bm), True),
        (f" {existing_in_unit}, а да је, након"
         f" тога, Решењем {terms.subject_gen} ", False),
        (f"број {spec.ceasing_decision_broj} од {spec.ceasing_decision_date}", True),
        (" именован за ", False),
        (spec.ceasing_role, True),
        (f" бирачког одбора у {ceasing_sastav_loc} саставу на бирачком месту број ", False),
        (str(spec.ceasing_bm), True),
        (f" {terms.in_unit}.", False),
    ], align=JUSTIFY)
    _add_para(doc, (
        f"Поводом наведеног, {terms.subject} је констатовала да"
        " је наступио случај из члана 15. Закона о избору народних"
        " посланика, те је, као надлежна за поступање, донела одлуку како"
        " је наведено у диспозитиву овог решења."
    ), align=JUSTIFY)
    _add_mixed_para(doc, [
        ("УПУТСТВО О ПРАВНОМ СРЕДСТВУ: Против овог решења ", False),
        ("именовани и подносилац проглашене изборне листе кандидата за"
         " народне посланике", True),
        (" могу поднети приговор Републичкој изборној комисији у року од"
         " 72 часа од објављивања овог решења на веб-презентацији"
         " Републичке изборне комисије.", False),
    ], align=JUSTIFY)

    # ── Signature block ────────────────────────────────────────────
    # Delegated-signer authorization — ONE shared setting used by every
    # resenje type (see config.election.DocumentInfo): if the chairman
    # has delegated signing authority for the period, the delegate signs
    # this cessation decision too. (Earlier revisions of this module
    # assumed only izmena documents could be delegate-signed — that
    # premise no longer holds now that Phase 3's own original decisions
    # support it too via the identical preamble pattern.)
    doc_info = election.document
    if doc_info.use_authorization and doc_info.authorized_name:
        signer_name = doc_info.authorized_name
        signer_role = doc_info.authorized_role or "овлашћено лице"
    else:
        signer_name = oik.chairman.name
        signer_role = oik.chairman.title or "председник"

    _add_para(doc, spec.document_broj, space_before_pt=12, align=JUSTIFY)
    try:
        prep = declension.place_preposition(mun_name)
        gen = declension.place_to_genitive(mun_name)
        place_line = f"{prep} {gen}, {spec.session_date}"
    except Exception:  # noqa: BLE001 — never let a declension edge case block the file
        place_line = f"У {mun_name}, {spec.session_date}"
    _add_para(doc, place_line, align=JUSTIFY)
    _add_para(doc, space_before_pt=24)
    for line in terms.signature_lines:
        _add_para(doc, line, align=CENTER)
    _add_para(doc, signer_role, align=CENTER)

    if signer_name:
        sig_tbl = doc.add_table(rows=1, cols=1)
        sig_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        sig_tbl.columns[0].width = _twips(2722)
        sig_cell = sig_tbl.rows[0].cells[0]
        sig_cell.width = _twips(2722)
        _set_cell_borders(sig_cell, top=False, bottom=True, left=False, right=False)
        sig_para = sig_cell.paragraphs[0]
        sig_para.paragraph_format.space_before = Pt(0)
        sig_para.paragraph_format.space_after = Pt(0)

    _add_para(doc, signer_name.upper(), bold=True, align=CENTER)

    return doc


def _safe_filename(name: str) -> str:
    cleaned = "".join(c if c.isalnum() or c in " _-" else "_" for c in name)
    return "_".join(cleaned.split()).strip("_") or "resenje"


def generate_prestanak(
    spec: PrestanakSpec, oik: OikConstants, election: ElectionConfig, output_dir: Path,
) -> Path:
    """Build + write one prestanak document. Returns the written path."""
    doc = build_prestanak_document(spec, oik, election)
    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / f"Prestanak_{_safe_filename(spec.person_name)}.docx"
    doc.save(out_path)
    return out_path
