"""Reads `Stalni_sastav.xlsx` (the RIK / municipality allocation template).

The template is provided by the OIK each election cycle. It contains:

  Sheet 'Raspodela':
    A header row with cell A containing 'БР БМ' (BM number),
    followed by 6 columns whose headers match `core.models.FUNCTIONS`
    (Председник, Заменик председника, Члан 1, Заменик члана 1,
     Члан 2, Заменик члана 2). Each data row holds the party
    abbreviation for that (BM, function) slot.

  Sheet 'bm' (optional in new world — OIK constants are authoritative):
    A header row (typically row 1), then per-BM rows with:
      column 1: BM number
      column 2: short name
      column 3: address
      column 4: full name (for решења)
    If present, we still read it and offer it as an initial seed
    for `oik_constants.polling_stations` on first install.

`build_blank_template` generates an empty version of this workbook so the
installer can drop a fill-in-the-blanks file into `sablon/` for users who
don't have the OIK-provided original yet.

Key change from legacy bo_system.py: column positions are discovered
from header text, not hard-coded. Sheet existence is checked
explicitly; missing sheets surface as ValidationReport issues, not
KeyErrors.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from biracki_odbor.config import candidacies as candidacies_mod
from biracki_odbor.core.xlsx_safety import safe_cell
from biracki_odbor.core.models import (
    FUNCTIONS,
    Allocation,
    Issue,
    PartyQuota,
    PollingStation,
    ValidationReport,
)
from biracki_odbor.validation.reports import Codes, message


SHEET_RASPODELA = "Raspodela"
SHEET_BM = "bm"

# Header text used when GENERATING a blank Raspodela sheet. Uses our
# canonical spelling (WITH spaces: "Заменик председника", "БР БМ") —
# NOT the real-world no-space RIK style ("Заменикпредседника",
# "Бројбирача") — because only the READER's tolerance (added 2026-07-10)
# understands both; already-deployed builds (v0.67.x and earlier) only
# recognize this canonical form. Once the tolerant reader has actually
# shipped in a release, it would be safe to switch the exporter to the
# real-world style too (see git history around 2026-07-10 for that
# attempt — reverted for exactly this reason, mid-client-test).
_RASPODELA_BM_HEADER = "БР БМ"
_RASPODELA_FUNCTION_HEADERS = FUNCTIONS


def build_blank_template(
    path: Path | str,
    *,
    n_polling_stations: int = 10,
    stations: list[PollingStation] | None = None,
    allocations: list[Allocation] | None = None,
) -> Path:
    """Write an empty Stalni_sastav.xlsx with the expected sheets and headers.

    Users fill in BM numbers, station info, and per-row party allocations
    themselves — or replace the file with the official OIK-provided version.
    The structure matches what `read_template` consumes.

    Pass `stations` (the OIK's actual configured polling stations) to seed
    both sheets with the REAL BM numbers/names instead of a synthetic
    1..N sequential range. Without this, a generated template whose BM
    numbers don't match `oik.polling_stations` (e.g. gaps, non-sequential
    numbering) round-trips straight into BM_UNKNOWN warnings at Phase 1 —
    the "there are BMs not in the list" complaint. `n_polling_stations` is
    kept only for the installer's first-run seed, before any OIK data exists.

    Pass `allocations` (from the in-app matrix builder, or a previously
    read template) to pre-fill Raspodela cells with actual party short
    names instead of leaving them blank — used when the operator edits
    the matrix in-app rather than round-tripping through Excel. Header
    text, column widths, and sheet layout are identical either way — the
    grid's shape is a legal constant (see test_legal_compliance.py),
    independent of whether cells happen to be filled at write time.
    """
    rows: list[tuple[int, str]]  # (bm_number, name)
    if stations:
        rows = [(s.bm_number, s.name) for s in sorted(stations, key=lambda s: s.bm_number)]
    else:
        rows = [(i + 1, "") for i in range(n_polling_stations)]

    alloc_by_slot: dict[tuple[int, str], str] = {
        (a.bm_number, a.function): a.party for a in (allocations or [])
    }

    wb = openpyxl.Workbook()
    bold = Font(bold=True)

    ws_bm = wb.active
    ws_bm.title = SHEET_BM
    for c, h in enumerate(("БР", "Назив", "Адреса", "Пун назив"), start=1):
        cell = ws_bm.cell(row=1, column=c, value=h)
        cell.font = bold
    for i, (bm_num, name) in enumerate(rows):
        ws_bm.cell(row=i + 2, column=1, value=bm_num)
        ws_bm.cell(row=i + 2, column=2, value=name)
    for col, width in zip("ABCD", (8, 24, 32, 36)):
        ws_bm.column_dimensions[col].width = width

    ws_r = wb.create_sheet(SHEET_RASPODELA)
    ws_r.cell(row=1, column=1, value=_RASPODELA_BM_HEADER).font = bold
    for c, func in enumerate(_RASPODELA_FUNCTION_HEADERS, start=2):
        cell = ws_r.cell(row=1, column=c, value=func)
        cell.font = bold
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
    for i, (bm_num, _name) in enumerate(rows):
        row = i + 2
        ws_r.cell(row=row, column=1, value=bm_num)
        for c, func in enumerate(_RASPODELA_FUNCTION_HEADERS, start=2):
            party = alloc_by_slot.get((bm_num, func))
            if party:
                ws_r.cell(row=row, column=c, value=party)
    ws_r.column_dimensions["A"].width = 8
    for c in range(2, 2 + len(FUNCTIONS)):
        ws_r.column_dimensions[get_column_letter(c)].width = 16
    ws_r.row_dimensions[1].height = 32

    p = Path(path)
    wb.save(p)
    wb.close()
    return p


@dataclass
class TemplateData:
    allocations: list[Allocation]
    polling_stations_from_template: list[PollingStation]  # optional seed
    parties: set[str]


def _normalize_header(v: object) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _collapse_whitespace(v: object) -> str:
    """Whitespace-insensitive comparison key.

    Real RIK-provided documents have been seen with headers like
    "Заменикпредседника" / "Бројбирача" — no internal space at all —
    where our canonical text has "Заменик председника" / expects "БР БМ".
    Collapsing all whitespace before comparing means either form matches
    the same slot, instead of a real allocation file silently failing
    HEADER_MISMATCH because of a missing space.
    """
    return "".join(_normalize_header(v).split())


_FUNCTIONS_BY_KEY = {_collapse_whitespace(f): f for f in FUNCTIONS}


def _find_header_row(ws, anchor_text: str = "БР БМ", max_scan: int = 200) -> int | None:
    """Find the header row by looking for the anchor in column A.

    Tolerant of real-world variants (e.g. "Бројбирача" instead of
    "БР БМ"): accepts the literal anchor, OR column-A text starting with
    "Број"/"БР" — the same prefix tolerance `read_bm_sheet` already
    applies to its own header check, applied here too since real RIK
    Raspodela sheets have been seen using that wording instead.

    max_scan defaults to 200, not a handful of rows: a real RIK-provided
    Stalni_sastav.xlsx puts a whole party-seat-quota summary table
    ("Посланичка група" / "Укупно места" etc., one row per party running,
    plus a "УКУПНО" total row and a blank separator) ABOVE the actual
    "БР БМ" matrix header — confirmed 2026-07-10 against a real file
    where that preamble pushed the real header down to row 20. The
    preamble's length varies with how many parties are running, so a
    small fixed scan window silently HEADER_MISMATCHed on any cycle with
    enough parties to push the header past it — this was the actual
    "BM not in the list" root cause, not a spacing/wording issue.
    """
    for r in range(1, min(ws.max_row, max_scan) + 1):
        v = ws.cell(row=r, column=1).value
        normalized = _normalize_header(v)
        if not normalized:
            continue
        if anchor_text in normalized or normalized.startswith(("Број", "БР")):
            return r
    return None


def _function_column_map(ws, header_row: int) -> dict[str, int]:
    """Return {function_name: column_index} by reading the header row."""
    mapping: dict[str, int] = {}
    for c in range(2, ws.max_column + 1):
        header = ws.cell(row=header_row, column=c).value
        canonical = _FUNCTIONS_BY_KEY.get(_collapse_whitespace(header))
        if canonical:
            mapping[canonical] = c
    return mapping


def read_template(template_path: Path | str) -> tuple[TemplateData, ValidationReport]:
    """Load the template. Returns (data, report). Never raises on user input."""
    report = ValidationReport()
    path = Path(template_path)

    if not path.exists():
        report.add(Issue(
            code=Codes.FILE_MISSING,
            message_sr=message(Codes.FILE_MISSING, file=str(path)),
            severity="error",
            file=str(path),
        ))
        return TemplateData([], [], set()), report

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except PermissionError:
        report.add(Issue(
            code=Codes.FILE_LOCKED,
            message_sr=message(Codes.FILE_LOCKED, file=str(path)),
            severity="error",
            file=str(path),
        ))
        return TemplateData([], [], set()), report

    try:
        # ── Raspodela ──
        if SHEET_RASPODELA not in wb.sheetnames:
            report.add(Issue(
                code=Codes.SHEET_MISSING,
                message_sr=message(Codes.SHEET_MISSING, file=path.name, sheet=SHEET_RASPODELA),
                severity="error",
                file=str(path), sheet=SHEET_RASPODELA,
            ))
            allocations: list[Allocation] = []
            parties: set[str] = set()
        else:
            allocations, parties = _read_raspodela(wb[SHEET_RASPODELA], path.name, report)

        # ── bm (optional) ──
        stations: list[PollingStation] = []
        if SHEET_BM in wb.sheetnames:
            stations = _read_bm(wb[SHEET_BM])
    finally:
        wb.close()

    return TemplateData(allocations, stations, parties), report


def _read_raspodela(
    ws,
    filename: str,
    report: ValidationReport,
) -> tuple[list[Allocation], set[str]]:
    header_row = _find_header_row(ws, "БР БМ")
    if header_row is None:
        report.add(Issue(
            code=Codes.HEADER_MISMATCH,
            message_sr=message(Codes.HEADER_MISMATCH, sheet=SHEET_RASPODELA, detail="БР БМ"),
            severity="error",
            file=filename, sheet=SHEET_RASPODELA, detail="БР БМ",
        ))
        return [], set()

    func_cols = _function_column_map(ws, header_row)
    missing = [f for f in FUNCTIONS if f not in func_cols]
    for m in missing:
        report.add(Issue(
            code=Codes.HEADER_MISMATCH,
            message_sr=message(Codes.HEADER_MISMATCH, sheet=SHEET_RASPODELA, detail=m),
            severity="error",
            file=filename, sheet=SHEET_RASPODELA, detail=m,
        ))
    if missing:
        return [], set()

    allocations: list[Allocation] = []
    parties: set[str] = set()
    for r in range(header_row + 1, ws.max_row + 1):
        raw_bm = ws.cell(row=r, column=1).value
        if raw_bm is None or (isinstance(raw_bm, str) and not raw_bm.strip()):
            break  # data ends at first empty BM number
        try:
            bm_num = int(raw_bm) if not isinstance(raw_bm, bool) else 0
        except (TypeError, ValueError):
            try:
                bm_num = int(float(raw_bm))
            except (TypeError, ValueError):
                report.add(Issue(
                    code=Codes.BM_NUMBER_NOT_INT,
                    message_sr=message(Codes.BM_NUMBER_NOT_INT, detail=str(raw_bm)),
                    severity="error",
                    file=filename, sheet=SHEET_RASPODELA, row=r, detail=str(raw_bm),
                ))
                continue

        for func, col in func_cols.items():
            party_cell = ws.cell(row=r, column=col).value
            party = (str(party_cell).strip() if party_cell is not None else "")
            if party:
                allocations.append(Allocation(
                    bm_number=bm_num, function=func, party=party,
                ))
                parties.add(party)

    if not allocations:
        report.add(Issue(
            code=Codes.NO_ALLOCATIONS,
            message_sr=message(Codes.NO_ALLOCATIONS),
            severity="error",
            file=filename, sheet=SHEET_RASPODELA,
        ))

    return allocations, parties


def _read_bm(ws) -> list[PollingStation]:
    """Best-effort read of the bm sheet for an initial seed."""
    stations: list[PollingStation] = []
    for r in range(2, ws.max_row + 1):
        raw_bm = ws.cell(row=r, column=1).value
        if raw_bm is None:
            break
        try:
            bm_num = int(raw_bm)
        except (TypeError, ValueError):
            try:
                bm_num = int(float(raw_bm))
            except (TypeError, ValueError):
                continue
        stations.append(PollingStation(
            bm_number=bm_num,
            name=str(ws.cell(row=r, column=2).value or "").strip(),
            address=str(ws.cell(row=r, column=3).value or "").strip(),
            full_name=str(ws.cell(row=r, column=4).value or "").strip(),
        ))
    return stations


def build_bm_workbook(
    stations: list[PollingStation],
    path: Path | str,
) -> Path:
    """Write a standalone bm.xlsx with the current polling stations.

    Single sheet, headers БР / Назив / Адреса / Пун назив, one row per
    station. Round-trips with `read_bm_sheet` for editing in Excel and
    importing back.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_BM
    bold = Font(bold=True)
    for c, h in enumerate(("БР", "Назив", "Адреса", "Пун назив"), start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = bold
    for i, ps in enumerate(stations, start=2):
        ws.cell(row=i, column=1, value=ps.bm_number)
        safe_cell(ws, i, 2, ps.name)
        safe_cell(ws, i, 3, ps.address)
        safe_cell(ws, i, 4, ps.full_name)
    for col, width in zip("ABCD", (8, 24, 32, 36)):
        ws.column_dimensions[col].width = width

    p = Path(path)
    wb.save(p)
    wb.close()
    return p


def build_parties_workbook(
    parties: dict[str, str],
    path: Path | str,
    *,
    contacts: dict | None = None,
) -> Path:
    """Write a standalone parties.xlsx.

    Four columns: Кратки назив / Пун назив / Е-пошта / Телефон. The
    email + phone columns are emitted empty when `contacts` is None;
    when provided, it's a dict[abbrev → PartyContact] (or anything
    with .email / .phone attrs). Round-trips with `read_parties_sheet`.
    `parties` should contain real parliamentary groups only — the
    nacelnik's display name lives in `OikConstants.nacelnik_full_name`,
    a separate field edited via its own UI control, not a row here.

    Older 2-column exports (pre-§3) still load via `read_parties_sheet`
    — the email/phone columns are treated as optional. New exports
    always emit all four columns.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "parties"
    bold = Font(bold=True)
    headers = ("Кратки назив", "Пун назив", "Е-пошта", "Телефон")
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = bold
    contacts = contacts or {}
    for i, (abbrev, full) in enumerate(parties.items(), start=2):
        safe_cell(ws, i, 1, abbrev)
        safe_cell(ws, i, 2, full)
        contact = contacts.get(abbrev)
        if contact is not None:
            safe_cell(ws, i, 3, getattr(contact, "email", "") or "")
            safe_cell(ws, i, 4, getattr(contact, "phone", "") or "")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 18

    p = Path(path)
    wb.save(p)
    wb.close()
    return p


def read_parties_sheet(
    path: Path | str,
) -> tuple[dict[str, str], dict[str, dict[str, str]], list[str]]:
    """Read party data from a parties.xlsx.

    Returns `(full_names, contacts, errors)`:
      - `full_names`: {abbrev: full_name}
      - `contacts`: {abbrev: {"email": ..., "phone": ...}} — only for
        rows that had at least one of email/phone populated. Empty for
        2-column legacy files.
      - `errors`: human-readable Cyrillic errors; empty on success.

    Requires a sheet named 'parties' OR an active sheet with the first
    two headers ('Кратки назив', 'Пун назив') in row 1. Email + phone
    headers are optional — older 2-column exports from before §3
    still load cleanly. Refuses to import an arbitrarily-shaped
    workbook so a misclick in the file picker can't silently corrupt
    the parties mapping.
    """
    p = Path(path)
    if not p.exists():
        return {}, {}, [f"Фајл не постоји: {p.name}"]
    try:
        wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    except PermissionError:
        return {}, {}, [f"Фајл је закључан: {p.name}"]
    except Exception as e:  # noqa: BLE001
        return {}, {}, [f"Не могу да отворим '{p.name}': {e}"]

    try:
        ws = wb["parties"] if "parties" in wb.sheetnames else wb.active
        if ws is None:
            return {}, {}, [f"Радни лист је празан у '{p.name}'"]

        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [str(c or "").strip() for c in first[:4]]
        if headers[:2] != ["Кратки назив", "Пун назив"]:
            return {}, {}, [
                f"Лист '{ws.title}' не изгледа као фајл странака."
                " Очекивано заглавље: A1='Кратки назив', B1='Пун назив'."
                f" Пронађено: {headers or '(празно)'}."
            ]

        result: dict[str, str] = {}
        contacts: dict[str, dict[str, str]] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            abbrev = (row[0] if len(row) > 0 else None)
            full = (row[1] if len(row) > 1 else None)
            email = (row[2] if len(row) > 2 else None)
            phone = (row[3] if len(row) > 3 else None)
            if abbrev is None:
                continue
            key = str(abbrev).strip()
            if not key:
                continue
            result[key] = str(full or "").strip()
            email_s = str(email or "").strip()
            phone_s = str(phone or "").strip()
            if email_s or phone_s:
                contacts[key] = {"email": email_s, "phone": phone_s}
        if not result:
            return {}, {}, [
                f"Нема валидних редова у листу '{ws.title}'"
                " (очекивано: колоне Кратки назив, Пун назив)."
            ]
        return result, contacts, []
    finally:
        wb.close()


def build_quotas_workbook(quotas: dict[str, PartyQuota], path: Path | str) -> Path:
    """Write a standalone quotas.xlsx.

    Four columns: Кратки назив / Председник / Заменик / Чл. + Зам. чл.
    (the same combined-member-bucket shape as `PartyQuota` — see its
    docstring in core/models.py). Round-trips with `read_quotas_sheet`.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "quotas"
    bold = Font(bold=True)
    headers = ("Кратки назив", "Председник", "Заменик", "Чл. + Зам. чл.")
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = bold
    for i, (abbrev, q) in enumerate(quotas.items(), start=2):
        ws.cell(row=i, column=1, value=abbrev)
        ws.cell(row=i, column=2, value=q.predsednik)
        ws.cell(row=i, column=3, value=q.zamenik)
        ws.cell(row=i, column=4, value=q.clan_i_zamenik)
    ws.column_dimensions["A"].width = 18
    for col in "BCD":
        ws.column_dimensions[col].width = 16

    p = Path(path)
    wb.save(p)
    wb.close()
    return p


def read_quotas_sheet(path: Path | str) -> tuple[dict[str, PartyQuota], list[str]]:
    """Read seat-quota data from a quotas.xlsx.

    Returns `(quotas, errors)` — `quotas` is `{abbrev: PartyQuota}`;
    `errors` are human-readable Cyrillic errors, empty on success.

    Requires a sheet named 'quotas' OR an active sheet whose first
    header is 'Кратки назив' (same tolerant-header pattern as
    `read_parties_sheet`). Non-numeric quota cells are treated as 0
    rather than rejecting the whole file — a manually-edited Excel
    file is more likely to have a stray blank cell than a genuinely
    invalid one.
    """
    p = Path(path)
    if not p.exists():
        return {}, [f"Фајл не постоји: {p.name}"]
    try:
        wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    except PermissionError:
        return {}, [f"Фајл је закључан: {p.name}"]
    except Exception as e:  # noqa: BLE001
        return {}, [f"Не могу да отворим '{p.name}': {e}"]

    try:
        ws = wb["quotas"] if "quotas" in wb.sheetnames else wb.active
        if ws is None:
            return {}, [f"Радни лист је празан у '{p.name}'"]

        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        header = str(first[0] if first else "" or "").strip()
        if header != "Кратки назив":
            return {}, [
                f"Лист '{ws.title}' не изгледа као фајл квота."
                " Очекивано заглавље: A1='Кратки назив'."
                f" Пронађено: {header or '(празно)'}."
            ]

        def _to_int(v: object) -> int:
            try:
                return int(v)  # type: ignore[arg-type]
            except (TypeError, ValueError):
                try:
                    return int(float(v))  # type: ignore[arg-type]
                except (TypeError, ValueError):
                    return 0

        result: dict[str, PartyQuota] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            abbrev = row[0] if len(row) > 0 else None
            if abbrev is None:
                continue
            key = str(abbrev).strip()
            if not key:
                continue
            result[key] = PartyQuota(
                predsednik=_to_int(row[1] if len(row) > 1 else 0),
                zamenik=_to_int(row[2] if len(row) > 2 else 0),
                clan_i_zamenik=_to_int(row[3] if len(row) > 3 else 0),
            )
        if not result:
            return {}, [
                f"Нема валидних редова у листу '{ws.title}'"
                " (очекивано: колона Кратки назив)."
            ]
        return result, []
    finally:
        wb.close()


def read_bm_sheet(path: Path | str) -> tuple[list[PollingStation], list[str]]:
    """Read polling stations from an Excel file's `bm` sheet (or active sheet).

    Prefers a sheet named 'bm'. Requires the first row to start with at
    least 'БР' (or 'Број') so we don't silently slurp a Stalni_sastav.xlsx
    Raspodela sheet or any other arbitrary 2-column workbook.

    Returns `(stations, error_messages)`. error_messages is Cyrillic.
    """
    p = Path(path)
    if not p.exists():
        return [], [f"Фајл не постоји: {p.name}"]
    try:
        wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    except PermissionError:
        return [], [f"Фајл је закључан: {p.name}"]
    except Exception as e:  # noqa: BLE001
        return [], [f"Не могу да отворим '{p.name}': {e}"]

    try:
        ws = wb[SHEET_BM] if SHEET_BM in wb.sheetnames else wb.active
        if ws is None:
            return [], [f"Радни лист је празан у '{p.name}'"]

        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [str(c or "").strip() for c in first[:4]]
        # Accept "БР", "Број", or "БМ" (e.g. "БМ #") in column A — first
        # token, since a standalone bm.xlsx exported by build_bm_workbook
        # uses "БР", but a real RIK-provided Stalni_sastav.xlsx 'bm' sheet
        # has been seen using "БМ #" instead (confirmed 2026-07-10 against
        # a real file — read_template()'s internal _read_bm() never
        # validated headers so this only affected the standalone
        # "Бирачка места → Увоз" import button, not the template upload).
        col_a = headers[0] if headers else ""
        if not col_a or col_a.split()[0] not in ("БР", "Број", "БМ"):
            return [], [
                f"Лист '{ws.title}' не изгледа као фајл бирачких места."
                " Очекивано: A1='БР' или 'Број', B1='Назив', C1='Адреса'."
                f" Пронађено: {headers or '(празно)'}."
            ]

        stations = _read_bm(ws)
        if not stations:
            return [], [
                f"Нема валидних редова у листу '{ws.title}' (очекивано: "
                f"колоне БР, Назив, Адреса, Пун назив са подацима од 2. реда)."
            ]
        return stations, []
    finally:
        wb.close()


# ── Candidacy registry I/O (prošireni-sastav A3) ──────────────────


_ENTITY_HEADERS = (
    "ID листе", "Назив", "Тип", "Ниво",
    "Редни бр. подношења", "Повезано са",
)
_CANDIDATE_HEADERS = ("ID листе", "Име и презиме", "ЈМБГ")


def build_candidacies_workbook(
    registry: "candidacies_mod.CandidaciesRegistry",
    path: Path | str,
) -> Path:
    """Write a standalone candidacies.xlsx with two sheets.

    Sheet 1 «Ентитети» (hidden): one row per (legal-entity × level).
    Columns cover the entity_id (operator slug, e.g.
    "sns-presidential"), the full ballot name, type (party / coalition
    / group_of_citizens), level (presidential / parliamentary /
    provincial / local), 1-based submission order at that level, and
    the optional cross-level link target. This is candidate data entry
    territory — Sheet 1 exists only to back the «ID листе» dropdown on
    Sheet 2, so it's hidden and Sheet 2 is what the operator actually
    sees on open. Still fully present and read normally by
    `read_candidacies_workbook` — hidden just means not the active tab.

    Sheet 2 «Кандидати» (active/visible): one row per person standing
    for office. Columns: entity_id (FK to Sheet 1, a dropdown of the
    already-known list IDs), full name, JMBG. The JMBG is the
    cross-check key that Phase B1 uses to block committee members who
    are also running.

    Round-trips with `read_candidacies_workbook` — exported and
    re-imported files preserve every field exactly.
    """
    wb = openpyxl.Workbook()

    # Sheet 1 — entities. `wb.active` initially holds the default
    # "Sheet" — rename it instead of creating + deleting.
    ws_e = wb.active
    ws_e.title = "Ентитети"
    bold = Font(bold=True)
    for c, h in enumerate(_ENTITY_HEADERS, start=1):
        cell = ws_e.cell(row=1, column=c, value=h)
        cell.font = bold
    for r, e in enumerate(registry.entities, start=2):
        ws_e.cell(row=r, column=1, value=e.entity_id)
        ws_e.cell(row=r, column=2, value=e.name)
        ws_e.cell(row=r, column=3, value=e.entity_type)
        ws_e.cell(row=r, column=4, value=e.level)
        ws_e.cell(row=r, column=5, value=e.submission_order)
        ws_e.cell(row=r, column=6, value=e.linked_to_entity_id)
    ws_e.column_dimensions["A"].width = 24
    ws_e.column_dimensions["B"].width = 48
    ws_e.column_dimensions["C"].width = 22
    ws_e.column_dimensions["D"].width = 18
    ws_e.column_dimensions["E"].width = 12
    ws_e.column_dimensions["F"].width = 24

    # Sheet 2 — candidates.
    ws_c = wb.create_sheet("Кандидати")
    for c, h in enumerate(_CANDIDATE_HEADERS, start=1):
        cell = ws_c.cell(row=1, column=c, value=h)
        cell.font = bold
    for r, cand in enumerate(registry.candidates, start=2):
        ws_c.cell(row=r, column=1, value=cand.endorsed_by_entity_id)
        ws_c.cell(row=r, column=2, value=cand.name)
        ws_c.cell(row=r, column=3, value=cand.jmbg)
    ws_c.column_dimensions["A"].width = 24
    ws_c.column_dimensions["B"].width = 36
    ws_c.column_dimensions["C"].width = 16

    n_entities = len(registry.entities)
    if n_entities:
        dv_entity = DataValidation(
            type="list",
            formula1=f"Ентитети!$A$2:$A${n_entities + 1}",
            allow_blank=True,
        )
        dv_entity.error = "Изаберите вредност из падајуће листе."
        dv_entity.errorTitle = "Неважећа вредност"
        ws_c.add_data_validation(dv_entity)
        dv_entity.add("A2:A1000")

    # Sheet 1 is reference data for Sheet 2's dropdown, not the primary
    # content — hide it and open the file on «Кандидати» instead.
    ws_e.sheet_state = "hidden"
    wb.active = wb.sheetnames.index("Кандидати")

    p = Path(path)
    wb.save(p)
    wb.close()
    return p


def read_candidacies_workbook(
    path: Path | str,
) -> tuple["candidacies_mod.CandidaciesRegistry", list[str]]:
    """Read a candidacies.xlsx (exported by `build_candidacies_workbook`
    or hand-prepared by the operator from RIK/PIK/OIK data).

    Returns `(registry, errors)`. Errors are Cyrillic and human-
    readable; non-fatal data issues (e.g. an unknown level value)
    still load the row and surface as an error so the operator can
    fix it. Fatal issues (file unreadable, wrong sheet layout) return
    an empty registry + an error list, mirroring `read_parties_sheet`.

    Refuses to import an arbitrarily-shaped workbook — the «Ентитети»
    + «Кандидати» sheet names + the canonical header rows must match,
    so a misclick on parties.xlsx can't silently corrupt the registry.
    """
    p = Path(path)
    if not p.exists():
        return candidacies_mod.CandidaciesRegistry(), [
            f"Фајл не постоји: {p.name}"
        ]
    try:
        wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    except PermissionError:
        return candidacies_mod.CandidaciesRegistry(), [
            f"Фајл је закључан: {p.name}"
        ]
    except Exception as e:  # noqa: BLE001
        return candidacies_mod.CandidaciesRegistry(), [
            f"Не могу да отворим '{p.name}': {e}"
        ]

    errors: list[str] = []
    entities: list[candidacies_mod.Entity] = []
    candidates: list[candidacies_mod.Candidate] = []
    try:
        if "Ентитети" not in wb.sheetnames:
            return candidacies_mod.CandidaciesRegistry(), [
                f"Фајл '{p.name}' нема лист 'Ентитети'."
            ]
        ws_e = wb["Ентитети"]
        first = next(ws_e.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [str(c or "").strip() for c in first[:len(_ENTITY_HEADERS)]]
        if headers != list(_ENTITY_HEADERS):
            return candidacies_mod.CandidaciesRegistry(), [
                f"Лист 'Ентитети' нема очекивано заглавље."
                f" Очекивано: {list(_ENTITY_HEADERS)}."
                f" Пронађено: {headers or '(празно)'}."
            ]
        for row_idx, row in enumerate(
            ws_e.iter_rows(min_row=2, values_only=True), start=2,
        ):
            if not row or row[0] is None:
                continue
            entity_id = str(row[0]).strip()
            if not entity_id:
                continue
            name = str(row[1] or "").strip() if len(row) > 1 else ""
            entity_type = str(row[2] or "").strip() if len(row) > 2 else ""
            level = str(row[3] or "").strip() if len(row) > 3 else ""
            order_raw = row[4] if len(row) > 4 else None
            linked = str(row[5] or "").strip() if len(row) > 5 else ""

            # Permissive int coerce — accept "1" / 1 / "" → 0.
            try:
                order = int(order_raw) if order_raw not in (None, "") else 0
            except (TypeError, ValueError):
                errors.append(
                    f"Ентитети, ред {row_idx}: 'Редни бр. подношења'"
                    f" није цео број ({order_raw!r})."
                )
                order = 0

            if entity_type and entity_type not in candidacies_mod.ENTITY_TYPES:
                errors.append(
                    f"Ентитети, ред {row_idx}: 'Тип' = {entity_type!r}"
                    f" није у дозвољеној листи {list(candidacies_mod.ENTITY_TYPES)}."
                )
            if level and level not in candidacies_mod.LEVEL_KEYS:
                errors.append(
                    f"Ентитети, ред {row_idx}: 'Ниво' = {level!r}"
                    f" није у дозвољеној листи {list(candidacies_mod.LEVEL_KEYS)}."
                )

            entities.append(candidacies_mod.Entity(
                entity_id=entity_id,
                name=name,
                entity_type=entity_type,
                level=level,
                submission_order=order,
                linked_to_entity_id=linked,
            ))

        if "Кандидати" in wb.sheetnames:
            ws_c = wb["Кандидати"]
            first_c = next(
                ws_c.iter_rows(min_row=1, max_row=1, values_only=True), (),
            )
            headers_c = [
                str(c or "").strip() for c in first_c[:len(_CANDIDATE_HEADERS)]
            ]
            if headers_c != list(_CANDIDATE_HEADERS):
                errors.append(
                    f"Лист 'Кандидати' нема очекивано заглавље."
                    f" Очекивано: {list(_CANDIDATE_HEADERS)}."
                    f" Пронађено: {headers_c or '(празно)'}."
                )
            else:
                for row_idx, row in enumerate(
                    ws_c.iter_rows(min_row=2, values_only=True), start=2,
                ):
                    if not row or row[0] is None:
                        continue
                    endorsed_by = str(row[0]).strip()
                    name = str(row[1] or "").strip() if len(row) > 1 else ""
                    # JMBGs often arrive as int (leading zero stripped).
                    # Pad via the canonical normalize helper.
                    from biracki_odbor.core.jmbg import normalize as jmbg_norm
                    jmbg = jmbg_norm(row[2]) if len(row) > 2 else ""
                    if not endorsed_by:
                        continue
                    candidates.append(candidacies_mod.Candidate(
                        name=name, jmbg=jmbg,
                        endorsed_by_entity_id=endorsed_by,
                    ))
        # Note: missing «Кандидати» sheet is not a fatal error — it's
        # legitimate to have entities recorded but no candidates yet
        # (e.g. mid-import when only the RIK entity decision is in
        # hand but the candidate lists haven't been processed).
    finally:
        wb.close()

    return (
        candidacies_mod.CandidaciesRegistry(
            entities=entities, candidates=candidates,
        ),
        errors,
    )


def build_entities_template(
    path: Path | str,
    registry: "candidacies_mod.CandidaciesRegistry | None" = None,
    *,
    available_levels: "Sequence[str] | None" = None,
    show_linked_column: bool = True,
) -> Path:
    """Write a fill-in-the-blanks template for the «Увези листе» import.

    Single sheet «Листе» — one row per electoral list (entity). Only
    «Назив» is required. «Тип» and «Ниво» are real Excel dropdowns
    (data validation) instead of free text, so operators can't typo
    past the exact values the reader expects.

    No «Кандидати» sheet: the reader paired with this template
    (`read_entities_sheet`, the «Увези листе» button) only ever reads
    the first sheet — writing a candidates sheet here would silently
    promise something this import path never delivers. Candidates
    have their own file — see `build_candidacies_workbook`.

    `registry` — when the operator already has lists entered (Учесници
    tab), pass those in: the sheet writes the ALREADY-ENTERED rows
    instead of the grey hint row, so "Generate template…" produces
    something the operator can actually hand off / re-import, not a
    blank form that silently discards what they already typed. The hint
    row only appears when there's nothing to show yet — operators
    starting from scratch still get the fill-in example, delete it
    before importing as before.

    `available_levels` — restrict the «Ниво» dropdown to the levels
    actually configured for this election cycle (e.g. municipalities
    outside Vojvodina never have "provincial"; a cycle with only local
    elections has nothing else to offer). Defaults to the full set of
    `candidacies_mod.LEVEL_KEYS` when not given.

    `show_linked_column` — «Повезано са» (cross-level linking) is only
    meaningful when 2+ election levels are active this cycle; the
    caller passes False to drop the column entirely for a single-level
    cycle instead of showing a field that can never do anything.
    """
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    hint_font = Font(italic=True, color="888888")
    hint_fill = PatternFill("solid", fgColor="F2F2F2")
    entities = registry.entities if registry is not None else []
    levels = list(available_levels) if available_levels else list(candidacies_mod.LEVEL_KEYS)
    entity_types = list(candidacies_mod.ENTITY_TYPES)

    # Hidden lookup sheet backing the «Тип»/«Ниво» dropdowns — Excel
    # list-validation against a long inline string hits a 255-char
    # formula limit, so the allowed values live in real cells instead
    # and the validations reference that range.
    ws_lu = wb.create_sheet("_lookup")
    ws_lu.sheet_state = "hidden"
    for i, t in enumerate(entity_types, start=1):
        ws_lu.cell(row=i, column=1, value=t)
    for i, lv in enumerate(levels, start=1):
        ws_lu.cell(row=i, column=2, value=lv)

    def _list_validation(col_letter: str, n: int) -> DataValidation:
        dv = DataValidation(
            type="list",
            formula1=f"_lookup!${col_letter}$1:${col_letter}${n}",
            allow_blank=True,
        )
        dv.error = "Изаберите вредност из падајуће листе."
        dv.errorTitle = "Неважећа вредност"
        return dv

    # ── Sheet 1: Листе ──────────────────────────────────────────────
    ws = wb.active
    ws.title = "Листе"

    entity_headers = list(_ENTITY_HEADERS)
    if not show_linked_column:
        entity_headers = entity_headers[:-1]  # drop «Повезано са»
    for c, h in enumerate(entity_headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = bold

    if entities:
        for r, e in enumerate(entities, start=2):
            ws.cell(row=r, column=1, value=e.entity_id)
            ws.cell(row=r, column=2, value=e.name)
            ws.cell(row=r, column=3, value=e.entity_type)
            ws.cell(row=r, column=4, value=e.level)
            ws.cell(row=r, column=5, value=e.submission_order)
            if show_linked_column:
                ws.cell(row=r, column=6, value=e.linked_to_entity_id)
    else:
        hint_values = [
            "(оставите празно — аутоматски)",
            "Пун назив листе (нпр. СРПСКА НАПРЕДНА СТРАНКА — АЛЕКСАНДАР ВУЧИЋ)",
            "(изаберите из падајуће листе)",
            "(изаберите из падајуће листе)",
            "1  (редни број на гласачком листићу)",
            "(ID повезане листе, или оставите празно)",
        ]
        if not show_linked_column:
            hint_values = hint_values[:-1]
        for c, v in enumerate(hint_values, start=1):
            cell = ws.cell(row=2, column=c, value=v)
            cell.font = hint_font
            cell.fill = hint_fill

    dv_type = _list_validation("A", len(entity_types))
    ws.add_data_validation(dv_type)
    dv_type.add("C2:C1000")
    dv_level = _list_validation("B", len(levels))
    ws.add_data_validation(dv_level)
    dv_level.add("D2:D1000")

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 56
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 14
    if show_linked_column:
        ws.column_dimensions["F"].width = 26

    p = Path(path)
    wb.save(p)
    wb.close()
    return p


def read_entities_sheet(
    path: Path | str,
    default_level: str = "",
) -> tuple[list, list[str]]:
    """Read electoral lists (entities) from the first sheet of any workbook.

    Lenient column detection by header keyword so it works on RIK-provided
    sheets without requiring the exact candidacies.xlsx layout.

    Required column: Назив / Naziv (list name).
    Optional columns — detected by keyword match (case-insensitive):
      ID       → entity_id  (auto-generated as 'lista-N' when absent/empty)
      Ниво     → level       (falls back to `default_level`)
      Тип      → entity_type (defaults to 'party')
      Редни    → submission_order
      Повезано → linked_to_entity_id

    Returns ``(entities, errors)``. Errors are non-fatal warnings; the
    row is still imported with sensible defaults when possible.
    """
    p = Path(path)
    if not p.exists():
        return [], [f"Фајл не постоји: {p.name}"]
    try:
        wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    except PermissionError:
        return [], [f"Фајл је закључан: {p.name}"]
    except Exception as e:  # noqa: BLE001
        return [], [f"Не могу да отворим '{p.name}': {e}"]

    errors: list[str] = []
    entities: list[candidacies_mod.Entity] = []
    try:
        ws = wb[wb.sheetnames[0]]
        first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [str(c or "").strip().lower() for c in first_row]

        def _col(*keywords: str) -> int | None:
            for i, h in enumerate(headers):
                if any(kw in h for kw in keywords):
                    return i
            return None

        col_name   = _col("назив", "naziv", "name", "пун назив")
        col_id     = _col("id")
        col_level  = _col("ниво", "nivo", "level")
        col_type   = _col("тип", "tip", "type")
        col_order  = _col("редни", "redni", "order", "submission", "бр. под")
        col_linked = _col("повезано", "linked", "connected")

        if col_name is None:
            return [], [
                f"Фајл '{p.name}': не могу да пронађем колону 'Назив'."
                " Прва врста мора имати заглавље са речима Назив/Naziv/Name."
            ]

        # Track generated IDs to guarantee uniqueness within this import.
        used_ids: set[str] = set()
        for row_idx, row in enumerate(
            ws.iter_rows(min_row=2, values_only=True), start=2,
        ):
            if not row:
                continue
            name = str(row[col_name] or "").strip() if col_name < len(row) else ""
            if not name:
                continue

            # entity_id — use sheet value or auto-generate.
            raw_id = (
                str(row[col_id] or "").strip()
                if col_id is not None and col_id < len(row)
                else ""
            )
            if raw_id:
                entity_id = raw_id
            else:
                # Generate a stable slug from row position so repeated
                # imports of the same sheet produce the same IDs.
                base = f"lista-{row_idx - 1}"
                entity_id = base
                n = 2
                while entity_id in used_ids:
                    entity_id = f"{base}-{n}"
                    n += 1
            used_ids.add(entity_id)

            level = (
                str(row[col_level] or "").strip()
                if col_level is not None and col_level < len(row)
                else ""
            ) or default_level
            if level and level not in candidacies_mod.LEVEL_KEYS:
                errors.append(
                    f"Ред {row_idx}: Ниво = {level!r} није у дозвољеној листи"
                    f" {list(candidacies_mod.LEVEL_KEYS)}."
                )

            entity_type = (
                str(row[col_type] or "").strip()
                if col_type is not None and col_type < len(row)
                else ""
            ) or "party"
            if entity_type not in candidacies_mod.ENTITY_TYPES:
                errors.append(
                    f"Ред {row_idx}: Тип = {entity_type!r} није у дозвољеној"
                    f" листи {list(candidacies_mod.ENTITY_TYPES)}."
                )

            order_raw = (
                row[col_order]
                if col_order is not None and col_order < len(row)
                else None
            )
            try:
                order = int(order_raw) if order_raw not in (None, "") else 0
            except (TypeError, ValueError):
                errors.append(
                    f"Ред {row_idx}: Редни бр. = {order_raw!r} није цео број."
                )
                order = 0

            linked = (
                str(row[col_linked] or "").strip()
                if col_linked is not None and col_linked < len(row)
                else ""
            )

            entities.append(candidacies_mod.Entity(
                entity_id=entity_id,
                name=name,
                entity_type=entity_type,
                level=level,
                submission_order=order,
                linked_to_entity_id=linked,
            ))
    finally:
        wb.close()

    return entities, errors
