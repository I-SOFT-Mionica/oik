"""Consume RIK's own multi-municipality Excel exports directly.

RIK distributes two files each cycle that already contain everything our
own blank-template workflow otherwise asks the operator to re-type:

  RASPODELA workbook — one giant sheet, one block of rows per
    municipality (`Управни округ:` / `Општина:` / `Шифра општине:`, a
    4-row party-quota mini-table, then a BM × function allocation grid in
    the exact same shape as our own `Stalni_sastav.xlsx` "Raspodela"
    sheet). Municipalities are told apart only by "шифра општине" — a
    name search is fragile (common short names collide as substrings),
    so every lookup here is by code, never by name.

  UKUPNO workbook — a small national summary. The only part we reuse is
    the "Напомена:" section at the bottom: official short↔long party
    name pairs, the same shape as `ElectionConfig.party_full_names`.

Design note: nothing in this module assumes a fixed number of
municipality blocks. The registry bundled with the app (`config/opstine.
py`) is simply whatever codes were present in the export we built it
from; matching a RASPODELA file at import time is a pure lookup — scan
whatever blocks that specific file actually contains for the one whose
code matches, and stop. Never validate against an expected count.

Real-world data-quality gotcha (confirmed by inspecting an actual 2023
RIK export, not a theoretical concern): one municipality's block used
`Општина` without the trailing colon, unlike every other block's
`Општина:`. Every marker match in this module is colon-tolerant for
exactly this reason — mirrors `core/template.py`'s own tolerance for
real-world header variance ("БР БМ" vs "Бројбирача").
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

import openpyxl

from biracki_odbor.core.models import Allocation, FUNCTIONS, PartyQuota, PollingStation

# The 3 quota-table row labels we read; "Укупно" (4th row) is a derived
# sum RIK includes for human readers — never stored, would be redundant
# with predsednik + zamenik + clan_i_zamenik. The 3rd row's label
# changed between cycles — "Чл.+Зам.чл." (2023 export) vs the spelled-
# out "Чланови +Заменици чланова" (fresh 2026 export, confirmed
# consistent across all 161 municipality blocks in that file) — so both
# are accepted rather than picking one and breaking the other cycle.
_QUOTA_ROW_LABELS = ("Председник", "Заменик")
_CLAN_ROW_LABELS = ("Чл.+Зам.чл.", "Чланови +Заменици чланова")


def _norm(v: object) -> str:
    return str(v).strip() if v is not None else ""


def _marker_value(row: tuple) -> str:
    """A marker row's own value (e.g. the code beside "Шифра општине:").
    Confirmed against a fresh 2026 RASPODELA export: RIK dropped the
    blank spacer column that used to sit between the label (column A)
    and its value — the value moved from column C into column B. Column
    B is checked first and column C is kept as a fallback rather than
    replaced outright, so an older archived file with the spacer intact
    still parses instead of silently reading an empty cell."""
    b = row[1] if len(row) > 1 else None
    c = row[2] if len(row) > 2 else None
    return _norm(b) or _norm(c)


def _collapse_ws(v: object) -> str:
    """Whitespace-insensitive comparison key — same rationale as
    `template.py`'s `_collapse_whitespace`: RIK's real header cells have
    embedded newlines (e.g. "Заменик\\nчлана 1") that must still match our
    canonical "Заменик члана 1"."""
    return "".join(_norm(v).split())


def _marker(v: object, label: str) -> bool:
    """True if `v` equals `label`, with or without a trailing colon."""
    s = _norm(v)
    return s == label or s == f"{label}:"


_FUNCTIONS_BY_KEY = {_collapse_ws(f): f for f in FUNCTIONS}


@dataclass
class OpstinaBlockRef:
    """One municipality's location inside a RASPODELA workbook."""

    sifra: str
    name: str
    okrug: str
    start_row: int  # 1-based row of this block's "Управни округ:" line


@dataclass
class RikMunicipalityImport:
    stations: list[PollingStation] = field(default_factory=list)
    allocations: list[Allocation] = field(default_factory=list)
    quotas: dict[str, PartyQuota] = field(default_factory=dict)
    # Short → long parliamentary group name, from the UKUPNO file's
    # «Напомена:» section — empty when no UKUPNO file was given.
    party_full_names: dict[str, str] = field(default_factory=dict)
    errors: list[str] = field(default_factory=list)


def _scan_blocks(ws) -> list[OpstinaBlockRef]:
    """One pass over the whole sheet, collecting every municipality
    block's starting position. Order of markers within a block is fixed
    (округ, шифра округа, општина, шифра општине) — RIK writes these
    mechanically, not by hand."""
    blocks: list[OpstinaBlockRef] = []
    cur_okrug = ""
    cur_start: int | None = None
    cur_name: str | None = None
    for row_i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        a = row[0] if row else None
        value = _marker_value(row) if row else ""
        if _marker(a, "Управни округ"):
            cur_okrug = value
            cur_start = row_i
            cur_name = None
        elif _marker(a, "Општина"):
            cur_name = value
        elif _marker(a, "Шифра општине") and cur_start is not None and cur_name:
            blocks.append(OpstinaBlockRef(
                sifra=value, name=cur_name, okrug=cur_okrug, start_row=cur_start,
            ))
    return blocks


def extract_opstina_registry(wb_path: Path | str) -> list[OpstinaBlockRef]:
    """Every municipality block found in `wb_path`, in file order.

    Used once to build the bundled `resources/opstine.json` registry, and
    reusable any time a fresh RIK export needs re-scanning (e.g. to spot-
    check a new cycle's file against the bundled list). Returns whatever
    the file actually contains — callers must not assume a fixed count.
    """
    wb = openpyxl.load_workbook(wb_path, data_only=True, read_only=True)
    try:
        ws = wb.worksheets[0]
        return _scan_blocks(ws)
    finally:
        wb.close()


def find_municipality_block(wb_path: Path | str, sifra: str) -> OpstinaBlockRef | None:
    """Locate the block for `sifra` in `wb_path`, or None if not present
    in this particular file."""
    sifra = _norm(sifra)
    for blk in extract_opstina_registry(wb_path):
        if blk.sifra == sifra:
            return blk
    return None


def _find_row_with_prefix(ws, start: int, end: int, prefix: str) -> int | None:
    for r in range(start, end + 1):
        v = _norm(ws.cell(row=r, column=1).value)
        if v.startswith(prefix):
            return r
    return None


def _find_row_equal(ws, start: int, end: int, label: str) -> int | None:
    for r in range(start, end + 1):
        if _norm(ws.cell(row=r, column=1).value) == label:
            return r
    return None


def _read_quota_table(ws, block_start: int, block_end: int) -> dict[str, PartyQuota]:
    predsednik_row = _find_row_equal(ws, block_start, block_end, "Председник")
    if predsednik_row is None or predsednik_row <= block_start:
        return {}
    header_row = predsednik_row - 1
    zamenik_row = predsednik_row + 1
    clan_row = predsednik_row + 2
    if _norm(ws.cell(row=zamenik_row, column=1).value) != "Заменик":
        return {}
    if _norm(ws.cell(row=clan_row, column=1).value) not in _CLAN_ROW_LABELS:
        return {}

    quotas: dict[str, PartyQuota] = {}
    max_col = ws.max_column
    # Column 2 (B), not 3 (C): RIK's 2026 export dropped the blank
    # spacer that used to sit between the row label and the first
    # party's column (see _marker_value). Starting at 2 stays correct
    # for an older file with the spacer intact too — the header row's
    # own column B is empty there, so `if not party: continue` just
    # skips it and column 3 onward still gets read.
    for col in range(2, max_col + 1):
        party = _norm(ws.cell(row=header_row, column=col).value)
        if not party:
            continue

        def _int_cell(row: int) -> int:
            v = ws.cell(row=row, column=col).value
            try:
                return int(v)
            except (TypeError, ValueError):
                return 0

        quotas[party] = PartyQuota(
            predsednik=_int_cell(predsednik_row),
            zamenik=_int_cell(zamenik_row),
            clan_i_zamenik=_int_cell(clan_row),
        )
    return quotas


def _read_allocation_grid(
    ws, block_start: int, block_end: int, errors: list[str],
) -> tuple[list[PollingStation], list[Allocation]]:
    header_row = _find_row_with_prefix(ws, block_start, block_end, "Број")
    if header_row is None:
        errors.append("Није пронађен ред заглавља расподеле („Број бирачког места“).")
        return [], []

    func_cols: dict[str, int] = {}
    for col in range(2, ws.max_column + 1):
        header = ws.cell(row=header_row, column=col).value
        canonical = _FUNCTIONS_BY_KEY.get(_collapse_ws(header))
        if canonical:
            func_cols[canonical] = col
    missing = [f for f in FUNCTIONS if f not in func_cols]
    if missing:
        errors.append(
            "Недостају колоне функција у мрежи расподеле: " + ", ".join(missing),
        )
        return [], []

    stations: list[PollingStation] = []
    allocations: list[Allocation] = []
    for r in range(header_row + 1, block_end + 1):
        raw_bm = ws.cell(row=r, column=1).value
        if raw_bm is None or (isinstance(raw_bm, str) and not raw_bm.strip()):
            break
        try:
            bm_num = int(raw_bm) if not isinstance(raw_bm, bool) else 0
        except (TypeError, ValueError):
            try:
                bm_num = int(float(raw_bm))
            except (TypeError, ValueError):
                errors.append(f"Ред {r}: број бирачког места '{raw_bm}' није цео број.")
                continue
        stations.append(PollingStation(bm_number=bm_num))
        for func, col in func_cols.items():
            cell = ws.cell(row=r, column=col).value
            party = _norm(cell)
            if party:
                allocations.append(Allocation(bm_number=bm_num, function=func, party=party))
    return stations, allocations


def read_municipality_allocation(
    wb_path: Path | str, sifra: str,
) -> RikMunicipalityImport:
    """Parse one municipality's block (quota table + allocation grid) out
    of a RASPODELA workbook, located purely by `sifra` — the caller is
    responsible for having a code (from the Опште tab picker) before
    calling this."""
    result = RikMunicipalityImport()
    wb = openpyxl.load_workbook(wb_path, data_only=True, read_only=True)
    try:
        ws = wb.worksheets[0]
        blocks = _scan_blocks(ws)
        target = next((b for b in blocks if b.sifra == _norm(sifra)), None)
        if target is None:
            result.errors.append(
                f"Општина са шифром '{sifra}' није пронађена у изабраном фајлу.",
            )
            return result
        idx = blocks.index(target)
        block_end = (blocks[idx + 1].start_row - 1) if idx + 1 < len(blocks) else ws.max_row

        result.quotas = _read_quota_table(ws, target.start_row, block_end)
        result.stations, result.allocations = _read_allocation_grid(
            ws, target.start_row, block_end, result.errors,
        )
        return result
    finally:
        wb.close()


_LEGEND_MARKERS = ("Напомена", "Легенда")


def read_party_full_names(ukupno_path: Path | str) -> dict[str, str]:
    """Short → long party name pairs from a RIK summary file's
    party-name legend section. Two cycles' shapes, both handled —
    confirmed against real files, not guessed:

    - Older UKUPNO workbook: marker "Напомена:" sits in column B, then
      short name in column B / long name in column C on each row after.
    - 2026 legend (ships two ways: standalone as its own file, e.g.
      "skracenice.xlsx", and as one sheet — "SRBIJA_UKUPNO", the 2nd of
      4 — inside the combined RASPODELA+summary+legend workbook): marker
      "Легенда:" sits in column A, short name in column A / long name in
      column B. Whichever column the marker itself is in is also where
      the short-name column starts — that invariant holds in both
      shapes, so the marker's own column drives the read instead of a
      hardcoded one.

    Every sheet is searched, not just the first: in the 2026 combined
    workbook the legend is the 2nd sheet, with the RASPODELA data (which
    has no legend at all) sitting first. An operator who points both the
    RASPODELA and UKUPNO file pickers at that same combined file must
    still get the real mapping, not a silent empty one from reading the
    wrong sheet.
    """
    wb = openpyxl.load_workbook(ukupno_path, data_only=True, read_only=True)
    try:
        for ws in wb.worksheets:
            marker_row: int | None = None
            marker_col: int | None = None
            for row_i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                for col_i, v in enumerate(row, start=1):
                    if any(_marker(v, m) for m in _LEGEND_MARKERS):
                        marker_row, marker_col = row_i, col_i
                        break
                if marker_row is not None:
                    break
            if marker_row is None:
                continue

            names: dict[str, str] = {}
            for r in range(marker_row + 1, ws.max_row + 1):
                short = _norm(ws.cell(row=r, column=marker_col).value)
                long_name = _norm(ws.cell(row=r, column=marker_col + 1).value)
                if not short and not long_name:
                    continue
                if short:
                    names[short] = long_name
            return names
        return {}
    finally:
        wb.close()


def import_municipality(
    raspodela_path: Path | str, sifra: str, *,
    ukupno_path: Path | str | None = None,
    progress: Callable[[str, int, int], None] | None = None,
) -> RikMunicipalityImport:
    """Orchestrates a full import for one municipality: parliamentary
    group names first (if a UKUPNO file was given), then the RASPODELA
    block itself. Meant to run off the GUI thread (see `gui.worker.
    start_worker`) — both files can be large enough that parsing them
    synchronously would freeze the window.

    `progress`, when given, is called with `(label, current, total)` —
    matches the signature `gui.worker.Worker` auto-detects and polls, so
    a caller wiring this through `start_worker` gets a real progress bar
    for free, no extra plumbing needed.
    """
    total_steps = 2 if ukupno_path else 1
    step = 0
    party_full_names: dict[str, str] = {}

    if ukupno_path:
        if progress:
            progress("Читање назива посланичких група...", step, total_steps)
        party_full_names = read_party_full_names(ukupno_path)
        step += 1

    if progress:
        progress("Читање расподеле по општини...", step, total_steps)
    result = read_municipality_allocation(raspodela_path, sifra)
    result.party_full_names = party_full_names

    if progress:
        progress("Готово", total_steps, total_steps)
    return result
