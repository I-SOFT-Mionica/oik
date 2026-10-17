"""Excel template export and import for RIK transitional-provision members.

One row per polling station:
  cols 1-2  : locked context (БМ број, Назив БМ)
  cols 3-7  : member data   (Члан — Име, ЈМБГ, Место/адреса, Тел, Е-пошта)
  cols 8-12 : deputy data   (Заменик — same five fields)
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter

from biracki_odbor.config.election import RikBmEntry, RikPerson
from biracki_odbor.config.oik import PollingStation
from biracki_odbor.core import phone as phone_mod
from biracki_odbor.core.xlsx_safety import safe_cell


# ── Style ─────────────────────────────────────────────────────────────

_HDR_FILL   = PatternFill("solid", fgColor="2F5496")
_LOCK_FILL  = PatternFill("solid", fgColor="D9E1F2")
_MEM_FILL   = PatternFill("solid", fgColor="EBF3FB")
_DEP_FILL   = PatternFill("solid", fgColor="FFF2CC")
_HDR_FONT   = Font(name="Arial", bold=True, size=10, color="FFFFFF")
_LOCK_FONT  = Font(name="Arial", size=10, color="2F5496")
_INPUT_FONT = Font(name="Arial", size=10)
_THIN       = Side(style="thin", color="BFBFBF")
_BORDER     = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_HEADER_ROW  = 1
_INSTRUCT_ROW = 2
_DATA_START  = 3

_COLS = [
    ("БМ број",              8,  "lock"),
    ("Назив БМ",             30, "lock"),
    ("Члан — Име и презиме", 26, "mem"),
    ("Члан — ЈМБГ",         16, "mem"),
    ("Члан — Место/адреса",  24, "mem"),
    ("Члан — Телефон",       16, "mem"),
    ("Члан — Е-пошта",       26, "mem"),
    ("Заменик — Име и презиме", 26, "dep"),
    ("Заменик — ЈМБГ",       16, "dep"),
    ("Заменик — Место/адреса", 24, "dep"),
    ("Заменик — Телефон",    16, "dep"),
    ("Заменик — Е-пошта",    26, "dep"),
]
_N_COLS = len(_COLS)


def _fill_for(kind: str) -> PatternFill:
    return {"lock": _LOCK_FILL, "mem": _MEM_FILL, "dep": _DEP_FILL}[kind]


def _font_for(kind: str) -> Font:
    return _LOCK_FONT if kind == "lock" else _INPUT_FONT


# ── Export ────────────────────────────────────────────────────────────

def export_rik_template(
    stations: list[PollingStation],
    existing: dict[int, RikBmEntry],
    out_path: Path,
) -> Path:
    """Write an xlsx template, pre-filling BM context and any existing member data."""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Чланови РИК-а"

    # Row 1 — headers
    for col_idx, (label, width, kind) in enumerate(_COLS, start=1):
        c = ws.cell(row=_HEADER_ROW, column=col_idx, value=label)
        c.font = _HDR_FONT
        c.fill = _HDR_FILL
        c.border = _BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[_HEADER_ROW].height = 28

    # Row 2 — instructions
    last_col = get_column_letter(_N_COLS)
    ws.merge_cells(f"A{_INSTRUCT_ROW}:{last_col}{_INSTRUCT_ROW}")
    instr = ws.cell(
        row=_INSTRUCT_ROW, column=1,
        value=(
            "Попуните поља за Члана и Заменика. "
            "Плава поља (БМ број, Назив БМ) не мењати — користе се за идентификацију. "
            "ЈМБГ: 13 цифара. Сачувајте фајл и увезите у апликацију."
        ),
    )
    instr.font = Font(name="Arial", size=9, italic=True, color="595959")
    instr.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[_INSTRUCT_ROW].height = 18

    # Data rows
    for rel_idx, station in enumerate(sorted(stations, key=lambda s: s.bm_number)):
        entry = existing.get(station.bm_number, RikBmEntry())
        m, d = entry.member, entry.deputy
        row_vals = [
            station.bm_number,
            station.name or "",
            m.name, m.jmbg, m.address, m.phone, m.email,
            d.name, d.jmbg, d.address, d.phone, d.email,
        ]
        row_idx = _DATA_START + rel_idx
        for col_idx, (val, (_, _, kind)) in enumerate(zip(row_vals, _COLS), start=1):
            c = safe_cell(ws, row_idx, col_idx, val)
            c.fill = _fill_for(kind)
            c.font = _font_for(kind)
            c.border = _BORDER
            c.alignment = Alignment(vertical="center")
            c.protection = Protection(locked=(kind == "lock"))

    ws.freeze_panes = f"A{_DATA_START}"
    ws.protection.sheet = True
    ws.protection.formatCells = False
    ws.protection.selectUnlockedCells = False

    wb.save(out_path)
    return out_path


# ── Import ────────────────────────────────────────────────────────────

def import_rik_xlsx(path: Path) -> tuple[dict[int, RikBmEntry], list[str]]:
    """Read a completed RIK template.

    Returns (members, errors). `errors` is non-empty when rows are
    skipped or malformed; callers should show these to the user.
    """
    path = Path(path)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        members: dict[int, RikBmEntry] = {}
        errors: list[str] = []

        for row in ws.iter_rows(min_row=_DATA_START, values_only=True):
            if not row or row[0] is None:
                continue
            try:
                bm_num = int(row[0])
            except (TypeError, ValueError):
                errors.append(f"Ред са неважећим бројем БМ: {row[0]!r} — прескочен")
                continue

            def _s(idx: int) -> str:
                v = row[idx] if idx < len(row) else None
                return str(v).strip() if v is not None else ""

            members[bm_num] = RikBmEntry(
                member=RikPerson(
                    name=_s(2), jmbg=_s(3), address=_s(4),
                    phone=phone_mod.normalize(row[5] if len(row) > 5 else None), email=_s(6),
                ),
                deputy=RikPerson(
                    name=_s(7), jmbg=_s(8), address=_s(9),
                    phone=phone_mod.normalize(row[10] if len(row) > 10 else None), email=_s(11),
                ),
            )
    finally:
        wb.close()

    return members, errors
